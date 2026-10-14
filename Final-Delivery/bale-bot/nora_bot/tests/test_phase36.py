# -*- coding: utf-8 -*-
"""
🧪 تست فاز ۳۶ — مرکز صدور · تاریخ شمسی · اکسل · عکس پرسنلی

🎯 گزارش شما:
   «قسمت آپلود اکسل و صدور گواهینامه کار نمیکنه، اکسل قبول نمیکنه»
   «چرا اطلاعات و تاریخ ... شمسی نیست؟ تاریخ ها خارجیه»
   «یه مرکز صدور گواهینامه تو قسمت کاربران اضافه کن»
   «همون لحظه برای همه ساخته نشه ... تا به سیستم فشار بیاد»
   «اگه پنل عکس نداشت ... یه تصویر آدمک ساده»
"""
import os
import sys
import json
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

os.environ.setdefault("BOT_TOKEN", "0:test")
os.environ.setdefault("SUPER_ADMIN_ID", "1")

_ok = _bad = 0
_fails = []


def check(name, cond, extra=""):
    global _ok, _bad
    if cond:
        _ok += 1
    else:
        _bad += 1
        _fails.append(name)
        print(f"  ❌ {name}  {extra}")


print("=" * 58)
print("🧪 تست فاز ۳۶ — مرکز صدور گواهینامه")
print("=" * 58)

# ══════════════════════════════════════════════════════════
print("\n[۱] 📅 ماژول تاریخ شمسی")
import jalali as J

check("امروز شمسی", J.today_str().startswith("۱۴"))
check("رفت‌وبرگشت تقویم",
      all(J.j2g(*J.g2j(*g)) == g for g in
          [(2026, 7, 29), (2024, 3, 20), (2000, 1, 1), (2025, 12, 31),
           (2020, 2, 29), (1999, 12, 31)]))
check("SQLite timestamp → شمسی",
      J.fmt("2026-07-29 14:32:11", "date") == "۱۴۰۵/۰۵/۰۷",
      J.fmt("2026-07-29 14:32:11", "date"))
check("ساعت هم می‌آید",
      J.fmt("2026-07-29 14:32:11", "datetime").endswith("۱۴:۳۲"))
check("سبک بلند", "مرداد" in J.fmt("2026-07-29", "long"))
check("ورودیِ از قبل شمسی دست نمی‌خورد",
      J.fmt("1405/05/07", "date") == "۱۴۰۵/۰۵/۰۷")
check("خالی → «—»", J.fmt("") == "—" and J.fmt(None) == "—")
check("زبالهٔ ورودی → «—»", J.fmt("سلام") == "—")
check("iso ارقام لاتین", J.fmt("2026-07-29", "iso") == "1405/05/07")
check("+۲۴ ماه", J.add_months("1405/05/07", 24) == "1407/05/07")
check("۳۱ اسفند → ۲۹/۳۰",
      J.add_months("1403/06/31", 6) in ("1403/12/30", "1403/12/29"))
check("+روز", J.add_days("1405/05/07", 10) == "1405/05/17")
check("اختلاف روز", J.diff_days("1405/05/17", "1405/05/07") == 10)
check("عدد مقایسه", J.to_num("1405/05/07") == 14050507)
check("اعتبارسنجی", J.valid("1405/05/07") and not J.valid("1405/13/01"))
check("۳۱ ماه دوم سال نامعتبر", not J.valid("1405/08/31"))
check("ارقام فارسی/لاتین",
      J.fa("123") == "۱۲۳" and J.en("۱۲۳") == "123")
check("روز هفته", J.weekday_name() in J.WEEKDAYS)
check("«پیش»", "پیش" in J.ago("2026-07-27 10:00:00")
      or J.ago("2026-07-27 10:00:00") == "همین حالا")
check("کبیسه", J.is_leap(1403) and not J.is_leap(1404))
check("روزهای ماه",
      J.month_days(1403, 1) == 31 and J.month_days(1403, 8) == 30)

# ── هیچ تاریخ میلادی خامی در رابط نماند ──
print("[۱ب] 🔍 تاریخ خام در رابط کاربر")
import re
import pathlib

RAW = re.compile(
    r"""\(\s*\w+\.get\(\s*['"](?:created_at|issued_at|paid_at|"""
    r"""approved_at|submitted_at|published_at)['"]\s*\)"""
    r"""\s*or\s*['"]['"]\s*\)\[:\d+\]""")
bad_files = []
for f in pathlib.Path(".").glob("*.py"):
    if f.name in ("jalali.py",):
        continue
    if RAW.search(f.read_text(encoding="utf-8")):
        bad_files.append(f.name)
check(f"هیچ تاریخ خامی نمانده ({bad_files})", not bad_files)

# ══════════════════════════════════════════════════════════
print("[۲] 📊 باگ اکسل")
import certbatch
import inspect as _ins

src = _ins.getsource(certbatch.handle_document)
check("دیگر download_file(file_id, dest) نیست",
      "download_file(doc.get(\"file_id\")" not in src)
check("از get_file_info استفاده می‌کند", "get_file_info" in src)
check("محتوا را می‌نویسد", "write_bytes" in src)

# اکسل واقعی بساز و بخوان
TMP = tempfile.mkdtemp(prefix="p36_")
try:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["نام و نام خانوادگی", "کد ملی", "موبایل", "نمره"])
    ws.append(["علی محمدی", "0012345679", "09121112233", "18"])
    ws.append(["زهرا کریمی", "0023456787", "09354445566", "20"])
    ws.append(["", "", "", ""])              # سطر خالی
    xl = f"{TMP}/list.xlsx"
    wb.save(xl)
    rows, cmap, errs = certbatch.parse_excel(xl)
    check("اکسل خوانده شد", len(rows) == 2, f"{len(rows)}")
    check("نام درست", rows and rows[0]["full_name"] == "علی محمدی")
    check("کد ملی درست", rows and rows[0]["national_id"] == "0012345679")
    check("ستون اضافه در extra",
          rows and "score" in (rows[0].get("extra") or {}))
    check("سطر خالی رد شد", len(rows) == 2)
except ImportError:
    check("openpyxl نصب است", False)

# ══════════════════════════════════════════════════════════
print("[۳] 📷 عکس پرسنلی")
import certphoto as cp

p = cp.make_silhouette(force=True)
check("آدمک ساخته شد", bool(p) and os.path.exists(p))
if p:
    from PIL import Image
    im = Image.open(p)
    check("مربع است", im.width == im.height)
    check("اندازهٔ معقول", im.width >= 300)

check("بدون کاربر → آدمک", bool(cp.photo_for(None)))
check("کاربر بی‌عکس → آدمک",
      cp.photo_for(None, {"profile_photo": None}) == p)
check("عکس تأییدنشده رد می‌شود",
      not cp.photo_approved({"profile_photo": p,
                             "profile_status": "pending"}))
check("عکس تأییدشده قبول",
      cp.photo_approved({"profile_photo": p,
                         "profile_status": "approved"}))
check("فایل گم‌شده رد می‌شود",
      not cp.photo_approved({"profile_photo": "/nope/x.png",
                             "profile_status": "approved"}))
_pp, _lbl, _real = cp.photo_state(None, {"profile_photo": None})
check("وضعیت خوانا", "آدمک" in _lbl and not _real)
check("۴ تنظیم عکس", len(cp.PHOTO_SETTINGS) == 4)
check("همه تنظیم‌ها توضیح دارند",
      all(r[5] for r in cp.PHOTO_SETTINGS if r[3] == "bool"))

# ══════════════════════════════════════════════════════════
print("[۴] 🏛️ مرکز صدور — ساختار")
import certcenter as cc
import database

database.init_db()
import certgen as _cg0
_cg0.init_cert_tables()          # جدول cert_templates لازم است
cc.init_tables()

for t in ("cert_issues", "cert_shares"):
    n = database.db_scalar(
        "SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name=?",
        (t,))
    check(f"جدول {t}", bool(n))

for fn in ("create_issue", "template_fields", "scan_template",
           "fill_from_event", "fill_from_segment", "fill_from_rows",
           "mapping_for", "build_one", "preview", "publish",
           "pending_for", "claim_by_link", "issue_ready", "link_of"):
    check(f"تابع {fn}", hasattr(cc, fn))

check("۴ راه انتخاب گیرنده", len(cc.AUDIENCES) == 4)
check("دسته‌های آماده", len(cc.SEGMENTS) >= 5)
check("تنظیمات مرکز کامل است", len(cc.CC_SETTINGS) >= 7,
      f"{len(cc.CC_SETTINGS)}")
check("تنظیم ساخت تنبل هست",
      any(r[0] == "cc_lazy_build" for r in cc.CC_SETTINGS))

# ══════════════════════════════════════════════════════════
print("[۵] 🏛️ چرخهٔ کامل")

# قالب آزمایشی
import certgen as cg
from docx.shared import Mm
from docx.enum.section import WD_ORIENT
import docx as _dx

d = _dx.Document()
s_ = d.sections[0]
s_.orientation = WD_ORIENT.LANDSCAPE
s_.page_width, s_.page_height = Mm(297), Mm(210)
d.add_paragraph("گواهی می‌شود {نام} با کد ملی {کدملی}")
d.add_paragraph("در دورهٔ {رویداد} به مدت {ساعت} ساعت")
d.add_paragraph("مدرس: {مدرس}   تاریخ: {تاریخ صدور}")
d.add_paragraph("{کیوآر}")
tpl_path = f"{TMP}/tpl.docx"
d.save(tpl_path)

tid = database.db_execute(
    "INSERT INTO cert_templates (name, file_path, kind, is_active) "
    "VALUES (?,?,?,1)", ("قالب تست فاز۳۶", tpl_path, "docx"))
check("قالب ثبت شد", bool(tid))

keys = cc.scan_template(tid)
check("پارامترها کشف شدند", len(keys) >= 5, f"{keys}")
ask, auto, imgs = cc.template_fields(tid)
check("«رویداد» پرسیدنی است", "event" in ask, f"ask={ask}")
check("«مدرس» پرسیدنی است", "teacher" in ask)
check("«نام» خودکار است", "name" in auto)
check("«کدملی» خودکار است", "national_id" in auto)
# 🔄 فاز ۳۹: «تاریخ صدور» عمداً از دستهٔ خودکار به دستهٔ
#    پرسیدنی منتقل شد.
#
#    🎯 «تو میگی برای پارامترهای ثابت مثل تاریخ برگزاری،
#        تاریخ صدور و شماره نامه ... من دیتا بدم»
#
#    ولی اختیاری است: خالی بگذارید، ربات خودش می‌سازد.
check("«تاریخ صدور» پرسیدنی شد",
      any(k in ask for k in ("issued_on", "issue_date")),
      f"ask={ask}")
check("«تاریخ صدور» از نوع سیستمی است",
      cc.field_kind("issued_on") == "system")
check("کیوآر تصویر است", "qr" in imgs)

iid = cc.create_issue(1, "طرح آزمایشی فاز ۳۶", tid)
check("طرح ساخته شد", bool(iid))
iss = cc.get_issue(iid)
check("لینک دارد", bool(iss.get("link_token")))
check("وضعیت پیش‌نویس", iss.get("status") == "draft")

ok, todo = cc.issue_ready(iss)
check("بدون پارامتر، آماده نیست", not ok)
check("دلیل روشن است", any("پارامتر" in t for t in todo), f"{todo}")

cc.set_value(iid, "event", "کارگاه امداد و نجات")
cc.set_value(iid, "teacher", "دکتر رضایی")
cc.set_value(iid, "hours", "۱۲")
# 🔄 فاز ۳۹: پارامترهای سیستمی هم در فهرست «پرسیدنی» هستند
#    (اختیاری‌اند ولی issue_ready پر بودنشان را می‌خواهد)
#    ⚠️ فقط ثابت‌ها لازم‌اند؛ سیستمی‌ها را خالی می‌گذاریم تا
#       ربات خودش بسازد.
_sysk = [k for k in ask if cc.field_kind(k) == "system"]
for _k in ask:
    if cc.field_kind(_k) == "system":
        continue
    if not str(cc.issue_values(cc.get_issue(iid)).get(_k) or "").strip():
        cc.set_value(iid, _k, "—")
iss = cc.get_issue(iid)
check("مقادیر ذخیره شدند",
      cc.issue_values(iss).get("event") == "کارگاه امداد و نجات")

ok, todo = cc.issue_ready(iss)
check("بدون گیرنده، هنوز آماده نیست", not ok)
check("دلیلش گیرنده است", any("گیرنده" in t for t in todo), f"{todo}")

# کاربران آزمایشی
for i, (nm, nid, mob) in enumerate([
        ("علی محمدی", "0012345679", "09121112233"),
        ("زهرا کریمی", "0023456787", "09354445566")], start=901):
    database.db_execute(
        "INSERT OR IGNORE INTO users (bale_user_id, first_name, "
        "last_name, national_id, mobile, profile_status) "
        "VALUES (?,?,?,?,?,'approved')",
        (i, nm.split()[0], nm.split()[1], nid, mob))

n = cc.fill_from_rows(iid, [
    {"full_name": "علی محمدی", "national_id": "0012345679",
     "mobile": "09121112233", "extra": {"نمره": "۱۸"}},
    {"full_name": "زهرا کریمی", "national_id": "0023456787",
     "mobile": "09354445566", "extra": {}},
    {"full_name": "مهمان بی‌حساب", "national_id": "0099999994",
     "mobile": "", "extra": {}},
])
check("۳ گیرنده اضافه شد", n == 3, f"{n}")
sh = cc.shares(iid)
check("۲ نفر به حساب وصل شدند",
      sum(1 for x in sh if x.get("bale_user_id")) == 2)
check("هرکس کد فردی دارد", all(x.get("person_code") for x in sh))
check("کدها یکتا هستند",
      len({x["person_code"] for x in sh}) == len(sh))

iss = cc.get_issue(iid)
ok, todo = cc.issue_ready(iss)
check("حالا آمادهٔ انتشار است", ok, f"{todo}")

# نقشهٔ پارامترها
mp, info = cc.mapping_for(iss, sh[0])
check("نام از کاربر", "علی" in str(mp.get("name")))
check("پارامتر مشترک", mp.get("event") == "کارگاه امداد و نجات")
check("پارامتر شخصی اکسل", str(mp.get("score") or "") == "۱۸"
      or str(mp.get("نمره") or "") == "۱۸")
check("تاریخ صدور شمسی", str(mp.get("issue_date", "")).startswith("۱۴"),
      str(mp.get("issue_date")))
check("کد فردی در نقشه", bool(info.get("person_code")))
check("اعتبار محاسبه شد", bool(info.get("expires_on")))
check("اعتبار شمسی است", str(info["expires_on"]).startswith("14"))

# ══════════════════════════════════════════════════════════
print("[۶] 🐢 ساخت تنبل — قلب درخواست")
n_before = database.db_scalar("SELECT COUNT(*) FROM certificates") or 0
sent, err = cc.publish(iid, notify=False)
n_after = database.db_scalar("SELECT COUNT(*) FROM certificates") or 0
check("انتشار موفق", not err, err)
check("🔑 هنگام انتشار هیچ فایلی ساخته نشد",
      n_after == n_before, f"{n_before}→{n_after}")
iss = cc.get_issue(iid)
check("وضعیت منتشرشده", iss.get("status") == "published")
check("زمان انتشار ثبت شد", bool(iss.get("published_at")))

pend = cc.pending_for(901)
check("کاربر گواهینامهٔ آماده دارد", len(pend) == 1, f"{len(pend)}")
check("عنوان طرح می‌آید", pend and pend[0].get("issue_title"))

# ══════════════════════════════════════════════════════════
print("[۷] 🎓 ساخت هنگام دریافت")
sh1 = [x for x in cc.shares(iid) if x.get("bale_user_id") == 901][0]
cert, err = cc.build_one(iid, sh1["id"], notify=False)
# 🧪 اگر موتور تبدیل موجود باشد فایل ساخته می‌شود؛ اگر نه، باید
#    خطای روشن بدهد. هر دو حالت را می‌سنجیم.
if cert:
    check("گواهینامه ساخته شد", True)
    check("به کاربر وصل است", cert.get("bale_user_id") == 901)
    check("کد فردی دارد", bool(cert.get("person_code")))
    check("منبع «مرکز»", cert.get("source") == "center")
    check("تاریخ صدور شمسی ذخیره شد",
          str(cert.get("issued_on") or "").startswith("14"))
    sh1b = database.db_execute("SELECT * FROM cert_shares WHERE id=?",
                               (sh1["id"],), fetch="one")
    check("وضعیت سهم = ساخته‌شده", sh1b.get("state") == "built")
    check("به گواهینامه لینک شد", sh1b.get("cert_id") == cert["id"])
    # دوبار ساخت → همان رکورد
    c2, _ = cc.build_one(iid, sh1["id"], notify=False)
    check("ساخت دوباره رکورد تازه نمی‌سازد",
          c2 and c2["id"] == cert["id"])
    n2 = database.db_scalar("SELECT COUNT(*) FROM certificates") or 0
    check("تعداد گواهینامه‌ها ثابت ماند", n2 == n_before + 1,
          f"{n_before + 1} vs {n2}")
    u = database.db_execute("SELECT total_certificates FROM users "
                            "WHERE bale_user_id=901", fetch="one")
    check("شمارندهٔ کاربر یک شد",
          (u or {}).get("total_certificates") == 1)
    pend2 = cc.pending_for(901)
    check("دیگر در فهرست انتظار نیست", len(pend2) == 0)
else:
    # 🧪 موتور تبدیل (LibreOffice) در این محیط نیست.
    #    مهم این است که **بی‌صدا نمیرد**: خطای روشن بدهد و
    #    وضعیت را «ناموفق» ثبت کند تا ادمین بفهمد.
    sh1b = database.db_execute("SELECT * FROM cert_shares WHERE id=?",
                               (sh1["id"],), fetch="one")
    check("ناموفق ولی با خطای روشن", bool(err), f"err={err!r}")
    check("وضعیت «ناموفق» ثبت شد",
          sh1b.get("state") == "failed", f"{sh1b.get('state')!r}")
    check("رکورد نیمه‌کاره نماند",
          (database.db_scalar("SELECT COUNT(*) FROM certificates")
           or 0) == n_before)

# ══════════════════════════════════════════════════════════
print("[۸] 🔗 لینک عمومی")
iss = cc.get_issue(iid)
tok = iss.get("link_token")
check("طرح با توکن پیدا می‌شود",
      (cc.by_token(tok) or {}).get("id") == iid)
check("توکن اشتباه چیزی نمی‌دهد", cc.by_token("nope") is None)

# ══════════════════════════════════════════════════════════
print("[۹] 🧩 استحکام")
check("طرح ناموجود → None", cc.get_issue(999999) is None)
check("شناسهٔ بی‌معنی → None", cc.get_issue("abc") is None)
_o, _t = cc.issue_ready(None)
check("طرح None امن است", not _o)
check("قالب ناموجود", cc.template_fields(999999) == ([], [], []))
check("پیش‌نمایش بدون طرح خطا نمی‌دهد",
      cc.preview(999999)[0] is None)
_c, _e = cc.build_one(999999, 1)
check("ساخت بدون طرح خطا نمی‌دهد", _c is None and bool(_e))
_n, _e2 = cc.publish(999999)
check("انتشار بدون طرح خطا نمی‌دهد", _n == 0 and bool(_e2))
check("گیرندهٔ تکراری دوبار ثبت نمی‌شود",
      cc.fill_from_rows(iid, [
          {"full_name": "علی محمدی", "national_id": "0012345679",
           "mobile": "09121112233"}]) == 0
      or len(cc.shares(iid)) == 3)

# ══════════════════════════════════════════════════════════
print("[۱۰] 🖥️ پنل")
for fn in ("show_home", "show_list", "start_new", "show_issue",
           "show_tpl_pick", "set_template", "show_params",
           "start_param", "show_audience", "pick_event",
           "pick_segment", "set_event", "set_segment", "ask_excel",
           "show_share_list", "show_preview", "confirm_publish",
           "do_publish", "renotify", "show_report", "show_link",
           "toggle_photo", "why_blocked", "confirm_delete",
           "do_delete", "user_get", "handle_text", "handle_document",
           "start_rename", "start_docno", "show_msg", "start_msg",
           "show_settings", "toggle_setting", "start_setting",
           "show_photo_panel", "toggle_photo_setting",
           "start_photo_setting", "show_silhouette", "show_stats"):
    check(f"تابع پنل {fn}", hasattr(cc, fn))

# روتر
main_src = pathlib.Path("main.py").read_text(encoding="utf-8")
check("مرکز در روتر هست", "cc_mod.show_home" in main_src)
check("جدول‌ها ساخته می‌شوند", "cc_mod.init_tables()" in main_src)
check("اکسل مرکز وصل است", 'state_name == "cc_excel"' in main_src)
check("متن مرکز وصل است", "cc_mod.handle_text" in main_src)
check("لینک عمیق وصل است", "cc_mod.claim_by_link" in main_src)
check("دکمهٔ دریافت کاربر", "ccget_" in main_src)
check("cb_home به مرکز می‌رود",
      '"cc_home", "cb_home"' in main_src)

cg_src = pathlib.Path("certgen.py").read_text(encoding="utf-8")
check("هاب دکمهٔ مرکز دارد",
      "🏛️ مرکز صدور گواهینامه" in cg_src)
check("هاب دکمهٔ عکس پرسنلی دارد",
      '"callback_data": "cc_photo"' in cg_src)

att_src = pathlib.Path("attendance.py").read_text(encoding="utf-8")
check("گواهینامه‌های من آماده‌ها را نشان می‌دهد",
      "pending_for" in att_src)

# ══════════════════════════════════════════════════════════
print("[۱۱] ⚙️ همه‌چیز از پنل قابل تغییر")
allset = list(cc.CC_SETTINGS) + list(cp.PHOTO_SETTINGS)
check("هیچ کلید تکراری نیست",
      len({r[0] for r in allset}) == len(allset))
check("همه نوع معتبر دارند",
      all(r[3] in ("bool", "num", "text") for r in allset))
check("همه برچسب فارسی دارند",
      all(r[2] and any("\u0600" <= c <= "\u06FF" for c in r[2])
          for r in allset))
check("همه پیش‌فرض دارند", all(r[4] is not None for r in allset))

print()
print("=" * 58)
print(f"✅ موفق: {_ok}   ❌ ناموفق: {_bad}")
if _fails:
    print("\nناموفق‌ها:")
    for f in _fails:
        print("  •", f)
print("=" * 58)
sys.exit(1 if _bad else 0)
