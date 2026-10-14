"""
تست فاز ۱۲ — لاگ فارسی، مالی، دسترسی ادمین، امتیاز سخت‌گیرانه

چهار درخواست کاربر:
  ۱) لاگ فعالیت ادمین به فارسی
  ۲) تراکنش و آمار مالی کامل با خروجی
  ۳) تعیین ادمین و نوع دسترسی پیشرفته
  ۴) نظام امتیازدهی سخت‌گیرانه
"""
import os
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

TMP = tempfile.mkdtemp(prefix="nora12_")
os.environ.update({
    "BOT_TOKEN": "1095412158:T", "SUPER_ADMIN_ID": "100000001",
    "SAFIR_API_KEY": "k", "SAFIR_BOT_ID": "317041514",
    "PAYMENT_PROVIDER_TOKEN": "WALLET-TEST-1111111111111111",
    "DB_PATH": f"{TMP}/t.db", "BACKUP_DIR": f"{TMP}/b", "IMPORT_DIR": f"{TMP}/i",
    "FILES_DIR": f"{TMP}/f", "PROFILES_DIR": f"{TMP}/f/p",
    "EVENTS_DIR": f"{TMP}/f/e", "LOG_DIR": f"{TMP}/l", "BROADCAST_RATE": "999",
})

import bale_api  # noqa: E402

SENT = []


def _api(m, p=None, **k):
    SENT.append({"m": m, "p": p or {}})
    if m == "getMe":
        return {"ok": True, "result": {"username": "nora_bot"}}
    return {"ok": True, "result": {"message_id": len(SENT) + 100}}


def _send(cid, text, kb=None, reply_to=None):
    SENT.append({"m": "sendMessage", "p": {"cid": cid, "text": text, "kb": kb}})
    return {"ok": True, "result": {"message_id": len(SENT) + 100}}


def _soe(cid, text, kb=None, mid=None, force_new=False):
    SENT.append({"m": "editMessageText", "p": {"cid": cid, "text": text, "kb": kb}})
    return {"ok": True, "result": {"message_id": mid or 1}}


def _file(method, key, cid, f, caption=None, kb=None):
    SENT.append({"m": method, "p": {"cid": cid, "caption": caption, "kb": kb}})
    return {"ok": True, "result": {"message_id": len(SENT) + 100}}


bale_api.api = _api
bale_api.api_async = _api
bale_api.send = _send
bale_api.send_or_edit = _soe
bale_api._send_file = _file
bale_api.send_photo = lambda c, p, cap=None, kb=None: _file("sendPhoto", "p", c, p, cap, kb)
bale_api.send_document = lambda c, d, cap=None, kb=None: _file("sendDocument", "d", c, d, cap, kb)

import config  # noqa: E402
import database  # noqa: E402
import auth  # noqa: E402
import forms as F  # noqa: E402
import payments as PAY  # noqa: E402
import events as EV  # noqa: E402
import attendance as AT  # noqa: E402
import occasions as OC  # noqa: E402
import permissions as PM  # noqa: E402
import activity_log as AL  # noqa: E402
import finance as FIN  # noqa: E402
import admin as AD  # noqa: E402
import main as M  # noqa: E402

for mod in (F, PAY, EV, AT, OC, AL, FIN, AD, M):
    for n in ("send", "send_or_edit", "send_photo", "send_document"):
        if hasattr(mod, n):
            setattr(mod, n, getattr(bale_api, n))

OK = FAIL = 0
FAILED = []


def check(name, cond, extra=""):
    global OK, FAIL
    if cond:
        OK += 1
        print(f"  ✅ {name}")
    else:
        FAIL += 1
        FAILED.append(name)
        print(f"  ❌ {name}   {extra}")


def texts():
    return [s["p"].get("text") or s["p"].get("caption") or "" for s in SENT]


def last():
    t = texts()
    return t[-1] if t else ""


def alltext():
    return " ".join(texts())


def cbdata():
    for s in reversed(SENT):
        kb = s["p"].get("kb") or {}
        out = [b.get("callback_data") for r in kb.get("inline_keyboard", [])
               for b in r if isinstance(b, dict) and b.get("callback_data")]
        if out:
            return out
    return []


def allcb():
    out = []
    for s in SENT:
        kb = s["p"].get("kb") or {}
        for r in kb.get("inline_keyboard", []):
            for b in r:
                if isinstance(b, dict) and b.get("callback_data"):
                    out.append(b["callback_data"])
    return out


def msg(uid, t):
    return {"message_id": 1, "from": {"id": uid, "first_name": "ت"},
            "chat": {"id": uid, "type": "private"}, "text": t}


def cb(uid, d):
    return {"id": "c", "from": {"id": uid, "first_name": "ت"}, "data": d,
            "message": {"message_id": 1, "chat": {"id": uid, "type": "private"}}}


S = 100000001
A_EV, A_FM, A_RP, A_FIN = 730001, 730002, 730003, 730004
U1, U2 = 740001, 740002

print("=" * 62)
print("🧪 تست فاز ۱۲")
print("=" * 62)

database.init_db()
EV.init_event_tables()
AT.init_attendance_tables()
PAY.init_payment_tables()
F.init_form_tables()
OC.init_occasion_tables()
PM.init_permission_tables()

auth.create_auth_user({"id": S, "first_name": "مدیر"}, "09120000000", "phone")
auth.ensure_super_admin(S)
for u, nm in ((A_EV, "رضا"), (A_FM, "مینا"), (A_RP, "کاوه"), (A_FIN, "لیلا"),
              (U1, "علی"), (U2, "سارا")):
    auth.create_auth_user({"id": u, "first_name": nm}, f"0913{u}", "phone")


# ═══════ ۱) لاگ فارسی ═══════
print("\n[۱] لاگ فعالیت فارسی")
check("جدول ترجمه دارد", len(AL.ACTION_FA) >= 40, str(len(AL.ACTION_FA)))
check("۸ دسته لاگ", len(AL.LOG_CATS) == 8, str(len(AL.LOG_CATS)))

_missing = []
for f in ROOT.glob("*.py"):
    import re
    src = f.read_text(encoding="utf-8")
    for m in re.finditer(r'log_action\([^,]+,\s*"([a-z_]+)"', src):
        if m.group(1) not in AL.ACTION_FA:
            _missing.append(m.group(1))
check("همه فعالیت‌های کد ترجمه دارند", not set(_missing),
      str(sorted(set(_missing))[:5]))

_icon, _name, _cat = AL.fa_action("event_created")
check("ترجمه event_created", _name == "ساخت رویداد", _name)
check("ترجمه دسته دارد", _cat == "رویداد")
_i2, _n2, _c2 = AL.fa_action("چیز_ناشناخته")
check("فعالیت ناشناس کرش نمی‌کند", isinstance(_n2, str))

AL.record(S, "event_created", {"event_id": 1})
AL.record(S, "form_publish", {"fid": 1})
AL.record(S, "receipt_approved", {"amount": 500000})
AL.record(U1, "form_submit", {"fid": 1})
AL.record(S, "admin_role", {"role": "event_admin"})
check("لاگ ثبت شد",
      database.db_scalar("SELECT COUNT(*) FROM user_actions") >= 5)

_row = database.db_execute("SELECT * FROM user_actions WHERE action_type="
                           "'receipt_approved' ORDER BY id DESC LIMIT 1",
                           fetch="one")
_desc = AL.describe(_row)
check("توصیف فارسی است", "تأیید رسید" in _desc, _desc)
check("مبلغ به ریال در توصیف", "500,000" in _desc, _desc)

_r2 = database.db_execute("SELECT * FROM user_actions WHERE action_type="
                          "'admin_role' ORDER BY id DESC LIMIT 1", fetch="one")
check("نقش در توصیف فارسی", "ادمین رویدادها" in AL.describe(_r2),
      AL.describe(_r2))

SENT.clear()
M.handle_callback(cb(S, "lg_list"))
check("صفحه لاگ", "لاگ فعالیت" in last())
check("متن فارسی نمایش", "ساخت رویداد" in alltext() or "تأیید رسید" in alltext())
check("انگلیسی خام نیست", "event_created" not in last())

SENT.clear()
M.handle_callback(cb(S, "lg_cats"))
check("فیلتر دسته", any(d.startswith("lg_c_") for d in cbdata()))
SENT.clear()
M.handle_callback(cb(S, "lg_c_رویداد_0"))
check("فیلتر رویداد", "رویداد" in last())
SENT.clear()
M.handle_callback(cb(S, "lg_adm_1"))
check("فیلتر فقط ادمین‌ها", "فقط ادمین" in last())
SENT.clear()
M.handle_callback(cb(S, "lg_stats"))
check("آمار فعالیت", "آمار فعالیت" in last())
check("پرتکرارترین‌ها", "پرتکرار" in last())
SENT.clear()
M.handle_callback(cb(S, "lg_export"))
check("خروجی لاگ", any(s["m"] == "sendDocument" for s in SENT))


# ═══════ ۲) مالی ═══════
print("\n[۲] مالی و تراکنش‌ها")
eid = database.db_execute("""INSERT INTO events (title,event_type,status,price,
    is_paid,created_by) VALUES ('کارگاه','workshop','published',500000,1,?)""", (S,))
eid2 = database.db_execute("""INSERT INTO events (title,event_type,status,price,
    is_paid,created_by) VALUES ('همایش','conference','published',1000000,1,?)""", (S,))
_tx = [(eid, U1, 500000, "paid"), (eid, U2, 500000, "paid"),
       (eid2, U1, 1000000, "paid"), (eid, U2, 300000, "pending"),
       (eid, U1, 200000, "failed"), (eid, U2, 100000, "refunded")]
for i, (e, u, amt, st) in enumerate(_tx):
    database.db_execute("""INSERT INTO transactions
        (bale_user_id,event_id,amount,status,payload)
        VALUES (?,?,?,?,?)""", (u, e, amt, st, f"ev{e}u{u}n{i}"))

check("تراکنش‌ها ثبت شد",
      database.db_scalar("SELECT COUNT(*) FROM transactions") == 6)
check("جمع پرداختی درست", FIN._sum() == 2000000, str(FIN._sum()))
check("تعداد موفق", FIN._count(f"status IN {FIN.PAID_STATES}") == 3)

SENT.clear()
M.handle_callback(cb(S, "fin_dash"))
_d = last()
check("داشبورد مالی", "داشبورد مالی" in _d)
check("درآمد به ریال", "2,000,000 ریال" in _d, _d[:120])
# فاز۱۷-د: همه‌جا ریال — همان عددی که در دیتابیس است
check("🔴 عدد ۱۰ برابر کوچک‌تر نیست", "200,000 ریال" not in _d, _d[:150])
check("تعداد موفق", "موفق: 3" in _d)
check("در انتظار", "در انتظار: 1" in _d)
check("ناموفق", "ناموفق: 1" in _d)
check("بازپرداخت", "بازپرداخت: 1" in _d)
check("نرخ موفقیت", "نرخ موفقیت" in _d)
check("میانگین تراکنش", "میانگین" in _d)

SENT.clear()
M.handle_callback(cb(S, "fin_tx_0_all"))
check("لیست تراکنش", "تراکنش‌ها" in last())
SENT.clear()
M.handle_callback(cb(S, "fin_tx_0_pending"))
check("فیلتر در انتظار", "در انتظار" in last())
SENT.clear()
M.handle_callback(cb(S, "fin_events"))
check("درآمد رویدادها", "تفکیک رویداد" in last())
# کارگاه ۲×۵۰۰k و همایش ۱×۱۰۰۰k → هر دو ۱٬۰۰۰٬۰۰۰ ریال (مساوی)
check("هر دو رویداد در گزارش", "کارگاه" in last() and "همایش" in last())
check("جمع درست", "2,000,000 ریال" in last(), last()[:90])
check("تعداد پرداخت نمایش", "2 پرداخت" in last() and "1 پرداخت" in last())
SENT.clear()
M.handle_callback(cb(S, "fin_methods"))
check("تفکیک روش پرداخت", "روش پرداخت" in last())
SENT.clear()
M.handle_callback(cb(S, "fin_pending"))
check("پرداخت‌های ناتمام", "ناتمام" in last())
check("مبلغ معلق", "300,000" in last(), last()[:100])
SENT.clear()
M.handle_callback(cb(S, "fin_forms"))
check("درآمد فرم‌ها", "تفکیک فرم" in last())

SENT.clear()
FIN.show_user_payments(S, S, U1)
check("پرداخت‌های یک کاربر", "پرداخت‌های" in last())
check("جمع کاربر درست", "1,500,000" in last(), last()[:130])

SENT.clear()
M.handle_callback(cb(S, "fin_export"))
check("خروجی مالی", any(s["m"] == "sendDocument" for s in SENT))
_exp = list((Path(TMP) / "f" / "exports").glob("finance_*.xlsx"))
check("فایل اکسل ساخته شد", len(_exp) >= 1)
if _exp:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(_exp[0])
        # 🆕 فاز ۳۰ — دو شیت تفکیکی اضافه شد
        check("حداقل ۴ شیت", len(wb.sheetnames) >= 4,
              str(wb.sheetnames))
        check("شیت تراکنش‌ها", "تراکنش‌ها" in wb.sheetnames)
        check("شیت خلاصه", "خلاصه مالی" in wb.sheetnames)
        check("راست‌چین", wb["تراکنش‌ها"].sheet_view.rightToLeft)
        # ستون‌ها با تفکیک کارمزد جابه‌جا شدند — با نام می‌گردیم
        _hdr = [c.value for c in wb["تراکنش‌ها"][1]]
        check("ستون مبلغ ریالی",
              any("ریال" in str(h) for h in _hdr), str(_hdr))
    except Exception as e:
        check("خواندن اکسل", False, str(e)[:60])

check("درصد تغییر مثبت", "▲" in FIN._pct(200, 100))
check("درصد تغییر منفی", "▼" in FIN._pct(50, 100))
check("بدون دوره قبل", "جدید" in FIN._pct(100, 0))
check("بدون تغییر", "بدون تغییر" in FIN._pct(100, 100))


# ═══════ ۳) دسترسی ادمین ═══════
print("\n[۳] نقش و دسترسی ادمین")
check("۸ نقش", len(PM.ROLES) == 8, str(len(PM.ROLES)))
check("۳۰+ دسترسی", len(PM.PERMISSIONS) >= 30, str(len(PM.PERMISSIONS)))
check("۸ دسته دسترسی", len(PM.PERM_CATS) == 8, str(PM.PERM_CATS))
_dup = len(PM.PERM_KEYS) != len(set(PM.PERM_KEYS))
check("کلید تکراری نیست", not _dup)

check("سوپرادمین همه‌چیز دارد", all(PM.has_perm(S, k) for k in PM.PERM_KEYS))
check("کاربر عادی هیچی ندارد", not PM.has_perm(U1, "ev_create"))

PM.set_role(A_EV, "event_admin", S)
check("نقش ادمین رویداد", PM.get_admin_role(A_EV)[0] == "event_admin")
check("✅ می‌تواند رویداد بسازد", PM.has_perm(A_EV, "ev_create"))
check("✅ حضور و غیاب دارد", PM.has_perm(A_EV, "attendance"))
check("🔴 فرم نمی‌تواند بسازد", not PM.has_perm(A_EV, "fm_create"))
check("🔴 مالی ندارد", not PM.has_perm(A_EV, "fin_approve"))
check("🔴 مدیریت ادمین ندارد", not PM.has_perm(A_EV, "sys_admins"))

PM.set_role(A_FM, "form_admin", S)
check("✅ ادمین فرم می‌سازد", PM.has_perm(A_FM, "fm_create"))
check("🔴 ادمین فرم رویداد نمی‌سازد", not PM.has_perm(A_FM, "ev_create"))

PM.set_role(A_RP, "report_admin", S)
check("✅ گزارش‌گیر می‌بیند", PM.has_perm(A_RP, "rp_view"))
check("✅ لاگ می‌بیند", PM.has_perm(A_RP, "rp_logs"))
check("🔴 گزارش‌گیر تغییر نمی‌دهد", not PM.has_perm(A_RP, "ev_create"))
check("🔴 گزارش‌گیر حذف نمی‌کند", not PM.has_perm(A_RP, "us_delete"))

PM.set_role(A_FIN, "finance_admin", S)
check("✅ ادمین مالی رسید تأیید می‌کند", PM.has_perm(A_FIN, "fin_approve"))
check("🔴 ادمین مالی فرم نمی‌سازد", not PM.has_perm(A_FIN, "fm_create"))

check("🔒 دسترسی سیستمی فقط سوپرادمین",
      all(not PM.has_perm(a, "sys_settings")
          for a in (A_EV, A_FM, A_RP, A_FIN)))
PM.set_role(A_EV, "custom", S, perms=["sys_admins", "ev_create"])
check("🔒 sys_admins حتی با custom داده نمی‌شود",
      not PM.has_perm(A_EV, "sys_admins"))
check("✅ بقیه custom کار می‌کند", PM.has_perm(A_EV, "ev_create"))

PM.toggle_perm(A_EV, "fm_create", S)
check("افزودن دسترسی تکی", PM.has_perm(A_EV, "fm_create"))
PM.toggle_perm(A_EV, "fm_create", S)
check("حذف دسترسی تکی", not PM.has_perm(A_EV, "fm_create"))
check("نقش به custom تغییر کرد", PM.get_admin_role(A_EV)[0] == "custom")

check("برچسب نقش", "ادمین فرم" in PM.role_label(A_FM), PM.role_label(A_FM))
check("برچسب سوپرادمین", "سوپر" in PM.role_label(S))
check("برچسب کاربر عادی", "عادی" in PM.role_label(U2))

# گارد عملی
SENT.clear()
F.start_create(A_EV, A_EV)
check("🔒 گارد فرم‌ساز عمل کرد", "دسترسی کافی ندارید" in last(), last()[:50])
SENT.clear()
EV.start_create(A_FM, A_FM)
check("🔒 گارد رویداد عمل کرد", "دسترسی" in last(), last()[:50])
SENT.clear()
AL.show_logs(A_EV, A_EV, 0)
check("🔒 گارد لاگ عمل کرد", "دسترسی کافی" in last())
SENT.clear()
FIN.show_dashboard(A_EV, A_EV)
check("🔒 گارد مالی عمل کرد", "دسترسی کافی" in last())
SENT.clear()
FIN.show_dashboard(A_FIN, A_FIN)
check("🔒 فاز۱۳: ادمین مالی هم داشبورد درآمد نمی‌بیند",
      "دسترسی" in last(), last()[:60])
SENT.clear()
FIN.show_dashboard(S, S)
check("✅ فقط سوپرادمین داشبورد درآمد می‌بیند", "داشبورد مالی" in last())
check("🔒 fin_view در SUPER_ONLY", "fin_view" in PM.SUPER_ONLY)
check("🔒 fin_export در SUPER_ONLY", "fin_export" in PM.SUPER_ONLY)
check("🔒 هیچ نقشی fin_view ندارد",
      all(not PM.has_perm(a, "fin_view") for a in (A_EV, A_FM, A_RP, A_FIN)))
check("✅ ادمین رسیدها هنوز رسید تأیید می‌کند",
      PM.has_perm(A_FIN, "fin_approve"))

# پنل مدیریت دسترسی
_tid = database.db_execute("SELECT id FROM users WHERE bale_user_id=?",
                           (A_FM,), fetch="one")["id"]
SENT.clear()
M.handle_callback(cb(S, f"pm_view_{_tid}"))
check("صفحه دسترسی", "دسترسی‌های" in last())
SENT.clear()
M.handle_callback(cb(S, f"pm_roles_{_tid}"))
check("انتخاب نقش", "انتخاب نقش" in last())
check("۷ نقش قابل انتخاب",
      sum(1 for d in cbdata() if d.startswith("pm_setr_")) == 7)
M.handle_callback(cb(S, f"pm_setr_broadcast_admin_{_tid}"))
check("نقش عوض شد", PM.get_admin_role(A_FM)[0] == "broadcast_admin")
SENT.clear()
M.handle_callback(cb(S, f"pm_cats_{_tid}"))
check("دسته‌های دسترسی", any(d.startswith("pm_cat_") for d in cbdata()))
SENT.clear()
M.handle_callback(cb(S, "pm_cat_رویدادها_" + str(_tid)))
check("دسترسی‌های دسته", "رویدادها" in last())
M.handle_callback(cb(S, f"pm_tog_ev_create_{_tid}"))
check("تاگل از پنل", PM.has_perm(A_FM, "ev_create"))

SENT.clear()
M.handle_callback(cb(S, f"pm_revoke_{_tid}"))
check("سلب دسترسی", not PM.has_perm(A_FM, "ev_create"))
check("is_admin صفر شد", database.db_execute(
    "SELECT is_admin FROM users WHERE bale_user_id=?", (A_FM,),
    fetch="one")["is_admin"] == 0)

# فقط سوپرادمین
SENT.clear()
AD.show_admin_perms(A_EV, A_EV, _tid)
check("🔒 فقط سوپرادمین دسترسی می‌دهد", "سوپرادمین" in last())
check("⛔ سوپرادمین قابل سلب نیست", not PM.revoke_admin(S))


# ═══════ ۴) امتیاز سخت‌گیرانه ═══════
print("\n[۴] نظام امتیازدهی سخت‌گیرانه")
check("امتیاز تولد ۱۵", database.get_setting("oc_birthday_points", "15") == "15")
_fw = database.db_execute("SELECT points FROM point_rules WHERE key="
                          "'first_workshop'", fetch="one")
check("اولین کارگاه ۲۰", _fw and _fw["points"] == 20, str(_fw))
_l5 = database.db_execute("SELECT points FROM point_rules WHERE key='loyal_5'",
                          fetch="one")
check("۵ حضور ۳۰", _l5 and _l5["points"] == 30)
_nw = database.db_execute("SELECT points FROM occasions WHERE key='nowruz'",
                          fetch="one")
check("نوروز ۱۰", _nw and _nw["points"] == 10)

_high = database.db_scalar("SELECT COUNT(*) FROM point_rules WHERE points > 100")
check("هیچ قانونی بیش از ۱۰۰ نیست", _high == 0, str(_high))
_ohigh = database.db_scalar("SELECT COUNT(*) FROM occasions WHERE points > 20")
check("هیچ مناسبتی بیش از ۲۰ نیست", _ohigh == 0, str(_ohigh))
check("ثبت‌نام صرف امتیاز ندارد", not database.db_execute(
    "SELECT id FROM point_rules WHERE trigger='event_register'", fetch="one"))

check("سقف تکی ۱۰۰", OC.POINTS_MAX_SINGLE == 100)
check("سقف روزانه ۵۰", OC.POINTS_MAX_DAILY == 50)
check("سقف ماهانه ۳۰۰", OC.POINTS_MAX_MONTHLY == 300)

_p, _ = OC.cap_points(U1, 500)
check("سقف تکی اعمال شد", _p == 50, str(_p))
_p2, _ = OC.cap_points(U1, 10)
check("مقدار کم دست‌نخورده", _p2 == 10)
_p3, _ = OC.cap_points(U1, 0)
check("صفر صفر می‌ماند", _p3 == 0)
_p4, _ = OC.cap_points(U1, -5)
check("منفی صفر می‌شود", _p4 == 0)

# سقف روزانه
database.db_execute("""INSERT INTO greeting_log
    (bale_user_id,kind,ref_key,jyear,points_given)
    VALUES (?,'rule','t1',0,45)""", (U2,))
_p5, _why = OC.cap_points(U2, 30)
check("سقف روزانه کاهش داد", _p5 == 5, f"{_p5}")
database.db_execute("""INSERT INTO greeting_log
    (bale_user_id,kind,ref_key,jyear,points_given)
    VALUES (?,'rule','t2',0,10)""", (U2,))
_p6, _why6 = OC.cap_points(U2, 30)
check("سقف روزانه پر شد", _p6 == 0 and "روزانه" in _why6, f"{_p6}/{_why6}")

check("امتیاز کسب‌شده امروز", OC.points_earned(U2, 1) == 55)

# سقف در عمل
database.db_execute("UPDATE point_rules SET points=999 WHERE key='profile_100'")
database.db_execute("UPDATE users SET complete_profile_percent=100 "
                    "WHERE bale_user_id=?", (U1,))
_before = database.db_scalar("SELECT total_points FROM users "
                             "WHERE bale_user_id=?", (U1,)) or 0
SENT.clear()
_n, _pts = OC.check_point_rules(U1, "profile_complete")
check("امتیاز نجومی مهار شد", _pts <= 50, str(_pts))
_after = database.db_scalar("SELECT total_points FROM users "
                            "WHERE bale_user_id=?", (U1,)) or 0
check("افزایش امتیاز محدود بود", _after - _before <= 50,
      f"{_before}->{_after}")

database.set_setting("points_max_daily", "10", "t", "occasions")
_p7, _ = OC.cap_points(U1, 100)
check("سقف از تنظیمات قابل تغییر", _p7 <= 10, str(_p7))
database.set_setting("points_max_daily", "50", "t", "occasions")


# ═══════ ۵) منو و مسیرها ═══════
print("\n[۵] منو و مسیرها")
SENT.clear()
M.handle_callback(cb(S, "admin_panel"))
_p = allcb()
check("فاز۱۳: پنل خلوت شد (حداکثر ۸ دکمه)", len(_p) <= 8, str(len(_p)))
for _d in ("admin_users", "admin_settings", "admin_stats"):
    check(f"دکمه {_d} در پنل", _d in _p, str(_p))
check("🔄 مالی از پنل اصلی حذف شد", "fin_dash" not in _p)
check("🔄 لاگ از پنل اصلی حذف شد", "lg_list" not in _p)

# مالی حالا زیر تنظیمات › پرداخت است
SENT.clear()
M.handle_callback(cb(S, "admin_settings"))
check("💳 پرداخت و مالی در تنظیمات", "set_payment" in allcb(), str(allcb()))
SENT.clear()
M.handle_callback(cb(S, "set_payment"))
_pay = allcb()
check("💰 داشبورد درآمد زیر پرداخت", "fin_dash" in _pay, str(_pay))
check("🧾 تراکنش‌ها زیر پرداخت", "fin_tx_0_all" in _pay)
check("🎛️ روش پیشرفته زیر پرداخت", "pay_adv" in _pay)

# لاگ حالا زیر مدیریت کاربران › ادمین‌ها است
SENT.clear()
M.handle_callback(cb(S, "um_admins"))
check("📋 لاگ زیر مدیریت ادمین‌ها", "lg_list" in allcb(), str(allcb()))

bad = ["pm_view_", "pm_view_abc", "pm_view_999999", "pm_roles_x",
       "pm_setr_bad_1", "pm_cat_ندارد_1", "pm_tog_evil_1",
       "pm_tog_sys_admins_1", "pm_revoke_999999",
       "fin_tx_x_all", "fin_tx_0_bad", "lg_c__0", "lg_p_x_همه",
       "lg_adm_x", "lg_c_ندارد_0"]
crashed = []
for d in bad:
    try:
        M.handle_callback(cb(S, d))
    except Exception as e:
        crashed.append(f"{d}: {type(e).__name__} {e}")
check("هیچ ورودی خرابی کرش نمی‌کند", not crashed, str(crashed[:2]))


# ═══════ ۶) یکپارچگی ═══════
print("\n[۶] یکپارچگی")
check("دسترسی نامعتبر ذخیره نشده", database.db_scalar(
    "SELECT COUNT(*) FROM admin_permissions WHERE role NOT IN "
    f"({','.join('?' * len(PM.ROLES))})", tuple(PM.ROLES)) == 0)
check("هر ادمین یک رکورد", database.db_scalar(
    """SELECT COUNT(*) FROM (SELECT bale_user_id FROM admin_permissions
       GROUP BY bale_user_id HAVING COUNT(*)>1)""") == 0)
check("امتیاز منفی نیست",
      database.db_scalar("SELECT COUNT(*) FROM users WHERE total_points<0") == 0)
check("مبلغ منفی نیست",
      database.db_scalar("SELECT COUNT(*) FROM transactions WHERE amount<0") == 0)
_orphan = database.db_scalar(
    """SELECT COUNT(*) FROM transactions t WHERE t.event_id IS NOT NULL
       AND NOT EXISTS (SELECT 1 FROM events e WHERE e.id=t.event_id)""")
check("تراکنش یتیم نیست", _orphan == 0, str(_orphan))

print("\n" + "=" * 62)
print(f"✅ موفق: {OK}    ❌ ناموفق: {FAIL}")
print("=" * 62)
if FAILED:
    print("\n❌ تست‌های ناموفق:")
    for f_ in FAILED:
        print(f"   • {f_}")
    sys.exit(1)
print("\n🎉 همه تست‌های فاز ۱۲ پاس شدند!")
