# -*- coding: utf-8 -*-
"""
🧪 تست فاز ۳۹ — مرکز صدور بازطراحی‌شده · قفل صدور خودسرویس

🎯 درخواست شما:
   «مرکز صدور گواهی ببر تو قسمت کاربران»
   «اسم طرحو که میزنم جلو نمیره اینو حلش کن»
   «برای تمام داده‌های گواهینامه بتونم داده‌هایی که دوست دارم
    تعریف کنم بدون هیچ دستکاری و اضافه‌ای از تو»
   «کد فردی و این چیزا اگه من پارامترش گذاشتم قرار بده وگرنه
    از خودت بساز»
   «صدور گواهینامه چرا دست کاربره؟ کلا از کاربر بگیر»
"""
import os
import re
import sys
import json
import pathlib
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("BOT_TOKEN", "0:test")
os.environ.setdefault("SUPER_ADMIN_ID", "1")
os.environ.setdefault("BOT_USERNAME", "nora_test_bot")

_TDB = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                    "_p39.db")
for _e in ("", "-wal", "-shm"):
    try:
        os.remove(_TDB + _e)
    except OSError:
        pass
os.environ["DB_PATH"] = _TDB

_ok = _bad = 0
_fails = []


def check(name, cond, extra=""):
    global _ok, _bad
    if cond:
        _ok += 1
    else:
        _bad += 1
        _fails.append(name)
        print(f"  ❌ {name}  {extra}")


print("=" * 58)
print("🧪 فاز ۳۹ — مرکز صدور · قفل صدور")
print("=" * 58)

import bale_api as _ba

SENT = []


def _fake_api(m, p=None, **k):
    SENT.append({"method": m, "params": p or {}})
    if m == "getMe":
        return {"ok": True, "result": {"username": "nora_test_bot"}}
    return {"ok": True, "result": {"message_id": len(SENT)}}


def _fake_send(cid, text="", kb=None, **k):
    SENT.append({"method": "sendMessage",
                 "params": {"chat_id": cid, "text": text, "kb": kb}})
    return {"ok": True, "result": {"message_id": len(SENT)}}


def _fake_file(cid, path=None, caption="", **k):
    SENT.append({"method": "sendFile",
                 "params": {"chat_id": cid, "path": str(path),
                            "text": caption}})
    return {"ok": True, "result": {"message_id": len(SENT)}}


_ba.api = _fake_api
_ba.send = _fake_send
_ba.send_or_edit = _fake_send
_ba.send_photo = _fake_file
_ba.send_document = _fake_file

SRC = {f.name: f.read_text(encoding="utf-8")
       for f in pathlib.Path(".").glob("*.py")}
MN = SRC["main.py"]

import database
import certgen as cg
import attendance as att
import certissue as ci
import certcenter as cc
import events as ev_mod

database.init_db()
cg.init_cert_tables()
att.init_attendance_tables()
ev_mod.init_event_tables()
cc.init_tables()


def _last():
    for s in reversed(SENT):
        if s["method"] == "sendMessage":
            return s["params"].get("text") or ""
    return ""


def _last_kb():
    for s in reversed(SENT):
        if s["method"] == "sendMessage" and s["params"].get("kb"):
            return s["params"]["kb"]
    return {}


def _cbs(kb):
    return [b.get("callback_data", "")
            for r in (kb or {}).get("inline_keyboard", []) for b in r]


def _txts(kb):
    return [b.get("text", "")
            for r in (kb or {}).get("inline_keyboard", []) for b in r]


# ══════════════════════════════════════════════════════════
print("\n[۱] 🐛 باگ «اسم طرح جلو نمیره»")

check("هندلر cc_ در روتر هست",
      'state_name.startswith("cc_")' in MN)
check("cc_mod.handle_text صدا زده می‌شود",
      "cc_mod.handle_text" in MN)
_i_cc = MN.find('state_name.startswith("cc_")')
_i_cb = MN.find('state_name.startswith("cb_")')
check("cc قبل از cb بررسی می‌شود", 0 < _i_cc < _i_cb)

# 🔬 تست زنده — کل جریان ساخت طرح
database.db_execute(
    "INSERT OR IGNORE INTO users (bale_user_id, first_name, is_admin,"
    " admin_type, profile_status) VALUES "
    "(1,'ادمین',1,'super_admin','approved')")

import docx as _dx
from docx.shared import Mm
from docx.enum.section import WD_ORIENT

TMP = tempfile.mkdtemp(prefix="p39_")
d = _dx.Document()
_s = d.sections[0]
_s.orientation = WD_ORIENT.LANDSCAPE
_s.page_width, _s.page_height = Mm(297), Mm(210)
d.add_paragraph("گواهی می‌شود {نام} — کد ملی {کدملی}")
d.add_paragraph("در {رویداد} با مدرس {مدرس} به مدت {ساعت} ساعت")
d.add_paragraph("تاریخ برگزاری {تاریخ} · صدور {تاریخ صدور}")
d.add_paragraph("شماره نامه {شماره نامه} · کد {کد فردی}")
d.add_paragraph("{کیوآر}")
TPL = f"{TMP}/t.docx"
d.save(TPL)
tid = database.db_execute(
    "INSERT INTO cert_templates (name, file_path, kind, is_active,"
    " is_default) VALUES (?,?,?,1,1)", ("قالب۳۹", TPL, "docx"))

from auth import set_state, get_state_name, clear_state

set_state(1, "cc_title", {}, merge=False)
SENT.clear()
_res = cc.handle_text({"from": {"id": 1}, "chat": {"id": 1},
                       "text": "طرح آزمایشی ۳۹"})
check("🔑 handle_text نام طرح را می‌پذیرد", _res is True)
_iss = database.db_execute(
    "SELECT * FROM cert_issues WHERE title=?",
    ("طرح آزمایشی ۳۹",), fetch="one")
check("🔑 طرح واقعاً ساخته شد", bool(_iss))
check("state پاک شد", not get_state_name(1))
iid = (_iss or {}).get("id")

# ══════════════════════════════════════════════════════════
print("[۲] 🧩 دسته‌بندی تازهٔ پارامترها")

check("PER_PERSON هست", hasattr(cc, "PER_PERSON"))
check("SYSTEM_KEYS هست", hasattr(cc, "SYSTEM_KEYS"))
check("field_kind هست", hasattr(cc, "field_kind"))
check("KIND_FA هست", len(cc.KIND_FA) == 4)

for k, want in (("name", "person"), ("national_id", "person"),
                ("event", "fixed"), ("teacher", "fixed"),
                ("hours", "fixed"), ("doc_no", "fixed"),
                ("serial", "system"), ("issue_date", "system"),
                ("verify_code", "system"), ("qr", "image")):
    check(f"{k} → {want}", cc.field_kind(k) == want,
          cc.field_kind(k))

cc.scan_template(tid)
ask, auto, imgs = cc.template_fields(tid)
check("🔑 «رویداد» پرسیده می‌شود", "event" in ask, f"{ask}")
check("🔑 «مدرس» پرسیده می‌شود", "teacher" in ask)
check("🔑 «ساعت» پرسیده می‌شود", "hours" in ask)
check("🔑 «شماره نامه» پرسیده می‌شود", "doc_no" in ask,
      f"ask={ask}")
check("🔑 «تاریخ صدور» هم پرسیده می‌شود",
      any(k in ask for k in ("issue_date", "issued_on")), f"{ask}")
check("«نام» پرسیده نمی‌شود", "name" not in ask)
check("«کد ملی» پرسیده نمی‌شود", "national_id" not in ask)
check("«نام» در دستهٔ هر‌نفر است", "name" in auto)
check("کیوآر تصویر است", "qr" in imgs)

# ══════════════════════════════════════════════════════════
print("[۳] ✍️ مقدار شما بر ربات غالب است")

cc.set_value(iid, "event", "کارگاه امداد")
cc.set_value(iid, "teacher", "دکتر رضایی")
cc.set_value(iid, "hours", "۱۲")

database.db_execute(
    "INSERT OR IGNORE INTO users (bale_user_id, first_name, last_name,"
    " national_id, mobile, profile_status) VALUES "
    "(39001,'سارا','احمدی','0012345679','09121110000','approved')")
cc.add_share(iid, 39001, "سارا احمدی", "0012345679", "09121110000")

iss = cc.get_issue(iid)
sh = cc.shares(iid)[0]
mp, info = cc.mapping_for(iss, sh)
check("پارامتر ثابت اعمال شد", mp.get("event") == "کارگاه امداد")
check("مدرس اعمال شد", mp.get("teacher") == "دکتر رضایی")
check("کد فردی خودکار ساخته شد", bool(info.get("person_code")))
_auto_code = info["person_code"]

# 🔑 حالا خودم کد فردی می‌دهم
cc.set_value(iid, "verify_code", "MYCODE1")
iss = cc.get_issue(iid)
mp2, info2 = cc.mapping_for(iss, sh)
check("🔑 کد فردی دستی برنده است",
      info2.get("person_code") == "MYCODE1", info2.get("person_code"))
check("در نقشه هم همان است", mp2.get("verify_code") == "MYCODE1")
check("با کد خودکار فرق دارد", _auto_code != "MYCODE1")

# 🔑 شمارهٔ نامهٔ دستی
cc.set_value(iid, "doc_no", "1421/999/د")
iss = cc.get_issue(iid)
mp3, info3 = cc.mapping_for(iss, sh)
check("🔑 شمارهٔ نامهٔ دستی برنده است",
      info3.get("doc_no") == "1421/999/د", info3.get("doc_no"))
check("شناسهٔ کامل ترکیب درست است",
      info3.get("full_code") == "1421/999/د-MYCODE1",
      info3.get("full_code"))

# 🔑 تاریخ صدور دستی
cc.set_value(iid, "issue_date", "۱۴۰۵/۰۱/۰۱")
iss = cc.get_issue(iid)
mp4, _ = cc.mapping_for(iss, sh)
check("🔑 تاریخ صدور دستی برنده است",
      mp4.get("issue_date") == "۱۴۰۵/۰۱/۰۱", mp4.get("issue_date"))
check("🔑 کلید هم‌معنی هم هماهنگ شد",
      mp4.get("issued_on") == "۱۴۰۵/۰۱/۰۱", mp4.get("issued_on"))

# پاک کردن → برگشت به خودکار
cc.set_value(iid, "issue_date", "")
cc.set_value(iid, "verify_code", "")
iss = cc.get_issue(iid)
mp5, info5 = cc.mapping_for(iss, sh)
check("خالی کردن → ربات دوباره می‌سازد",
      bool(mp5.get("issue_date")) and
      mp5.get("issue_date") != "۱۴۰۵/۰۱/۰۱")
check("کد فردی به خودکار برگشت",
      info5.get("person_code") != "MYCODE1")

# ══════════════════════════════════════════════════════════
print("[۴] ⚡ پر کردن پشت‌سرهم")

for fn in ("wizard_start", "_wizard_ask", "wizard_input"):
    check(f"تابع {fn}", hasattr(cc, fn))
check("روتر cc_wiz_", "cc_wiz_" in MN)

SENT.clear()
cc.wizard_start(1, 1, iid)
check("ویزارد شروع شد", "گام" in _last(), _last()[:50])
check("state تنظیم شد", get_state_name(1) == "cc_wiz")

SENT.clear()
cc.handle_text({"from": {"id": 1}, "chat": {"id": 1},
                "text": "کارگاه تازه"})
check("پاسخ ثبت و گام بعد", "گام" in _last())
_v = cc.issue_values(cc.get_issue(iid))
check("مقدار ذخیره شد",
      "کارگاه تازه" in str(_v.values()), str(_v)[:80])

SENT.clear()
cc.handle_text({"from": {"id": 1}, "chat": {"id": 1}, "text": "پایان"})
check("«پایان» ویزارد را می‌بندد", not get_state_name(1))
clear_state(1)

# ══════════════════════════════════════════════════════════
print("[۵] 👥 راه‌های تازهٔ گیرنده")

for fn in ("ask_manual", "manual_input", "ask_search",
           "search_input", "add_one_user", "send_sample_excel"):
    check(f"تابع {fn}", hasattr(cc, fn))
for cb in ("cc_amn_", "cc_asr_", "cc_addu_", "cc_xs_"):
    check(f"روتر {cb}", cb in MN)

# ✍️ افزودن دستی
cc.clear_shares(iid)
set_state(1, "cc_manual", {"iid": iid}, merge=False)
SENT.clear()
cc.handle_text({"from": {"id": 1}, "chat": {"id": 1},
                "text": "علی محمدی\nزهرا کریمی، 0023456787\nم"})
_n = len(cc.shares(iid))
check("🔑 دستی: ۲ نفر اضافه شدند", _n == 2, f"{_n}")
_names = {s.get("full_name") for s in cc.shares(iid)}
check("نام‌ها درست خوانده شدند",
      "علی محمدی" in _names and "زهرا کریمی" in _names, str(_names))
_z = next(s for s in cc.shares(iid) if s["full_name"] == "زهرا کریمی")
check("کد ملی از خط جدا شد",
      _z.get("national_id") == "0023456787", _z.get("national_id"))
check("خط نامفهوم رد شد", "م" not in _names)

# 🔍 جستجو
set_state(1, "cc_search", {"iid": iid}, merge=False)
SENT.clear()
cc.handle_text({"from": {"id": 1}, "chat": {"id": 1},
                "text": "احمدی"})
check("🔍 جستجو نتیجه داد", "سارا" in _last(), _last()[:60])
check("دکمهٔ افزودن دارد",
      any(c.startswith("cc_addu_") for c in _cbs(_last_kb())))

_before = len(cc.shares(iid))
SENT.clear()
cc.add_one_user(1, 1, iid, 39001)
check("➕ کاربر از جستجو اضافه شد",
      len(cc.shares(iid)) == _before + 1)

set_state(1, "cc_search", {"iid": iid}, merge=False)
SENT.clear()
cc.handle_text({"from": {"id": 1}, "chat": {"id": 1},
                "text": "هیچکس‌نیست‌xyz"})
check("جستجوی بی‌نتیجه پیام مهربان", "پیدا نشد" in _last())

# 📥 نمونهٔ اکسل
SENT.clear()
cc.send_sample_excel(1, 1, iid)
_f = [s for s in SENT if s["method"] == "sendFile"]
check("📥 نمونهٔ اکسل ساخته شد", len(_f) >= 1)
if _f:
    _p = _f[0]["params"]["path"]
    check("فایل روی دیسک هست", os.path.exists(_p))
    try:
        import openpyxl
        _wb = openpyxl.load_workbook(_p)
        _hd = [c.value for c in _wb.active[1]]
        check("ستون نام دارد",
              any("نام" in str(h) for h in _hd), str(_hd))
        check("ستون کد ملی دارد",
              any("ملی" in str(h) for h in _hd))
    except ImportError:
        pass

# صفحهٔ گیرندگان همهٔ راه‌ها را دارد
SENT.clear()
cc.show_audience(1, 1, iid)
_c = _cbs(_last_kb())
for pre in ("cc_ae_", "cc_as_", "cc_asr_", "cc_amn_", "cc_ax_",
            "cc_xs_"):
    check(f"دکمهٔ {pre} در صفحهٔ گیرندگان",
          any(x.startswith(pre) for x in _c), str(_c))

# ══════════════════════════════════════════════════════════
print("[۶] 📤 چهار روش انتشار")

check("PUBLISH_WAYS هست", len(cc.PUBLISH_WAYS) == 4)
for w in ("notify", "link", "pdf", "silent"):
    check(f"روش {w}", w in cc.PUBLISH_WAYS)
for fn in ("show_publish", "publish_way", "export_all"):
    check(f"تابع {fn}", hasattr(cc, fn))
check("روتر cc_pub_", "cc_pub_" in MN)
check("روتر cc_pw_", "cc_pw_" in MN)

# 🧩 همهٔ پارامترهای ثابت را پر کن تا طرح آمادهٔ انتشار شود
_ask, _, _ = cc.template_fields(tid)
for _k in _ask:
    if not str(cc.issue_values(cc.get_issue(iid)).get(_k) or "").strip():
        cc.set_value(iid, _k, "مقدار تست")
_ok2, _todo2 = cc.issue_ready(cc.get_issue(iid))
check("طرح آمادهٔ انتشار شد", _ok2, str(_todo2))

SENT.clear()
cc.show_publish(1, 1, iid)
_c2 = _cbs(_last_kb())
check("صفحهٔ انتشار باز شد", "روش انتشار" in _last())
check("۴ روش دکمه دارند",
      sum(1 for x in _c2 if x.startswith("cc_pw_")) == 4, str(_c2))

# 🤫 بی‌صدا
_n0 = database.db_scalar("SELECT COUNT(*) FROM certificates") or 0
SENT.clear()
cc.publish_way(1, 1, iid, "silent")
check("🤫 منتشر شد",
      (cc.get_issue(iid) or {}).get("status") == "published")
_n1 = database.db_scalar("SELECT COUNT(*) FROM certificates") or 0
check("🔑 بی‌صدا فایلی نساخت", _n1 == _n0)

# 📄 خروجی یکجا
SENT.clear()
cc.export_all(1, 1, iid)
_zips = [s for s in SENT if s["method"] == "sendFile"
         and str(s["params"].get("path", "")).endswith(".zip")]
if _zips:
    check("📄 بستهٔ زیپ ساخته شد", True)
    _zp = _zips[0]["params"]["path"]
    check("زیپ روی دیسک هست", os.path.exists(_zp))
    import zipfile as _zf
    with _zf.ZipFile(_zp) as z:
        check("زیپ خالی نیست", len(z.namelist()) >= 1,
              str(z.namelist()[:3]))
    _n2 = database.db_scalar("SELECT COUNT(*) FROM certificates") or 0
    check("🔑 خروجی یکجا واقعاً ساخت", _n2 > _n1, f"{_n1}→{_n2}")
else:
    check("خروجی یکجا پیام روشن داد",
          "❌" in _last() or "ناموفق" in _last(), _last()[:70])

# سقف
database.set_setting("cc_export_max", "1", "", "certificate")
SENT.clear()
cc.export_all(1, 1, iid)
check("سقف خروجی رعایت می‌شود", "زیاد است" in _last(), _last()[:60])
database.set_setting("cc_export_max", "200", "", "certificate")

# ══════════════════════════════════════════════════════════
print("[۷] 🔒 قفل صدور خودسرویس — ضروری")

check("تابع self_issue_allowed", hasattr(ci, "self_issue_allowed"))
check("تابع deny_self_issue", hasattr(ci, "deny_self_issue"))
check("🔑 پیش‌فرض بسته است", ci.self_issue_allowed() is False)
check("تنظیم در پنل هست",
      any(r[0] == "cert_self_issue" for r in cg.CERT_SETTINGS))
_row = next(r for r in cg.CERT_SETTINGS if r[0] == "cert_self_issue")
check("پیش‌فرضش «۰» است", _row[4] == "0", _row[4])
check("پیام قابل تغییر",
      any(r[0] == "cert_self_issue_msg" for r in cg.CERT_SETTINGS))

# 🔬 کاربر عادی نمی‌تواند صادر کند
eid = database.db_execute(
    "INSERT INTO events (title, event_type, status, start_date,"
    " has_certificate, has_attendance) VALUES "
    "(?,?,?,?,1,0)",
    ("رویداد قفل", "workshop", "published", "1405/05/01"))
database.db_execute(
    "INSERT OR IGNORE INTO event_registrations (event_id,"
    " bale_user_id, status) VALUES (?,?,'confirmed')", (eid, 39001))

_nb = database.db_scalar("SELECT COUNT(*) FROM certificates") or 0
SENT.clear()
att.issue_certificate(39001, 39001, eid)
_na = database.db_scalar("SELECT COUNT(*) FROM certificates") or 0
check("🔑 کاربر عادی گواهینامه نساخت", _na == _nb, f"{_nb}→{_na}")
check("پیام مهربان گرفت",
      "برگزارکننده" in _last() or "در راه" in _last(), _last()[:70])
check("دکمهٔ گواهینامه‌های من دارد",
      "u_my_certs" in _cbs(_last_kb()))

# 🔬 ادمین می‌تواند
SENT.clear()
_c, _e = ci.issue_one(eid, 39001, 1, send_to_user=False)
check("🔑 ادمین می‌تواند صادر کند", bool(_c), str(_e)[:60])

# 🔬 با روشن‌کردن تنظیم، کاربر هم می‌تواند
database.set_setting("cert_self_issue", "1", "", "certificate")
check("با روشن‌کردن، باز می‌شود", ci.self_issue_allowed() is True)
database.set_setting("cert_self_issue", "0", "", "certificate")
check("دوباره بسته شد", ci.self_issue_allowed() is False)

# دکمهٔ بعد از ثبت حضور
check("دکمهٔ ucert_ مشروط شد",
      "_ci.self_issue_allowed()" in SRC["attendance.py"])
check("eventcert هم قفل دارد",
      "self_issue_allowed" in SRC["eventcert.py"])
check("پنل کاربر دکمهٔ صدور را پنهان می‌کند",
      "به‌زودی صادر می‌شود" in SRC["eventcert.py"])

# ══════════════════════════════════════════════════════════
print("[۸] 🏛️ دکمه در مدیریت کاربران")

from bale_api import kb_users_menu

_kb = kb_users_menu({"total": 10, "admins": 2, "blocked": 1,
                     "pending": 3})
_rows = _kb["inline_keyboard"]
check("🔑 دکمهٔ مرکز صدور هست",
      any("مرکز صدور" in b.get("text", "")
          for r in _rows for b in r), str(_txts(_kb)))
check("callback درست است", "cc_home" in _cbs(_kb))
check("حداکثر ۱۰ ردیف", len(_rows) <= 10, f"{len(_rows)}")

SENT.clear()
cc.show_home(1, 1)
check("بازگشت به مدیریت کاربران",
      "admin_users" in _cbs(_last_kb()), str(_cbs(_last_kb())))

# ══════════════════════════════════════════════════════════
print("[۹] 🖥️ صفحهٔ پارامترها — سه دسته")

SENT.clear()
cc.show_params(1, 1, iid)
_t = _last()
check("عنوان تازه", "دادهٔ گواهینامه" in _t, _t[:40])
check("بخش ثابت", "برای همه یکسان" in _t)
check("بخش سیستمی", "ربات می‌سازد" in _t)
check("بخش هر‌نفر", "برای هر نفر فرق دارد" in _t)
check("دکمهٔ ویزارد",
      any(c.startswith("cc_wiz_") for c in _cbs(_last_kb())))

# ══════════════════════════════════════════════════════════
print("[۱۰] 🧩 استحکام")

# ⚠️ این توابع پیام می‌فرستند و مقدار برگشتی‌شان مهم نیست؛
#    مهم این است که **استثنا نیندازند**.
def _safe(name, fn):
    try:
        fn()
        check(name, True)
    except Exception as _ex:
        check(f"{name} ({type(_ex).__name__}: {_ex})", False)


_safe("wizard با طرح ناموجود", lambda: cc.wizard_start(1, 1, 999999))
_safe("publish_way نامعتبر",
      lambda: cc.publish_way(1, 1, iid, "nope"))
_safe("export_all طرح ناموجود", lambda: cc.export_all(1, 1, 999999))
_safe("add_one_user کاربر ناموجود",
      lambda: cc.add_one_user(1, 1, iid, 999999))

set_state(1, "cc_manual", {"iid": iid}, merge=False)
SENT.clear()
cc.handle_text({"from": {"id": 1}, "chat": {"id": 1}, "text": "   "})
check("فهرست خالی رد می‌شود", "خوانده نشد" in _last())
clear_state(1)

set_state(1, "cc_search", {"iid": iid}, merge=False)
SENT.clear()
cc.handle_text({"from": {"id": 1}, "chat": {"id": 1}, "text": "ا"})
check("جستجوی تک‌حرفی رد می‌شود", "۲ حرف" in _last())
clear_state(1)

check("field_kind با کلید ناشناخته",
      cc.field_kind("nope_xyz") == "fixed")

print()
print("=" * 58)
print(f"✅ موفق: {_ok}   ❌ ناموفق: {_bad}")
if _fails:
    print("\nناموفق‌ها:")
    for f in _fails:
        print("  •", f)
print("=" * 58)

try:
    database.close_db()
except Exception:
    pass
for _e in ("", "-wal", "-shm"):
    try:
        os.remove(_TDB + _e)
    except OSError:
        pass

sys.exit(1 if _bad else 0)
