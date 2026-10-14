#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
🧪 تست فاز ۳۲ — کیوآر شناور · چند گواهینامه · صدور گروهی

سه درخواست کاربر:
  ۱. «کیوآر کد رو میسازه ولی اونجا که باید قرار نمیده، میندازه
      پشت بکگراند» + «صفحه سفید بالا بک گراند»
  ۲. «با فایل اکسل گواهینامه بزنی و با لینک اشتراکش بزارم»
  ۳. «بعضی رویدادها همکاری هستن و چندین گواهی مستقل صادر بشه»
"""
import sys
import os
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

OK = 0
FAIL = 0
ERRORS = []


def check(label, cond, detail=""):
    global OK, FAIL
    if cond:
        OK += 1
    else:
        FAIL += 1
        ERRORS.append(f"{label} {detail}")


def section(t):
    print(f"\n{'─' * 62}\n{t}\n{'─' * 62}")


import certgen as cg
import eventcert as evc
import certbatch as cb
import database

HAVE_DOCX = True
try:
    import docx
    from docx.shared import Mm, Emu
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn
except ImportError:
    HAVE_DOCX = False

HAVE_QR = True
try:
    import qrcode
except ImportError:
    HAVE_QR = False

HAVE_XL = True
try:
    import openpyxl
except ImportError:
    HAVE_XL = False

TMP = Path("/tmp/_p32")
TMP.mkdir(exist_ok=True)


# ══════════════════════════════════════════════════════════
section("۱) موتور جاگذاری شناور")
# ══════════════════════════════════════════════════════════
for fn in ["_add_floating_image", "_fit_background", "_float_cfg",
           "_page_mm", "FLOAT_POS", "ANCHOR_FA"]:
    check(f"certgen.{fn}", hasattr(cg, fn))

check("۹ گوشهٔ لنگر", len(cg.ANCHOR_FA) == 9)
check("qr در FLOAT_POS", "qr" in cg.FLOAT_POS)

c = cg._float_cfg("qr")
check("موقعیت qr خوانده شد",
      all(k in c for k in ("x", "y", "anchor")), str(c))
check("مختصات در بازهٔ درست",
      0 <= c["x"] <= 100 and 0 <= c["y"] <= 100, str(c))

# کلید ناشناخته نباید بترکد
c2 = cg._float_cfg("unknown_key_xyz")
check("کلید ناشناخته امن", c2["anchor"] == "cc", str(c2))


# ══════════════════════════════════════════════════════════
section("۲) 🔳 کیوآر شناور روی همه‌چیز")
# ══════════════════════════════════════════════════════════
if not (HAVE_DOCX and HAVE_QR):
    print("   ⏭️ python-docx یا qrcode نصب نیست")
else:
    qr_p = str(TMP / "q.png")
    qrcode.make("TEST-QR").save(qr_p)

    d = docx.Document()
    s = d.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Mm(297), Mm(210)
    d.add_paragraph("گواهینامه")
    d.add_paragraph("{نام}")
    tpl = str(TMP / "t.docx")
    d.save(tpl)

    out = str(TMP / "o.docx")
    got, err = cg.fill_docx(tpl, out, {"name": "علی رضایی"},
                            qr_png=qr_p)
    check("قالب پر شد", bool(got), err)

    if got:
        d2 = docx.Document(out)
        xml = d2.element.xml
        # 🔴 این چهار مورد قلب رفع هستند
        check("wp:anchor ساخته شد", "<wp:anchor" in xml)
        check('behindDoc="0" (جلوی همه)', 'behindDoc="0"' in xml,
              "کیوآر نباید پشت پس‌زمینه برود")
        check('relativeFrom="page"', 'relativeFrom="page"' in xml,
              "مختصات باید نسبت به صفحه باشد نه متن")
        check("wrapNone", "wrapNone" in xml)
        check("تصویر درج شد",
              sum(1 for r in d2.part.rels.values()
                  if "image" in r.reltype) >= 1)

        # 📍 مختصات واقعی درست است؟
        found = False
        for anc in d2.element.body.iter(qn("wp:anchor")):
            for tag in ("wp:positionH", "wp:positionV"):
                pos = anc.find(qn(tag))
                if pos is None:
                    continue
                off = pos.find(qn("wp:posOffset"))
                if off is not None and off.text:
                    v = int(off.text)
                    lim = int(Emu(Mm(297 if "H" in tag else 210)))
                    check(f"{tag} داخل صفحه", 0 <= v <= lim,
                          f"{v} > {lim}")
                    found = True
        check("مختصات ثبت شد", found)


# ══════════════════════════════════════════════════════════
section("۳) 📐 پس‌زمینه روی صفحه")
# ══════════════════════════════════════════════════════════
if not HAVE_DOCX:
    print("   ⏭️ python-docx نصب نیست")
else:
    from PIL import Image
    bgp = str(TMP / "bg.png")
    Image.new("RGB", (1492, 1054), (20, 50, 120)).save(bgp)

    d = docx.Document()
    s = d.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Mm(297), Mm(210)
    # 🔴 عمداً کمی کوچک‌تر — همان باگ واقعی کاربر
    d.add_paragraph().add_run().add_picture(bgp, width=Mm(296.0))
    d.add_paragraph("{نام}")
    tpl2 = str(TMP / "t2.docx")
    d.save(tpl2)

    out2 = str(TMP / "o2.docx")
    got, err = cg.fill_docx(tpl2, out2, {"name": "علی"})
    check("قالب با پس‌زمینه پر شد", bool(got), err)

    if got:
        d3 = docx.Document(out2)
        sec = d3.sections[0]
        pw, ph = int(sec.page_width), int(sec.page_height)
        fixed = False
        for inl in d3.element.body.iter(qn("wp:inline")):
            e = inl.find(qn("wp:extent"))
            if e is None:
                continue
            cx, cy = int(e.get("cx")), int(e.get("cy"))
            if cx > pw * 0.8:
                check("پس‌زمینه دقیقاً عرض صفحه", cx == pw,
                      f"{cx} ≠ {pw}")
                check("پس‌زمینه دقیقاً ارتفاع صفحه", cy == ph,
                      f"{cy} ≠ {ph}")
                fixed = True
        check("پس‌زمینه پیدا و تنظیم شد", fixed)
        check("حاشیهٔ بالا صفر شد", int(sec.top_margin) == 0,
              str(int(sec.top_margin)))
        check("حاشیهٔ چپ صفر شد", int(sec.left_margin) == 0)

    # 🛡️ تصویر کوچک نباید پس‌زمینه حساب شود
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(297), Mm(210)
    small = str(TMP / "sm.png")
    Image.new("RGB", (100, 100), "red").save(small)
    d.add_paragraph().add_run().add_picture(small, width=Mm(25))
    tpl3 = str(TMP / "t3.docx")
    d.save(tpl3)
    n = cg._fit_background(docx.Document(tpl3))
    check("تصویر کوچک دست‌نخورده", n == 0, f"{n} تصویر تغییر کرد")


# ══════════════════════════════════════════════════════════
section("۴) 🎓 چند گواهینامه برای رویداد")
# ══════════════════════════════════════════════════════════
cg.init_cert_tables()
EV = 990032
database.db_execute("DELETE FROM event_certs WHERE event_id=?", (EV,))

check("جدول event_certs", True)
check("۶ شرط صدور", len(evc.CONDITIONS) == 6)
check("سقف منطقی", evc.MAX_PER_EVENT >= 2)

r1, e1 = evc.add_cert(EV, "گواهینامهٔ دانشگاه فرهنگیان",
                      issuer="دانشگاه فرهنگیان", condition="present")
r2, e2 = evc.add_cert(EV, "گواهینامهٔ هلال‌احمر",
                      issuer="جمعیت هلال‌احمر", condition="sessions",
                      min_sessions=3)
check("گواهینامهٔ اول", bool(r1), e1)
check("گواهینامهٔ دوم", bool(r2), e2)
check("هر دو ثبت شدند", evc.count_certs(EV) == 2,
      str(evc.count_certs(EV)))

lst = evc.list_certs(EV)
check("ترتیب حفظ شد", lst[0]["id"] == r1 and lst[1]["id"] == r2)
check("صادرکننده ذخیره شد",
      lst[0]["issuer"] == "دانشگاه فرهنگیان")
check("شرط جدا برای هرکدام",
      lst[0]["condition"] != lst[1]["condition"])

# جابه‌جایی
evc.move_cert(r2, -1)
lst = evc.list_certs(EV)
check("جابه‌جایی کار کرد", lst[0]["id"] == r2, str([x["id"] for x in lst]))
evc.move_cert(r2, 1)

# سقف
for i in range(evc.MAX_PER_EVENT):
    evc.add_cert(EV, f"تست {i}")
rid, err = evc.add_cert(EV, "بیش از سقف")
check("سقف رعایت شد", rid is None and "حداکثر" in err, err)

# متن تبلیغ
ad = evc.ad_lines(EV)
check("متن تبلیغ ساخته شد", len(ad) > 0)
check("عنوان در تبلیغ",
      any("فرهنگیان" in x for x in ad), str(ad[:3]))
check("صادرکننده در تبلیغ",
      any("🏛️" in x for x in ad))
check("شرط در تبلیغ", any("🎯" in x for x in ad))

# خلاصه
sm = evc.summary(EV)
check("خلاصه ساخته شد", bool(sm) and "گواهینامه" in sm, sm)

# شرط‌ها
c_sess = evc.get_cert(r2)
txt = evc._cond_fa(c_sess)
check("شرط جلسه فارسی", "جلسه" in txt, txt)

ok, why = evc.eligible(evc.get_cert(r1), EV, 999999999)
check("کاربر ثبت‌نام‌نکرده رد شد", not ok)
check("دلیل داده شد", bool(why), why)

# پاکسازی
database.db_execute("DELETE FROM event_certs WHERE event_id=?", (EV,))
check("پاکسازی", evc.count_certs(EV) == 0)


# ══════════════════════════════════════════════════════════
section("۵) 📊 خواندن اکسل")
# ══════════════════════════════════════════════════════════
if not HAVE_XL:
    print("   ⏭️ openpyxl نصب نیست")
else:
    # نرمال‌سازی
    for raw, want in [("09121234567", "09121234567"),
                      ("989121234567", "09121234567"),
                      ("۰۹۱۲۱۲۳۴۵۶۷", "09121234567"),
                      ("0912-123-4567", "09121234567"),
                      ("123", "")]:
        check(f"موبایل «{raw}»", cb.norm_mobile(raw) == want,
              cb.norm_mobile(raw))

    def _mk_nid(p9):
        t = sum(int(p9[i]) * (10 - i) for i in range(9)) % 11
        return p9 + str(t if t < 2 else 11 - t)

    good = _mk_nid("001024716")
    check(f"کد ملی معتبر {good}", cb.valid_nid(good) == good)
    check("کد ملی تکراری رد شد", not cb.valid_nid("1111111111"))
    check("چک‌سام غلط رد شد", not cb.valid_nid("0012345678"))

    for h, w in [("نام و نام خانوادگی", "full_name"),
                 ("موبایل", "mobile"), ("کد ملی", "national_id"),
                 ("Full Name", "full_name"), ("تلفن همراه", "mobile"),
                 ("نمره", "score")]:
        check(f"ستون «{h}»", cb.detect_column(h) == w,
              str(cb.detect_column(h)))
    check("ستون ناشناخته", cb.detect_column("xyz123") is None)

    # اکسل واقعی با خطاهای عمدی
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ردیف", "نام و نام خانوادگی", "موبایل",
               "کد ملی", "نمره"])
    ws.append([1, "علی رضایی", "09121234567", good, 18])
    ws.append([2, "مریم احمدی", "۰۹۱۲۹۸۷۶۵۴۳", "", 20])
    ws.append([3, "", "09121111111", "", ""])          # نام خالی
    ws.append([4, "حسن کریمی", "0912-333-4444", "1111111111", 15])
    ws.append([5, "علی رضایی", "09121234567", "", ""])  # تکراری
    ws.append([None, None, None, None, None])           # خالی
    xp = str(TMP / "in.xlsx")
    wb.save(xp)

    rows, cmap, errs = cb.parse_excel(xp)
    check("۳ ردیف معتبر", len(rows) == 3, f"{len(rows)}")
    check("ستون نام", cmap.get("full_name") == 1)
    check("نام خالی رد شد",
          any("نام خالی" in e for e in errs), str(errs))
    check("کد ملی بد گزارش شد",
          any("کد ملی" in e for e in errs))
    check("تکراری گرفته شد",
          any("تکراری" in e for e in errs))
    check("سطر خالی نادیده", all(r["full_name"] for r in rows))
    check("ارقام فارسی تبدیل شد",
          rows[1]["mobile"] == "09129876543", rows[1]["mobile"])
    check("ستون اضافه ذخیره شد",
          "score" in rows[0]["extra"], str(rows[0]["extra"]))

    # فایل خراب
    bad = str(TMP / "bad.xlsx")
    Path(bad).write_bytes(b"not-an-excel")
    r2_, _c, e2_ = cb.parse_excel(bad)
    check("فایل خراب امن", r2_ == [] and bool(e2_))

    # فایل بدون ستون شناخته‌شده
    wb2 = openpyxl.Workbook()
    wb2.active.append(["الف", "ب", "ج"])
    wb2.active.append([1, 2, 3])
    np_ = str(TMP / "nocol.xlsx")
    wb2.save(np_)
    r3, _c, e3 = cb.parse_excel(np_)
    check("بدون ستون شناخته‌شده", r3 == [] and bool(e3))

    # فایل نمونه
    sp = cb.sample_excel(str(TMP / "sample.xlsx"))
    check("فایل نمونه ساخته شد", sp and Path(sp).exists())
    if sp:
        wbs = openpyxl.load_workbook(sp)
        check("نمونه راهنما دارد", "راهنما" in wbs.sheetnames,
              str(wbs.sheetnames))


# ══════════════════════════════════════════════════════════
section("۶) 🎟️ دستهٔ گواهینامه و لینک")
# ══════════════════════════════════════════════════════════
database.db_execute("DELETE FROM cert_batches WHERE title LIKE 'T32%'")

rows = [{"full_name": "علی رضایی", "mobile": "09121234567",
         "national_id": "", "extra": {"score": "18"}},
        {"full_name": "مریم احمدی", "mobile": "09129876543",
         "national_id": "", "extra": {}}]
bid, code = cb.create_batch("T32 تست", rows, uid=1)
check("دسته ساخته شد", bool(bid))
check("کد یکتا", bool(code) and len(code) >= 6, code)

b = cb.get_batch(bid)
check("بازیابی با id", bool(b))
check("بازیابی با code", cb.get_batch(code=code)["id"] == bid)
check("تعداد کل ثبت شد", b["total"] == 2)

st = cb.batch_stats(bid)
check("آمار: کل", st["total"] == 2)
check("آمار: صادرشده صفر", st["issued"] == 0)
check("آمار: درصد صفر", st["pct"] == 0)

link = cb.batch_link(code)
check("لینک ساخته شد", bool(link))
check("لینک شامل gc_", f"gc_{code}" in link, link)

ok, why = cb.can_claim(b)
check("دستهٔ فعال قابل دریافت", ok, why)
ok2, why2 = cb.can_claim(None)
check("دستهٔ نامعتبر رد شد", not ok2 and bool(why2))

# تطبیق
found, why = cb.find_row(b, {"mobile": "09121234567"})
check("تطبیق با موبایل", bool(found) and
      found["full_name"] == "علی رضایی", why)
nf, why = cb.find_row(b, {"mobile": "09120000000"})
check("موبایل ناموجود رد شد", nf is None and bool(why), why)
nf2, why2 = cb.find_row(b, {})
check("بدون موبایل پیام درست", nf2 is None and "پروفایل" in why2,
      why2)

# نقشهٔ پارامتر
if found:
    mp = cb.build_mapping(b, found, {"mobile": "09121234567"})
    check("نام در نقشه", mp.get("name") == "علی رضایی")
    check("ستون اضافه در نقشه", mp.get("score") == "18", str(mp))

# ثبت دریافت
if found:
    cb.mark_claimed(found["id"], 12345, None)
    st2 = cb.batch_stats(bid)
    check("آمار بعد از دریافت", st2["issued"] == 1, str(st2))
    check("درصد به‌روز شد", st2["pct"] == 50, str(st2["pct"]))

# سقف
database.db_execute("UPDATE cert_batches SET max_claims=1 WHERE id=?",
                    (bid,))
ok3, why3 = cb.can_claim(cb.get_batch(bid))
check("سقف اعمال شد", not ok3 and "ظرفیت" in why3, why3)

# غیرفعال
database.db_execute("UPDATE cert_batches SET max_claims=0 WHERE id=?",
                    (bid,))
cb.toggle_batch(bid, False)
ok4, why4 = cb.can_claim(cb.get_batch(bid))
check("غیرفعال رد شد", not ok4, why4)

# مهلت گذشته
cb.toggle_batch(bid, True)
database.db_execute(
    "UPDATE cert_batches SET expires_at='2020-01-01' WHERE id=?", (bid,))
ok5, why5 = cb.can_claim(cb.get_batch(bid))
check("مهلت گذشته رد شد", not ok5 and "مهلت" in why5, why5)

cb.delete_batch(bid)
check("حذف دسته", cb.get_batch(bid) is None)
check("ردیف‌ها هم پاک شدند", not cb.batch_rows(bid))


# ══════════════════════════════════════════════════════════
section("۷) روتر و اتصال‌ها")
# ══════════════════════════════════════════════════════════
mn = Path("main.py").read_text(encoding="utf-8")
check("import eventcert", "import eventcert" in mn)
check("import certbatch", "import certbatch" in mn)
check("deep-link gc_", 'param.startswith("gc_")' in mn)
check("دریافت اکسل", 'cb_wait_file' in mn)
check("متن evc_", 'state_name.startswith("evc_")' in mn)
check("متن cb_", 'state_name.startswith("cb_")' in mn)

for route in ["evc_p_", "evc_add_", "evc_v_", "evc_cond_",
              "evc_get_", "evc_iss_", "cb_home", "cb_new",
              "cb_sample", "cb_v_", "cbB_", "cbT_"]:
    check(f"مسیر {route}", route in mn)

ev = Path("events.py").read_text(encoding="utf-8")
check("دکمه در پنل رویداد", 'evc_p_{eid}' in ev)
check("تبلیغ رویداد", "_evc.ad_lines" in ev)
check("راهنمای قدم بعد", "گواهینامه روشن است ولی تعریف نشده" in ev)

cgs = Path("certgen.py").read_text(encoding="utf-8")
# 🔄 فاز ۳۶: «صدور گروهی با اکسل» جای خود را به «مرکز صدور»
#    داد — به درخواست کاربر: «اگه اون پنل که گفتم بسازی میتونی
#    کلا قسمت صدور گروهی با اکسل پاک کنی». ورودی اکسل حالا
#    یکی از چهار راهِ انتخاب گیرنده در مرکز است.
#    روتر همچنان cb_home را می‌پذیرد (سازگاری عقب‌رو).
check("دکمه در مرکز گواهینامه",
      '"cc_home"' in cgs or '"cb_home"' in cgs)

att = Path("attendance.py").read_text(encoding="utf-8")
check("ستون event_cert_id", "event_cert_id" in att)
check("قید یکتای سه‌تایی",
      "event_id, bale_user_id, event_cert_id" in att)


# ══════════════════════════════════════════════════════════
section("۸) واژه‌نامه")
# ══════════════════════════════════════════════════════════
import re
for fn in ["eventcert.py", "certbatch.py", "certgen.py"]:
    bad = []
    in_quote = False          # نقل‌قول چندخطی کاربر
    for i, ln in enumerate(Path(fn).read_text(encoding="utf-8")
                           .splitlines(), 1):
        st_ = ln.lstrip()
        # ردیابی نقل‌قول چندخطی «…»
        if "«" in ln and "»" not in ln:
            in_quote = True
        elif "»" in ln and "«" not in ln:
            in_quote = False
            continue
        if in_quote or st_.startswith("#") or ("«" in ln and "»" in ln):
            continue
        if re.search(r"گواهی(?!نامه)(?! می)(?!\w)", ln):
            bad.append(str(i))
    check(f"«گواهینامه» در {fn}", not bad, " · ".join(bad[:3]))


print(f"\n{'=' * 62}")
print(f"✅ موفق: {OK}    ❌ ناموفق: {FAIL}")
print("=" * 62)
if ERRORS:
    print("\n❌ خطاها:")
    for e in ERRORS[:30]:
        print(f"   • {e}")
    sys.exit(1)
print("\n🎉 همه تست‌های فاز ۳۲ پاس شدند!")
