#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
🧪 تست فاز ۳۰ — تفکیک کامل درآمد

🎯 «توی داشبورد درآمد ریز همه چی مشخص کن، مثلا درآمد کارمزد
    جدا مشخص نشده، درآمدها را همه تفکیک و کامل و با جزئیات بیار»

مشکل قبل: فقط SUM(amount) نشان داده می‌شد. معلوم نبود چقدرش
کارمزد است، چقدرش سهم برگزارکننده، و بازپرداخت هم کم نمی‌شد.
"""
import sys
import os
import random
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

OK = 0
FAIL = 0
ERRORS = []


def check(label, cond, detail=""):
    global OK, FAIL
    if cond:
        OK += 1
    else:
        FAIL += 1
        ERRORS.append(f"{label} {detail}")


def section(t):
    print(f"\n{'─' * 62}\n{t}\n{'─' * 62}")


import finance as F


# ══════════════════════════════════════════════════════════
section("۱) توابع تفکیک موجودند")
# ══════════════════════════════════════════════════════════
for fn in ["_agg", "_refunded", "_by_source", "_by_method",
           "_fee_detail", "_infra_split"]:
    check(f"تابع {fn}", hasattr(F, fn))

for fn in ["show_fees", "show_breakdown", "show_period"]:
    check(f"صفحهٔ {fn}", hasattr(F, fn))

check("SOURCES تعریف شده", hasattr(F, "SOURCES"))
check("PERIODS تعریف شده", hasattr(F, "PERIODS"))
if hasattr(F, "SOURCES"):
    check("۴ منبع درآمد", len(F.SOURCES) == 4,
          str(list(F.SOURCES)))
    for k in ["event", "form", "plan", "other"]:
        check(f"منبع {k}", k in F.SOURCES)
if hasattr(F, "PERIODS"):
    check("۴ بازهٔ زمانی", len(F.PERIODS) == 4)


# ══════════════════════════════════════════════════════════
section("۲) ساخت دادهٔ آزمایشی")
# ══════════════════════════════════════════════════════════
import database
import events

try:
    events.init_event_tables()
    check("جدول تراکنش آماده", True)
except Exception as e:
    check("جدول تراکنش آماده", False, str(e))

TAG = f"t30_{int(time.time())}"
database.db_execute(
    "DELETE FROM transactions WHERE payload LIKE 't30_%'")

# دادهٔ کنترل‌شده — اعداد دقیق می‌دانیم
PLAN = [
    # (منبع, روش, پایه, کارمزد, وضعیت, امتیاز)
    ("event", "custom_gateway", 100000, 2500, "paid", 0),
    ("event", "custom_gateway", 200000, 5000, "paid", 0),
    ("event", "card", 300000, 0, "paid", 0),
    ("form", "bale_wallet", 150000, 3000, "paid", 0),
    ("form", "bale_wallet", 250000, 5000, "paid", 0),
    ("form", "points", 50000, 0, "paid", 500),
    ("plan", "custom_gateway", 400000, 10000, "paid", 0),
    ("other", "onsite", 80000, 0, "paid", 0),
    # غیرموفق‌ها نباید در درآمد بیایند
    ("event", "custom_gateway", 999000, 9999, "pending", 0),
    ("form", "card", 888000, 0, "failed", 0),
    # بازپرداخت باید کم شود
    ("event", "custom_gateway", 120000, 3000, "refunded", 0),
]

for i, (src, m, base, fee, st, pts) in enumerate(PLAN):
    ev = 1 if src == "event" else None
    fm = 1 if src == "form" else None
    pl = 1 if src == "plan" else None
    database.db_execute(
        "INSERT INTO transactions (bale_user_id,event_id,form_id,plan_id,"
        "payload,amount,status,method,fee_amount,base_amount,points_used,"
        "created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,datetime('now'))",
        (9000 + i, ev, fm, pl, f"{TAG}_{i}", base + fee, st, m,
         fee, base, pts))

check("دادهٔ آزمایشی ثبت شد", True)

# اعداد مورد انتظار (فقط paid)
E_GROSS = sum(b + f for s, m, b, f, st, p in PLAN if st == "paid")
E_FEE = sum(f for s, m, b, f, st, p in PLAN if st == "paid")
E_BASE = sum(b for s, m, b, f, st, p in PLAN if st == "paid")
E_PTS = sum(p for s, m, b, f, st, p in PLAN if st == "paid")
E_N = sum(1 for s, m, b, f, st, p in PLAN if st == "paid")
E_REF = sum(b + f for s, m, b, f, st, p in PLAN if st == "refunded")


# ══════════════════════════════════════════════════════════
section("۳) _agg — تجمیع دقیق")
# ══════════════════════════════════════════════════════════
w = f"payload LIKE '{TAG}_%'"
a = F._agg(w)

check("ناخالص درست", a["gross"] == E_GROSS,
      f"{a['gross']:,} ≠ {E_GROSS:,}")
check("کارمزد درست", a["fee"] == E_FEE,
      f"{a['fee']:,} ≠ {E_FEE:,}")
check("پایه درست", a["base"] == E_BASE,
      f"{a['base']:,} ≠ {E_BASE:,}")
check("امتیاز درست", a["points"] == E_PTS,
      f"{a['points']} ≠ {E_PTS}")
check("تعداد درست", a["n"] == E_N, f"{a['n']} ≠ {E_N}")

# 🔴 حیاتی: معادلهٔ حسابداری
check("پایه + کارمزد = ناخالص", a["base"] + a["fee"] == a["gross"],
      f"{a['base']:,} + {a['fee']:,} ≠ {a['gross']:,}")

# 🔴 غیرموفق‌ها نباید در درآمد بیایند
#    مقایسه با «اگر همه شمرده می‌شدند» — نه با تک عدد
_all_sum = sum(b + f for s, m, b, f, st, p in PLAN)
check("جمع فقط موفق‌هاست", a["gross"] == E_GROSS,
      f"{a['gross']:,} ≠ {E_GROSS:,}")
check("pending/failed/refunded کنار گذاشته شدند",
      a["gross"] < _all_sum,
      f"{a['gross']:,} باید < {_all_sum:,} باشد")
check("تعداد فقط موفق‌ها", a["n"] == E_N, f"{a['n']} ≠ {E_N}")


# ══════════════════════════════════════════════════════════
section("۴) _refunded — بازپرداخت")
# ══════════════════════════════════════════════════════════
r = F._refunded(w)
check("بازپرداخت درست", r == E_REF, f"{r:,} ≠ {E_REF:,}")
check("خالص = ناخالص − بازپرداخت",
      a["gross"] - r == E_GROSS - E_REF)


# ══════════════════════════════════════════════════════════
section("۵) _by_source — تفکیک منبع")
# ══════════════════════════════════════════════════════════
srcs = {s["key"]: s for s in F._by_source()}
check("منابع پیدا شدند", len(srcs) >= 3, str(list(srcs)))

for key in ["event", "form", "plan"]:
    exp_g = sum(b + f for s, m, b, f, st, p in PLAN
                if s == key and st == "paid")
    if key in srcs:
        # چون دادهٔ قبلی هم ممکن است باشد، فقط ≥ چک می‌کنیم
        check(f"منبع {key} ثبت شد", srcs[key]["gross"] >= exp_g,
              f"{srcs[key]['gross']:,} < {exp_g:,}")

# مرتب‌سازی نزولی
vals = [s["gross"] for s in F._by_source()]
check("مرتب‌سازی نزولی", vals == sorted(vals, reverse=True))

# هر منبع باید معادلهٔ حسابداری را رعایت کند
for s in F._by_source():
    check(f"معادله در {s['key']}",
          s["base"] + s["fee"] == s["gross"],
          f"{s['base']:,}+{s['fee']:,}≠{s['gross']:,}")


# ══════════════════════════════════════════════════════════
section("۶) _fee_detail — ریز کارمزد")
# ══════════════════════════════════════════════════════════
det = F._fee_detail()
check("کارمزد تفکیک شد", len(det) >= 1)

# فقط روش‌هایی که کارمزد دارند
for d in det:
    check(f"کارمزد {d['key']} مثبت", d["fee"] > 0)
    check(f"نام {d['key']} فارسی", bool(d["name"]))
    check(f"آیکن {d['key']}", bool(d["icon"]))

# مجموع ریز باید با کل بخواند
tot_det = sum(d["fee"] for d in det)
tot_all = F._agg()["fee"]
check("جمع ریز = جمع کل کارمزد", tot_det == tot_all,
      f"{tot_det:,} ≠ {tot_all:,}")

# مرتب‌سازی
fvals = [d["fee"] for d in det]
check("کارمزد مرتب نزولی", fvals == sorted(fvals, reverse=True))


# ══════════════════════════════════════════════════════════
section("۷) _by_method — تفکیک روش")
# ══════════════════════════════════════════════════════════
meths = F._by_method()
check("روش‌ها تفکیک شدند", len(meths) >= 3)
for m in meths:
    check(f"روش {m['key']} نام دارد", bool(m["name"]))
    check(f"روش {m['key']} تعداد دارد", m["n"] > 0)

# جمع روش‌ها = کل
tot_m = sum(m["gross"] for m in meths)
check("جمع روش‌ها = کل ناخالص", tot_m == F._agg()["gross"],
      f"{tot_m:,} ≠ {F._agg()['gross']:,}")


# ══════════════════════════════════════════════════════════
section("۸) رندر صفحات")
# ══════════════════════════════════════════════════════════
import permissions as perm
import bale_api

CAP = []
_o_soe, _o_s = bale_api.send_or_edit, bale_api.send
_o_req = perm.require

bale_api.send_or_edit = lambda cid, t, kb=None, **k: CAP.append(t)
bale_api.send = lambda cid, t, kb=None, **k: CAP.append(t)
F.send_or_edit = bale_api.send_or_edit
F.send = bale_api.send
perm.require = lambda *a, **k: True
F.perm = perm

try:
    # داشبورد
    CAP.clear()
    F.show_dashboard(1, 1)
    txt = CAP[-1] if CAP else ""
    check("داشبورد رندر شد", bool(txt))
    check("داشبورد: ناخالص", "ناخالص" in txt)
    check("داشبورد: کارمزد", "کارمزد" in txt)
    check("داشبورد: بهای خدمات", "بهای خدمات" in txt)
    check("داشبورد: منابع درآمد", "منابع درآمد" in txt)
    check("داشبورد: نوار درصد", "▓" in txt or "░" in txt)

    # ریز کارمزد
    CAP.clear()
    F.show_fees(1, 1)
    txt = CAP[-1] if CAP else ""
    check("صفحهٔ کارمزد رندر شد", bool(txt))
    check("کارمزد: مجموع", "مجموع کارمزد" in txt)
    check("کارمزد: سهم از گردش", "سهم از" in txt)
    check("کارمزد: تفکیک روش", "تفکیک روش" in txt)
    check("کارمزد: روند", "روند کارمزد" in txt)

    # تفکیک کامل
    CAP.clear()
    F.show_breakdown(1, 1)
    txt = CAP[-1] if CAP else ""
    check("تفکیک رندر شد", bool(txt))
    check("تفکیک: صورت درآمد", "صورت درآمد" in txt)
    check("تفکیک: تقسیم", "تقسیم درآمد" in txt)
    check("تفکیک: سهم برگزارکننده", "برگزارکننده" in txt)
    check("تفکیک: درآمد ما", "درآمد ما" in txt)
    check("تفکیک: به تفکیک منبع", "تفکیک منبع" in txt)
    check("تفکیک: به تفکیک روش", "تفکیک روش" in txt)

    # گزارش دوره‌ای — همهٔ بازه‌ها
    for k in F.PERIODS:
        CAP.clear()
        F.show_period(1, 1, k)
        txt = CAP[-1] if CAP else ""
        check(f"دوره {k} رندر شد", bool(txt))
        check(f"دوره {k} عنوان دارد", "گزارش" in txt)

    # بازهٔ نامعتبر نباید بترکد
    CAP.clear()
    F.show_period(1, 1, "bogus")
    check("بازهٔ نامعتبر خطا نمی‌دهد", bool(CAP))
except Exception as e:
    check("رندر صفحات", False, str(e))
finally:
    bale_api.send_or_edit, bale_api.send = _o_soe, _o_s
    perm.require = _o_req


# ══════════════════════════════════════════════════════════
section("۹) روتر")
# ══════════════════════════════════════════════════════════
mn = Path("main.py").read_text(encoding="utf-8")
check("مسیر fin_fees", '"fin_fees"' in mn)
check("مسیر fin_break", '"fin_break"' in mn)
check("مسیر fin_period", '"fin_period"' in mn)
check("مسیر fin_per_", 'fin_per_' in mn)

# 🔴 ترتیب حیاتی: fin_per_ باید قبل از fin_pending باشد
i_per = mn.find('startswith("fin_per_")')
i_pend = mn.find('data == "fin_pending"')
check("fin_per_ قبل از fin_pending", 0 < i_per < i_pend,
      f"per={i_per} pending={i_pend}")

fi = Path("finance.py").read_text(encoding="utf-8")
check("دکمهٔ ریز کارمزد", '"fin_fees"' in fi)
check("دکمهٔ تفکیک کامل", '"fin_break"' in fi)
check("دکمهٔ گزارش دوره‌ای", '"fin_period"' in fi)


# ══════════════════════════════════════════════════════════
section("۱۰) خروجی اکسل")
# ══════════════════════════════════════════════════════════
try:
    import openpyxl                       # noqa: F401
    DOC = []
    _o_doc = bale_api.send_document
    bale_api.send_document = lambda cid, p, cap=None, **k: (
        DOC.append(p), {"ok": True})[1]
    bale_api.send = lambda cid, t, kb=None, **k: None
    F.send_document = bale_api.send_document
    F.send = bale_api.send
    perm.require = lambda *a, **k: True
    F.perm = perm
    try:
        F.export_finance(1, 1)
        check("فایل اکسل ساخته شد", bool(DOC))
        if DOC:
            wb = openpyxl.load_workbook(DOC[0])
            names = wb.sheetnames
            for want in ["ریز کارمزد", "منابع درآمد"]:
                check(f"شیت «{want}»", want in names, str(names))
            # ستون‌های تفکیکی در شیت تراکنش‌ها
            ws = wb[names[0]]
            hdr = [c.value for c in ws[1]]
            for col in ["ناخالص (ریال)", "کارمزد (ریال)",
                        "خالص (ریال)", "منبع", "روش پرداخت"]:
                check(f"ستون «{col}»", col in hdr, str(hdr))
            # شیت خلاصه باید کارمزد داشته باشد
            if "خلاصه مالی" in names:
                vals = [r[0] for r in wb["خلاصه مالی"]
                        .iter_rows(values_only=True)]
                check("خلاصه: کارمزد",
                      any("کارمزد" in str(v) for v in vals))
                check("خلاصه: بهای خدمات",
                      any("بهای خدمات" in str(v) for v in vals))
            os.remove(DOC[0])
    finally:
        bale_api.send_document = _o_doc
        bale_api.send_or_edit, bale_api.send = _o_soe, _o_s
        perm.require = _o_req
except ImportError:
    print("   ⏭️ openpyxl نصب نیست")


# ══════════════════════════════════════════════════════════
section("۱۱) حالت‌های مرزی")
# ══════════════════════════════════════════════════════════
# پایگاه خالی
e = F._agg("1=0")
check("خالی: ناخالص صفر", e["gross"] == 0)
check("خالی: تقسیم بر صفر نمی‌شود", e["n"] == 0)

# بازهٔ بی‌نتیجه
e2 = F._agg("created_at > datetime('now','+10 years')")
check("آینده: صفر", e2["gross"] == 0)

# _infra_split نباید بترکد
try:
    F._infra_split()
    check("_infra_split امن", True)
except Exception as ex:
    check("_infra_split امن", False, str(ex))


# ══════════════════════════════════════════════════════════
section("۱۲) پاکسازی")
# ══════════════════════════════════════════════════════════
database.db_execute(f"DELETE FROM transactions WHERE payload LIKE '{TAG}_%'")
left = database.db_scalar(
    f"SELECT COUNT(*) FROM transactions WHERE payload LIKE '{TAG}_%'")
check("دادهٔ آزمایشی پاک شد", not left)


print(f"\n{'=' * 62}")
print(f"✅ موفق: {OK}    ❌ ناموفق: {FAIL}")
print("=" * 62)
if ERRORS:
    print("\n❌ خطاها:")
    for e in ERRORS[:30]:
        print(f"   • {e}")
    sys.exit(1)
print("\n🎉 همه تست‌های فاز ۳۰ پاس شدند!")
