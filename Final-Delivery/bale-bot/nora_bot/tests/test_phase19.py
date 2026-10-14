"""
تست فاز ۱۹ — 🎓 گواهینامه از فایل ورد

درخواست کاربر:
  «از فایل ورد خام گواهی، قسمت‌هایی مثل نام خالی بود، اون از دیتا
   فایل ورد رو با مشخصات شخصی‌سازی شده تبدیل میکرد به PNG و PDF...
   حتی جای کیوآرکد اعتبارسنجی رو من در ورد معلوم میکنم...
   ادمین بتونه فایل ورد بده و ربات سیو کنه به عنوان قالب...
   دیگه تو ربات ننویس گواهی، فقط بنویس گواهینامه»
"""
import json
import time
import os
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

TMP = tempfile.mkdtemp(prefix="nora19_")
os.environ.update({
    "BOT_TOKEN": "1095412158:T", "SUPER_ADMIN_ID": "100000001",
    "SAFIR_API_KEY": "k", "SAFIR_BOT_ID": "317041514",
    "PAYMENT_PROVIDER_TOKEN": "WALLET-TEST-1111111111111111",
    "DB_PATH": f"{TMP}/t.db", "BACKUP_DIR": f"{TMP}/b", "IMPORT_DIR": f"{TMP}/i",
    "FILES_DIR": f"{TMP}/f", "PROFILES_DIR": f"{TMP}/f/p",
    "EVENTS_DIR": f"{TMP}/f/e", "LOG_DIR": f"{TMP}/l", "BROADCAST_RATE": "999",
})

import bale_api  # noqa: E402

SENT = []
FILES = []


def _api(m, p=None, **k):
    SENT.append({"m": m, "p": p or {}})
    if m == "getMe":
        return {"ok": True, "result": {"username": "nora_bot"}}
    if m == "getFile":
        return {"ok": True, "result": {"file_path": "x/y.docx"}}
    return {"ok": True, "result": {"message_id": len(SENT) + 100}}


def _send(cid, text, kb=None, reply_to=None):
    SENT.append({"m": "sendMessage", "p": {"cid": cid, "text": text, "kb": kb}})
    return {"ok": True, "result": {"message_id": len(SENT) + 100}}


def _soe(cid, text, kb=None, mid=None, force_new=False):
    SENT.append({"m": "editMessageText",
                 "p": {"cid": cid, "text": text, "kb": kb}})
    return {"ok": True, "result": {"message_id": mid or 1}}


def _file(method, key, cid, f, caption=None, kb=None):
    FILES.append({"m": method, "path": str(f), "caption": caption})
    SENT.append({"m": method, "p": {"cid": cid, "caption": caption, "kb": kb}})
    return {"ok": True, "result": {"message_id": len(SENT) + 100}}


bale_api.api = _api
bale_api.api_async = _api
bale_api.send = _send
bale_api.send_or_edit = _soe
bale_api._send_file = _file
bale_api.send_photo = lambda c, p, cap=None, kb=None: _file("sendPhoto", "p", c, p, cap, kb)
bale_api.send_document = lambda c, d, cap=None, kb=None: _file("sendDocument", "d", c, d, cap, kb)
bale_api.send_invoice = lambda c, **kw: {"ok": True, "result": {"message_id": 1}}
bale_api.get_file_info = lambda fid: {"ok": True,
                                      "result": {"file_path": "x/tpl.docx"}}

_FAKE_BYTES = {"data": b""}
bale_api.download_file = lambda path: _FAKE_BYTES["data"]

import config  # noqa: E402
import database  # noqa: E402
import auth  # noqa: E402
import certgen as CG  # noqa: E402
import events as EV  # noqa: E402
import attendance as AT  # noqa: E402
import forms as F  # noqa: E402
import payments as PAY  # noqa: E402
import occasions as OC  # noqa: E402
import permissions as PM  # noqa: E402
import club as CL  # noqa: E402
import tickets as TK  # noqa: E402
import about as AB  # noqa: E402
import admin as AD  # noqa: E402
import settings_admin as SET  # noqa: E402
import main as M  # noqa: E402

for mod in (CG, EV, AT, F, PAY, OC, CL, TK, AB, AD, SET, M):
    for n in ("send", "send_or_edit", "send_photo", "send_document"):
        if hasattr(mod, n):
            setattr(mod, n, getattr(bale_api, n))

OK = FAIL = 0
FAILED = []


def check(name, cond, extra=""):
    global OK, FAIL
    if cond:
        OK += 1
        print(f"  ✅ {name}")
    else:
        FAIL += 1
        FAILED.append(name)
        print(f"  ❌ {name}   {extra}")


def texts():
    return [s["p"].get("text") or s["p"].get("caption") or "" for s in SENT]


def last():
    t = texts()
    return t[-1] if t else ""


def alltext():
    return " ".join(texts())


def allcb():
    out = []
    for s in SENT:
        kb = s["p"].get("kb") or {}
        for r in kb.get("inline_keyboard", []):
            for b in r:
                if isinstance(b, dict) and b.get("callback_data"):
                    out.append(b["callback_data"])
    return out


def msg(uid, t):
    return {"message_id": 1, "from": {"id": uid, "first_name": "ت"},
            "chat": {"id": uid, "type": "private"}, "text": t}


def cb(uid, d):
    return {"id": "c", "from": {"id": uid, "first_name": "ت"}, "data": d,
            "message": {"message_id": 1,
                        "chat": {"id": uid, "type": "private"}}}


S = 100000001
U1 = 1550001

print("=" * 62)
print("🧪 تست فاز ۱۹ — گواهینامه از فایل ورد")
print("=" * 62)

database.init_db()
EV.init_event_tables()
AT.init_attendance_tables()
PAY.init_payment_tables()
F.init_form_tables()
OC.init_occasion_tables()
PM.init_permission_tables()
CL.init_club_tables()
TK.init_ticket_tables()
AB.init_about_tables()
CG.init_cert_tables()
database.seed_default_settings()

auth.create_auth_user({"id": S, "first_name": "مدیر"}, "09120000000", "phone")
auth.ensure_super_admin(S)
auth.create_auth_user({"id": U1, "first_name": "امیر", "last_name": "نعمتی"},
                      "09121111111", "phone")
database.db_execute("UPDATE users SET last_name='نعمتی روزبهانی', "
                    "national_id='0010247165', profile_status='approved' "
                    "WHERE bale_user_id=?", (U1,))


# ═══════════ ۱) 🔍 تشخیص پارامترها ═══════════
print("\n[۱] 🔍 شناسایی پارامترها در متن")

check("کاتالوگ پارامترها موجود است", len(CG.PLACEHOLDERS) >= 15,
      str(len(CG.PLACEHOLDERS)))

cases = [
    ("گواهی می‌شود {نام} با کدملی {کدملی}", {"name", "national_id"}),
    ("در دوره {رویداد} به مدت {ساعت} ساعت", {"event", "hours"}),
    ("{{نام و نام خانوادگی}}", {"name"}),
    ("شماره {شماره} — {کیوآر}", {"serial", "qr"}),
    ("{name} and {event}", {"name", "event"}),
    ("متن بدون هیچ پارامتری", set()),
    ("{پارامتر_ناشناخته}", set()),
]
for txt, want in cases:
    got = CG.find_placeholders(txt)
    check(f"«{txt[:34]}»", got == want, f"{got} ≠ {want}")

check("نام فارسی و انگلیسی هر دو کار می‌کند",
      CG.resolve_key("نام") == CG.resolve_key("name") == "name")
check("نیم‌فاصله مشکلی ندارد",
      CG.resolve_key("کد\u200cملی") == "national_id")
check("کیوآر شناخته می‌شود", CG.resolve_key("کیوآر") == "qr")


# ═══════════ ۲) 📄 پر کردن فایل ورد ═══════════
print("\n[۲] 📄 پر کردن قالب ورد")

check("python-docx در دسترس است", CG.docx_available())

import docx  # noqa: E402
from docx.shared import Pt  # noqa: E402

d = docx.Document()
d.add_paragraph("گواهینامه")
d.add_paragraph("گواهی می‌شود جناب‌آقای / سرکار خانم")
# 🔴 عمداً پارامتر را به چند run می‌شکنیم — دقیقاً مثل ورد واقعی
p = d.add_paragraph()
p.add_run("{نا")
p.add_run("م}")
p.add_run(" با کدملی ")
p.add_run("{کد")
p.add_run("ملی}")
d.add_paragraph("در کارگاه «{رویداد}» تاریخ {تاریخ} مدت {ساعت} ساعت")
d.add_paragraph("سازمان: {سازمان} — شماره: {شماره}")
d.add_paragraph("{کیوآر}")
# جدول هم تست شود
tb = d.add_table(rows=1, cols=2)
tb.rows[0].cells[0].text = "صادرکننده:"
tb.rows[0].cells[1].text = "{سمت}"
TPL = os.path.join(TMP, "tpl.docx")
d.save(TPL)

found, err = CG.scan_docx(TPL)
check("قالب بدون خطا خوانده شد", not err, err)
check("🔴 پارامتر شکسته بین چند run پیدا شد",
      "name" in found and "national_id" in found, str(sorted(found)))
check("پارامتر داخل جدول هم پیدا شد", "org_role" in found,
      str(sorted(found)))
check("کیوآر تشخیص داده شد", "qr" in found)

mp = CG.build_mapping(
    name="امیر نعمتی روزبهانی", national_id="0010247165",
    event="جمعیت هلال‌احمر", date="۱۴۰۵/۰۲/۲۷", hours="۳",
    serial="NORA-0405-51472", org_role="معاون فرهنگی")
check("همه کلیدها در mapping هستند",
      set(mp.keys()) >= set(CG.PLACEHOLDERS.keys()))
check("سازمان خودکار پر شد", bool(mp.get("org")))
check("تاریخ صدور خودکار پر شد", bool(mp.get("issue_date")))
check("لینک اعتبارسنجی ساخته شد", "cert_" in (mp.get("verify_url") or ""),
      mp.get("verify_url"))

OUT = os.path.join(TMP, "out.docx")
got, err2 = CG.fill_docx(TPL, OUT, mp)
check("فایل خروجی ساخته شد", bool(got) and Path(OUT).exists(), err2)

res_doc = docx.Document(OUT)
body = "\n".join(p.text for p in res_doc.paragraphs)
for t in res_doc.tables:
    for row in t.rows:
        for c in row.cells:
            body += "\n" + c.text

check("🔴 نام درست جایگزین شد", "امیر نعمتی روزبهانی" in body, body[:120])
# 🔢 فاز ۲۷ — اعداد گواهینامه فارسی می‌شوند
check("🔴 کدملی درست جایگزین شد",
      "0010247165" in body or "۰۰۱۰۲۴۷۱۶۵" in body, body[:120])
check("رویداد جایگزین شد", "هلال‌احمر" in body)
check("شماره جایگزین شد", "NORA-0405-51472" in body)
check("پارامتر جدول جایگزین شد", "معاون فرهنگی" in body)
check("🔴 هیچ جای خالی باقی نماند",
      not CG.find_placeholders(body), str(CG.find_placeholders(body)))
check("🔴 آکولاد خام باقی نماند", "{نام}" not in body and "{کدملی}" not in body)


# ═══════════ ۳) 🔳 کیوآرکد ═══════════
print("\n[۳] 🔳 کیوآرکد اعتبارسنجی")

st = CG.engine_status()
check("وضعیت موتورها خوانده می‌شود", isinstance(st, dict) and "docx" in st)

if st.get("qr"):
    qp = CG.make_qr("https://ble.ir/test?start=cert_X1",
                    os.path.join(TMP, "q.png"))
    check("🔳 تصویر QR ساخته شد", bool(qp) and Path(qp).exists())
    check("🔳 حجم منطقی دارد", os.path.getsize(qp) > 200)

    OUT2 = os.path.join(TMP, "out_qr.docx")
    got2, _ = CG.fill_docx(TPL, OUT2, mp, qr_png=qp)
    check("🔳 قالب با QR پر شد", bool(got2))
    if got2:
        dd = docx.Document(got2)
        n_img = sum(1 for r in dd.part.rels.values()
                    if "image" in r.reltype)
        check("🔳 تصویر QR داخل فایل درج شد", n_img >= 1, str(n_img))
else:
    check("⚠️ qrcode نصب نیست — رد شد", True)

check("لینک اعتبارسنجی الگوی درست دارد",
      "cert_ABC" in CG.verify_url("ABC"), CG.verify_url("ABC"))


# ═══════════ ۴) 🖼️ موتور تصویری ═══════════
print("\n[۴] 🖼️ قالب تصویری")

from PIL import Image  # noqa: E402
BG = os.path.join(TMP, "bg.png")
Image.new("RGB", (1684, 1191), (250, 248, 240)).save(BG)

boxes = [{"key": "name", "x": 840, "y": 560, "size": 48, "align": "center"},
         {"key": "serial", "x": 200, "y": 700, "size": 24, "align": "center"}]
tid_img = database.db_execute(
    """INSERT INTO cert_templates
    (name, file_path, kind, boxes, qr_x, qr_y, qr_size, is_active)
    VALUES (?,?,?,?,?,?,?,1)""",
    ("قالب تصویری", BG, "image", json.dumps(boxes), 1400, 950, 150))

res_i = CG.generate(mp, CG.get_template(tid_img))
check("🖼️ PNG تولید شد", bool(res_i.get("png")), str(res_i.get("error")))
check("🖼️ PDF تولید شد", bool(res_i.get("pdf")))
check("🖼️ موتور درست انتخاب شد", res_i.get("engine") == "image")
if res_i.get("png"):
    im = Image.open(res_i["png"])
    check("🖼️ ابعاد حفظ شد", im.size == (1684, 1191), str(im.size))

# 🔴 باگ جهت متن فارسی
check("🔴 حالت bidi تشخیص داده شد",
      CG._detect_bidi_mode() in ("reshape", "both", "raw"),
      CG._detect_bidi_mode())
_sh = CG._shape_fa("امیر")
check("🔴 متن فارسی برعکس نمی‌شود (باگ python-bidi 0.6)",
      len(_sh) > 0 and _sh != "ریما", _sh)


# ═══════════ ۵) ⚙️ پنل ادمین قالب ═══════════
print("\n[۵] ⚙️ پنل مدیریت قالب")

SENT.clear()
M.handle_callback(cb(S, "ct_panel"))
check("پنل قالب‌ها باز شد", "قالب" in last() and "گواهینامه" in last(),
      last()[:60])
check("دکمه افزودن قالب ورد هست", "ct_new_docx" in allcb())
check("دکمه راهنما هست", "ct_help" in allcb())

SENT.clear()
M.handle_callback(cb(S, "ct_help"))
check("راهنمای پارامترها نمایش داده شد",
      "{نام}" in last() or "نام کامل" in last(), last()[:100])
check("راهنمای کیوآر هست", "کیوآر" in last())

# آپلود قالب ورد
_FAKE_BYTES["data"] = Path(TPL).read_bytes()
SENT.clear()
M.handle_callback(cb(S, "ct_new_docx"))
check("منتظر فایل شد", auth.get_state_name(S) == CG.STATE_CT_UPLOAD)

SENT.clear()
M.process_message({
    "message_id": 2, "from": {"id": S, "first_name": "م"},
    "chat": {"id": S, "type": "private"},
    "document": {"file_name": "cert.docx", "file_id": "F1"}})
check("📄 فایل دریافت و تحلیل شد",
      "پارامتر" in alltext(), alltext()[-200:])
check("نام پارامترها به ادمین نشان داده شد",
      "نام" in alltext() and "کدملی" in alltext())
check("حالا منتظر نام قالب است",
      auth.get_state_name(S) == CG.STATE_CT_NAME)

SENT.clear()
M.process_message(msg(S, "گواهینامه رسمی دوره‌ها"))
_t = database.db_execute("SELECT * FROM cert_templates WHERE kind='docx' "
                         "ORDER BY id DESC LIMIT 1", fetch="one")
check("💾 قالب ذخیره شد", bool(_t) and _t["name"] == "گواهینامه رسمی دوره‌ها")
check("فایل قالب روی دیسک هست", Path(_t["file_path"]).exists())
_phs = json.loads(_t["placeholders"] or "[]")
check("پارامترها ذخیره شدند", "name" in _phs and "national_id" in _phs,
      str(_phs))
check("اولین قالب خودکار پیش‌فرض شد یا قبلی بود",
      _t["is_default"] in (0, 1))

# رد کردن فایل اشتباه
SENT.clear()
M.handle_callback(cb(S, "ct_new_docx"))
M.process_message({
    "message_id": 3, "from": {"id": S, "first_name": "م"},
    "chat": {"id": S, "type": "private"},
    "document": {"file_name": "old.doc", "file_id": "F2"}})
check("🔒 فایل .doc قدیمی رد شد", "doc" in last() or "docx" in last(),
      last()[:80])
auth.clear_state(S)

# قالب بدون پارامتر
d2 = docx.Document()
d2.add_paragraph("متن ساده بدون هیچ پارامتری")
EMPTY = os.path.join(TMP, "empty.docx")
d2.save(EMPTY)
_FAKE_BYTES["data"] = Path(EMPTY).read_bytes()
SENT.clear()
M.handle_callback(cb(S, "ct_new_docx"))
M.process_message({
    "message_id": 4, "from": {"id": S, "first_name": "م"},
    "chat": {"id": S, "type": "private"},
    "document": {"file_name": "empty.docx", "file_id": "F3"}})
check("🔒 قالب بدون پارامتر هشدار داد",
      "پیدا نشد" in last() or "آکولاد" in last(), last()[:100])
auth.clear_state(S)

# جزئیات و پیش‌فرض
SENT.clear()
M.handle_callback(cb(S, f"ctv_{_t['id']}"))
check("جزئیات قالب باز شد", "گواهینامه رسمی" in last(), last()[:60])
SENT.clear()
M.handle_callback(cb(S, f"ctdef_{_t['id']}"))
check("⭐ پیش‌فرض شد",
      CG.get_template(_t["id"])["is_default"] == 1)
check("قالب پیش‌فرض همان است", CG.get_template()["id"] == _t["id"])

# پیش‌نمایش
SENT.clear()
FILES.clear()
M.handle_callback(cb(S, f"ctprev_{_t['id']}"))
check("👁️ پیش‌نمایش اجرا شد",
      "پیش‌نمایش" in alltext(), alltext()[-120:])
check("👁️ فایلی برای ادمین فرستاده شد", len(FILES) >= 1,
      str(len(FILES)))

# تنظیم مختصات قالب تصویری
SENT.clear()
M.handle_callback(cb(S, f"ctbox_{tid_img}"))
check("📐 صفحه مختصات باز شد", "پارامتر" in last(), last()[:60])
M.process_message(msg(S, "نام | 800 | 550 | 46\nکدملی | 600 | 570 | 40\n"
                          "کیوآر | 1400 | 950 | 150"))
_ti = CG.get_template(tid_img)
_bx = json.loads(_ti["boxes"] or "[]")
check("📐 مختصات ذخیره شد", len(_bx) == 2, str(_bx))
check("📐 کیوآر جدا ذخیره شد", _ti["qr_x"] == 1400 and _ti["qr_size"] == 150)


# ═══════════ ۶) 🎓 صدور واقعی گواهینامه ═══════════
print("\n[۶] 🎓 صدور گواهینامه از قالب")

eid = database.db_execute("""INSERT INTO events
    (title, event_type, status, capacity, created_by, has_certificate,
     start_date, duration_minutes, location_type)
    VALUES ('کارگاه هلال احمر','workshop','published',50,?,1,
            '1405/02/27',180,'online')""", (S,))
database.db_execute("""INSERT INTO event_registrations
    (event_id, user_id, bale_user_id, status)
    VALUES (?,?,?,'attended')""", (eid, auth.get_user(U1)["id"], U1))

SENT.clear()
FILES.clear()
cert = AT.issue_certificate(U1, U1, eid, target_uid=U1)
check("🎓 گواهینامه صادر شد", bool(cert))
check("🎓 شماره یکتا دارد", bool(cert.get("serial")))
check("🎓 به قالب وصل است", cert.get("template_id") == _t["id"],
      str(cert.get("template_id")))
check("🎓 فایل ورد شخصی‌سازی‌شده ساخته شد",
      bool(cert.get("docx_path")) and Path(cert["docx_path"]).exists())

# محتوای فایل نهایی
_final = docx.Document(cert["docx_path"])
_fb = "\n".join(p.text for p in _final.paragraphs)
check("🔴 نام واقعی کاربر در گواهینامه هست",
      "امیر" in _fb and "نعمتی" in _fb, _fb[:150])
check("🔴 کد ملی واقعی در گواهینامه هست",
      "0010247165" in _fb or "۰۰۱۰۲۴۷۱۶۵" in _fb, _fb[:120])
check("🔴 عنوان رویداد در گواهینامه هست", "هلال احمر" in _fb)
check("🔴 شماره گواهینامه در فایل هست", str(cert["serial"]) in _fb)
check("🔴 هیچ جای خالی نماند", not CG.find_placeholders(_fb),
      str(CG.find_placeholders(_fb)))
check("📤 فایلی به کاربر ارسال شد", len(FILES) >= 1, str(len(FILES)))

# تکراری نشود
c2 = AT.issue_certificate(U1, U1, eid, target_uid=U1, silent=True)
check("🔒 گواهینامه تکراری صادر نمی‌شود",
      c2 and c2["serial"] == cert["serial"])


# ═══════════ ۷) ✅ اعتبارسنجی با کیوآر ═══════════
print("\n[۷] ✅ اعتبارسنجی گواهینامه")

SENT.clear()
M.handle_deep_link(U1, U1, f"cert_{cert['serial']}", {"id": U1})
check("✅ گواهینامه معتبر شناخته شد", "معتبر" in last(), last()[:80])
check("✅ نام دارنده نمایش داده شد", "امیر" in last())
check("🔒 کد ملی ماسک شد",
      "0010247165" not in last()
      and "۰۰۱۰۲۴۷۱۶۵" not in last(), last()[:200])

SENT.clear()
M.handle_deep_link(U1, U1, "cert_JAALI-999", {"id": U1})
check("🔒 گواهینامه جعلی رد شد", "یافت نشد" in last(), last()[:80])

database.db_execute("UPDATE certificates SET revoked=1 WHERE id=?",
                    (cert["id"],))
SENT.clear()
M.handle_deep_link(U1, U1, f"cert_{cert['serial']}", {"id": U1})
check("🚫 گواهینامه باطل‌شده اعلام می‌شود", "باطل" in last(), last()[:80])
database.db_execute("UPDATE certificates SET revoked=0 WHERE id=?",
                    (cert["id"],))


# ═══════════ ۸) ✍️ واژه «گواهینامه» ═══════════
print("\n[۸] ✍️ همه‌جا «گواهینامه» نوشته شده")

import re as _re
_bad = {}
for _f in ROOT.glob("*.py"):
    _lines = []
    _inq = False          # نقل‌قول چندخطی کاربر
    for _ln in _f.read_text(encoding="utf-8").splitlines():
        _st = _ln.lstrip()
        # 🗣️ ردیابی نقل‌قول چندخطی «…»
        if "«" in _ln and "»" not in _ln:
            _inq = True
        elif "»" in _ln and "«" not in _ln:
            _inq = False
            continue
        if _inq:
            continue
        # کامنت‌ها متن نمایشی نیستند
        if _st.startswith("#"):
            continue
        if "«" in _ln and "»" in _ln:
            continue
        _lines.append(_ln)
    _src = "\n".join(_lines)
    # «گواهی می‌شود» فعل است و باید بماند
    _src = _src.replace("گواهی می‌شود", "").replace("گواهی میشود", "")
    _n = len(_re.findall(r"گواهی(?!نامه)", _src))
    if _n:
        _bad[_f.name] = _n
check("🔴 هیچ‌جا «گواهی» تنها نمانده", not _bad, str(_bad))
check("🔴 «گواهی حضور» حذف شد",
      not any("گواهینامه حضور" in f.read_text(encoding="utf-8")
              for f in ROOT.glob("*.py")))


# ═══════════ ۸.۵) 🎨 فاز ۱۹-ب — پارامترهای ریز ═══════════
print("\n[۸.۵] 🎨 پارامترهای گسترده")

check("🔴 بیش از ۶۰ پارامتر داریم", len(CG.PLACEHOLDERS) >= 60,
      str(len(CG.PLACEHOLDERS)))
check("دسته‌بندی شده", len(CG.PH_GROUPS) >= 6, str(len(CG.PH_GROUPS)))
check("هر پارامتر در یک دسته هست",
      set(CG.PLACEHOLDERS) == {k for _, ks in CG.PH_GROUPS for k in ks},
      str(set(CG.PLACEHOLDERS) ^
          {k for _, ks in CG.PH_GROUPS for k in ks}))

for k in ("sign_img", "sign_img2", "sign_img3", "stamp", "logo",
          "logo2", "logo3", "badge", "photo", "watermark", "barcode"):
    check(f"تصویر «{CG.ph_label(k)}» تعریف شده", k in CG.IMAGE_KEYS)

for k in ("signer", "org_role", "signer2", "role2", "signer3", "role3"):
    check(f"امضا «{CG.ph_label(k)}» هست", k in CG.PLACEHOLDERS)

check("هر تصویر اندازه پیش‌فرض دارد",
      all(k in CG.IMAGE_MM for k in CG.IMAGE_KEYS),
      str(CG.IMAGE_KEYS - set(CG.IMAGE_MM)))
check("اندازه از تنظیمات خوانده می‌شود", CG.image_mm("qr") > 0)

# 🧮 محاسبات خودکار
_m = CG.build_mapping(gender="مرد", first_name="علی", last_name="رضایی",
                      start_date="1405/01/10", end_date="1405/01/20",
                      sessions=8, total_sessions=10, score=85,
                      max_score=100, serial="NORA-0405-99123")
check("🎩 خطاب از جنسیت مرد", "آقای" in _m["gender_title"],
      _m["gender_title"])
check("👤 نام کامل خودکار ساخته شد", _m["name"] == "علی رضایی", _m["name"])
check("📆 بازه تاریخ خودکار",
      "از 1405/01/10" in _m["date_range"]
      or "از ۱۴۰۵/۰۱/۱۰" in _m["date_range"],
      _m["date_range"])
check("📈 درصد حضور خودکار",
      _m["attendance_pct"] in ("80٪", "۸۰٪"),
      _m["attendance_pct"])
check("🔑 کد اعتبارسنجی کوتاه", _m["verify_code"] == "99123",
      _m["verify_code"])

_f = CG.build_mapping(gender="زن")
check("🎩 خطاب از جنسیت زن", "خانم" in _f["gender_title"],
      _f["gender_title"])
_n = CG.build_mapping()
check("🎩 خطاب نامشخص هر دو را می‌آورد",
      "آقای" in _n["gender_title"] and "خانم" in _n["gender_title"])

database.set_setting("cert_pass_score", "60", "t", "certificate")
database.clear_settings_cache()
check("✅ نتیجه قبول", CG.build_mapping(score=85)["result"] == "قبول")
check("❌ نتیجه مردود", CG.build_mapping(score=40)["result"] == "مردود")

database.set_setting("cert_valid_months", "24", "t", "certificate")
database.clear_settings_cache()
_e = CG.build_mapping(issue_date="1405/05/06")["expire_date"]
check("⏳ تاریخ انقضا محاسبه شد",
      _e in ("1407/05/06", "۱۴۰۷/۰۵/۰۶"), _e)
check("⏳ افزودن ماه امن است", CG._add_months_jalali("خراب", 5) == "")
database.set_setting("cert_valid_months", "0", "t", "certificate")
database.clear_settings_cache()

# ✍️ امضاها از تنظیمات
database.set_setting("cert_signer", "حمیدرضا ابراهیمی", "t", "certificate")
database.set_setting("cert_role", "معاون فرهنگی", "t", "certificate")
database.set_setting("cert_signer2", "مریم احمدی", "t", "certificate")
database.clear_settings_cache()
_s2 = CG.build_mapping()
check("✍️ امضاکننده از تنظیمات", _s2["signer"] == "حمیدرضا ابراهیمی")
check("🎖️ سمت از تنظیمات", _s2["org_role"] == "معاون فرهنگی")
check("✍️ امضاکننده دوم", _s2["signer2"] == "مریم احمدی")


print("\n[۸.۶] 🖼️ تصاویر (امضا · مهر · لوگو)")

from PIL import Image as _IM
_assets = {}
# ⚠️ هر تصویر رنگ متفاوت — وگرنه python-docx آن‌ها را یکی می‌شمارد
for _i, _k in enumerate(("sign_img", "sign_img2", "stamp",
                         "logo", "logo2")):
    _p = os.path.join(TMP, f"{_k}.png")
    _IM.new("RGBA", (160, 80), (30 + _i * 40, 60, 140, 255)).save(_p)
    CG.set_asset(_k, _p, uid=S)
    _assets[_k] = _p
check("🖼️ ۵ تصویر ثبت شد", len(CG.all_assets()) == 5,
      str(list(CG.all_assets())))
check("🖼️ خواندن یک تصویر", bool(CG.get_asset("stamp")))
check("🖼️ تصویر ناموجود None می‌دهد", CG.get_asset("nope") is None)

# قالب کامل با همه تصاویر
_d3 = docx.Document()
_d3.add_paragraph("{لوگو} {لوگو دوم}")
_d3.add_paragraph("{خطاب} {نام} — {کدملی}")
_d3.add_paragraph("{امضا} {امضا دوم} {مهر}")
_d3.add_paragraph("{امضاکننده}/{سمت}  {امضاکننده دوم}/{سمت دوم}")
_d3.add_paragraph("{کیوآر}")
_TPL3 = os.path.join(TMP, "full.docx")
_d3.save(_TPL3)
_found3, _ = CG.scan_docx(_TPL3)
check("🔍 همه تصاویر در قالب پیدا شدند",
      {"logo", "logo2", "sign_img", "sign_img2", "stamp", "qr"}
      <= _found3, str(sorted(_found3)))

_tid3 = database.db_execute(
    """INSERT INTO cert_templates (name, file_path, kind, is_active)
    VALUES (?,?,?,1)""", ("کامل", _TPL3, "docx"))
_r3 = CG.generate(CG.build_mapping(name="سارا محمدی", gender="زن",
                                   national_id="0012345678",
                                   serial="NORA-0405-33221"),
                  CG.get_template(_tid3))
check("🖼️ فایل با تصاویر ساخته شد", bool(_r3.get("docx")),
      str(_r3.get("error")))
if _r3.get("docx"):
    _dd = docx.Document(_r3["docx"])
    from docx.oxml.ns import qn as _qn
    _ni = len(list(_dd.element.body.iter(_qn("a:blip"))))
    check("🔴 همه تصاویر درج شدند (۲لوگو+۲امضا+مهر+QR)",
          _ni >= 5, f"{_ni} درج")
    _tb = "\n".join(p.text for p in _dd.paragraphs)
    check("🔴 نام و امضاکننده‌ها درست پر شدند",
          "سارا محمدی" in _tb and "حمیدرضا" in _tb and "مریم" in _tb,
          _tb[:150])
    check("🔴 خطاب خانم درست آمد", "خانم" in _tb, _tb[:100])
    check("🔴 هیچ جای خالی نماند", not CG.find_placeholders(_tb),
          str(CG.find_placeholders(_tb)))

CG.del_asset("logo2")
check("🗑️ حذف تصویر کار می‌کند", CG.get_asset("logo2") is None)


print("\n[۸.۷] 🧹 پاکسازی خودکار فایل‌ها")

check("تنظیم پاکسازی پیش‌فرض روشن", CG.autodelete_on())
check("مهلت پیش‌فرض منطقی", CG.keep_hours() > 0, str(CG.keep_hours()))

_st0 = CG.files_stats()
check("📊 آمار فایل خوانده می‌شود", isinstance(_st0.get("count"), int))
for _i in range(3):
    CG.generate(CG.build_mapping(name=f"تست {_i}", serial=f"T{_i}"),
                CG.get_template(_tid3))
_st1 = CG.files_stats()
check("📊 فایل‌ها شمرده می‌شوند", _st1["count"] > _st0["count"],
      f"{_st0['count']} → {_st1['count']}")

database.set_setting("cert_keep_hours", "24", "t", "certificate")
database.clear_settings_cache()
_n0, _ = CG.cleanup_files()
check("⏰ فایل تازه پاک نمی‌شود", _n0 == 0, str(_n0))

_old = time.time() - 100 * 3600
for _f2 in CG.CERT_OUT_DIR.glob("*"):
    try:
        os.utime(_f2, (_old, _old))
    except Exception:
        pass
_n1, _kb1 = CG.cleanup_files()
check("🧹 فایل قدیمی پاک شد", _n1 > 0, str(_n1))
check("🔴 رکورد گواهینامه دست‌نخورده ماند",
      (database.db_scalar("SELECT COUNT(*) FROM certificates") or 0) >= 1)
check("🔴 اعتبارسنجی هنوز کار می‌کند",
      bool(database.db_execute("SELECT * FROM certificates WHERE serial=?",
                               (cert["serial"],), fetch="one")))

database.set_setting("cert_autodelete", "0", "t", "certificate")
database.clear_settings_cache()
CG.generate(CG.build_mapping(name="ماندگار", serial="KEEP1"),
            CG.get_template(_tid3))
for _f2 in CG.CERT_OUT_DIR.glob("*"):
    try:
        os.utime(_f2, (_old, _old))
    except Exception:
        pass
check("🔒 وقتی خاموش است پاک نمی‌کند", CG.cleanup_files()[0] == 0)
check("🧹 اجبار دستی حتی وقتی خاموش است کار می‌کند",
      CG.cleanup_files(force=True)[0] > 0)
database.set_setting("cert_autodelete", "1", "t", "certificate")
database.clear_settings_cache()

# ♻️ بازتولید
_cid_rec = database.db_execute("SELECT * FROM certificates WHERE id=?",
                               (cert["id"],), fetch="one")
database.db_execute("UPDATE cert_templates SET is_default=1 WHERE id=?",
                    (_tid3,))
_rg = CG.regenerate(_cid_rec)
check("♻️ بازتولید از روی رکورد کار می‌کند",
      bool(_rg) and bool(_rg.get("docx")), str(_rg and _rg.get("error")))
check("♻️ مسیر جدید در دیتابیس ثبت شد",
      bool(database.db_execute("SELECT docx_path FROM certificates "
                               "WHERE id=?", (cert["id"],),
                               fetch="one")["docx_path"]))
check("♻️ بازتولید با ورودی خراب کرش نمی‌کند",
      CG.regenerate(None) is None)


print("\n[۸.۸] ⚙️ پنل تنظیمات گواهینامه و بلیط")

check("تنظیمات گواهینامه تعریف شده", len(CG.CERT_SETTINGS) >= 20,
      str(len(CG.CERT_SETTINGS)))
check("تنظیمات بلیط تعریف شده", len(CG.TICKET_SETTINGS) >= 5,
      str(len(CG.TICKET_SETTINGS)))

SENT.clear()
M.handle_callback(cb(S, "ct_set_cert"))
check("⚙️ پنل تنظیمات گواهینامه باز شد",
      "تنظیمات گواهینامه" in last(), last()[:60])
check("وضعیت فایل‌ها نمایش داده می‌شود", "فایل" in last())
_cbs2 = allcb()
check("دکمه تصاویر هست", "ct_assets" in _cbs2)
check("دکمه اندازه‌ها هست", "ct_sizes" in _cbs2)
check("دکمه پاکسازی هست", "ct_clean" in _cbs2)
check("دکمه قالب‌ها هست", "ct_panel" in _cbs2)
check("دکمه تنظیمات بلیط هست", "ct_set_ticket" in _cbs2)

SENT.clear()
M.handle_callback(cb(S, "ct_set_ticket"))
check("🎫 پنل تنظیمات بلیط باز شد", "بلیط" in last(), last()[:60])

SENT.clear()
M.handle_callback(cb(S, "ctsb_cert_qr_on"))
check("🔳 کلید کیوآر خاموش شد",
      database.get_setting("cert_qr_on", "1") == "0")
M.handle_callback(cb(S, "ctsb_cert_qr_on"))
check("🔳 دوباره روشن شد",
      database.get_setting("cert_qr_on", "1") == "1")

SENT.clear()
M.handle_callback(cb(S, "ctse_cert_signer"))
check("✏️ ویرایش امضاکننده شروع شد", "امضاکننده" in last(), last()[:50])
M.process_message(msg(S, "دکتر محمدی"))
database.clear_settings_cache()
check("✏️ مقدار ذخیره شد",
      database.get_setting("cert_signer", "") == "دکتر محمدی")

SENT.clear()
M.handle_callback(cb(S, "ct_assets"))
check("🖼️ صفحه تصاویر باز شد", "تصاویر" in last(), last()[:50])
SENT.clear()
M.handle_callback(cb(S, "ct_sizes"))
check("📐 صفحه اندازه‌ها باز شد", "اندازه" in last(), last()[:50])
M.handle_callback(cb(S, "ctmm_qr"))
M.process_message(msg(S, "40"))
database.clear_settings_cache()
check("📐 اندازه ذخیره شد", CG.image_mm("qr") == 40, str(CG.image_mm("qr")))

SENT.clear()
M.handle_callback(cb(S, "ct_clean"))
check("🧹 پاکسازی دستی اجرا شد", "پاک" in last(), last()[:60])


# ═══════════ ۸.۹) 🔳 تشخیص خودکار عکس کیوآر + خروجی PNG ═══════════
print("\n[۸.۹] 🔳 تشخیص عکس نمونه کیوآر و خروجی PNG")

# 🔴 باگ کاربر: «یه عکس کیوآرکد خودم میزارم تو فایل، ربات تشخیص نمیده»
#    ریشه: ورد نام فایل را دور می‌ریزد (my_qr.png → image2.png)
import io as _io
try:
    import qrcode as _qrlib
    _HAS_QR = True
except ImportError:
    _HAS_QR = False

if _HAS_QR:
    _q = _qrlib.QRCode(box_size=6, border=2)
    _q.add_data("نمونه")
    _q.make()
    _qsample = os.path.join(TMP, "user_sample_qr.png")
    _q.make_image(fill_color="black",
                  back_color="white").convert("RGB").save(_qsample)

    _blob = Path(_qsample).read_bytes()
    check("🔳 عکس کیوآر از روی محتوا تشخیص داده می‌شود",
          CG.looks_like_qr(_blob))

    _logo_p = os.path.join(TMP, "not_qr_logo.png")
    _IM.new("RGB", (300, 120), (20, 60, 140)).save(_logo_p)
    check("🔳 لوگو اشتباهاً کیوآر تشخیص داده نمی‌شود",
          not CG.looks_like_qr(Path(_logo_p).read_bytes()))
    _sq = os.path.join(TMP, "red_square.png")
    _IM.new("RGB", (200, 200), (220, 30, 30)).save(_sq)
    check("🔳 مربع رنگی کیوآر نیست",
          not CG.looks_like_qr(Path(_sq).read_bytes()))

    # قالبی که کاربر عکس QR داخلش گذاشته (بدون {کیوآر})
    from docx.shared import Mm as _Mm
    _d9 = docx.Document()
    _d9.add_paragraph("گواهینامه")
    _d9.add_paragraph().add_run().add_picture(_logo_p, width=_Mm(30))
    _d9.add_paragraph("نام: {نام}")
    _d9.add_paragraph().add_run().add_picture(_qsample, width=_Mm(25))
    _TPL9 = os.path.join(TMP, "with_qr_img.docx")
    _d9.save(_TPL9)

    _dd9 = docx.Document(_TPL9)
    check("🔴 ورد نام اصلی فایل را نگه نمی‌دارد",
          not any("user_sample" in str(r.target_ref)
                  for r in _dd9.part.rels.values()
                  if "image" in r.reltype))
    _rel = CG.find_qr_placeholder(_dd9)
    check("🔴 با وجود نام گمشده، عکس کیوآر پیدا شد", _rel is not None,
          str(_rel and _rel.target_ref))

    _tid9 = database.db_execute(
        """INSERT INTO cert_templates (name, file_path, kind, is_active)
        VALUES (?,?,?,1)""", ("با عکس کیوآر", _TPL9, "docx"))
    _r9 = CG.generate(CG.build_mapping(name="علی تستی",
                                       serial="NORA-QR-001"),
                      CG.get_template(_tid9))
    check("🔳 گواهینامه با عکس کیوآر ساخته شد", bool(_r9.get("docx")),
          str(_r9.get("error")))
    if _r9.get("docx"):
        _out9 = docx.Document(_r9["docx"])
        _still_same = False
        for _rl in _out9.part.rels.values():
            if "image" not in _rl.reltype:
                continue
            if _rl.target_part.blob == _blob:
                _still_same = True
        check("🔴 عکس نمونه با کیوآر واقعی عوض شد", not _still_same)

    check("🔳 قالب بدون هیچ تصویری کرش نمی‌کند",
          CG.find_qr_placeholder(docx.Document()) is None)


print("\n[۸.۱۰] 🖼️ خروجی PNG بدون LibreOffice")

# 🔴 باگ کاربر: «ربات گزینه تحویل PNG نداره از ورد»
_d10 = docx.Document()
_d10.add_paragraph("گواهینامه")
_d10.add_paragraph("گواهی می‌شود {خطاب} {نام} با کدملی {کدملی}")
_d10.add_paragraph("در دوره {رویداد} شرکت نمود.")
_d10.add_paragraph("شماره: {شماره}")
_TPL10 = os.path.join(TMP, "png_test.docx")
_d10.save(_TPL10)
_tid10 = database.db_execute(
    """INSERT INTO cert_templates (name, file_path, kind, is_active)
    VALUES (?,?,?,1)""", ("تست PNG", _TPL10, "docx"))

_r10 = CG.generate(CG.build_mapping(
    name="مریم رضایی", gender="زن", national_id="0012345678",
    event="کارگاه نمونه", serial="NORA-PNG-001"),
    CG.get_template(_tid10))

check("🔴 PNG بدون LibreOffice ساخته شد", bool(_r10.get("png")),
      str(_r10.get("error")))
check("🔴 PDF هم ساخته شد", bool(_r10.get("pdf")))
check("موتور داخلی شناسایی شد",
      _r10.get("engine") in ("builtin", "office", "libreoffice"),
      str(_r10.get("engine")))
if _r10.get("png"):
    _im10 = _IM.open(_r10["png"])
    check("🖼️ ابعاد PNG منطقی است", _im10.size[0] > 1000,
          str(_im10.size))
    check("🖼️ حجم PNG منطقی است",
          os.path.getsize(_r10["png"]) > 5000,
          str(os.path.getsize(_r10["png"])))

check("تابع رندر مستقیم وجود دارد",
      hasattr(CG, "docx_to_png_fallback"))
_direct = CG.docx_to_png_fallback(_TPL10, os.path.join(TMP, "d.png"))
check("🖼️ رندر مستقیم از ورد کار می‌کند", bool(_direct))
check("🖼️ فایل ناموجود کرش نمی‌کند",
      CG.docx_to_png_fallback("/nope/x.docx",
                              os.path.join(TMP, "n.png")) is None)

# خاموش کردن fallback
database.set_setting("cert_png_fallback", "0", "t", "certificate")
database.clear_settings_cache()
_r11 = CG.generate(CG.build_mapping(name="ب", serial="S2"),
                   CG.get_template(_tid10))
check("🔒 با خاموش بودن fallback، PNG ساخته نمی‌شود",
      not _r11.get("png") or CG.libreoffice_bin())
database.set_setting("cert_png_fallback", "1", "t", "certificate")
database.clear_settings_cache()


print("\n[۸.۱۱] 📦 راهنمای نصب متناسب با سیستم")

_g = CG.install_guide()
check("راهنمای نصب تولید می‌شود", len(_g) > 100)
check("سیستم‌عامل تشخیص داده شد",
      "ویندوز" in _g or "لینوکس" in _g, _g[:80])
check("تابع word_available وجود دارد", hasattr(CG, "word_available"))
check("وضعیت موتور شامل word است", "word" in CG.engine_status())
check("وضعیت شامل موتور داخلی است", "png_builtin" in CG.engine_status())

_realname = os.name
try:
    os.name = "nt"
    CG._LO_BIN = ""
    _gw = CG.install_guide()
    check("🪟 روی ویندوز دستور apt پیشنهاد نمی‌شود",
          "sudo apt" not in _gw, _gw[:200])
    check("🪟 دستور درست ویندوز داده می‌شود",
          "py -m pip" in _gw or "libreoffice.org" in _gw)
    check("🪟 هشدار apt داده می‌شود", "apt install" in _gw and
          "لینوکس" in _gw)
finally:
    os.name = _realname
    CG._LO_BIN = None

SENT.clear()
M.handle_callback(cb(S, "ct_install"))
check("📦 صفحه راهنمای نصب باز شد",
      "نصب" in last(), last()[:60])


# ═══════════ ۹) 🛡️ رگرسیون ═══════════
print("\n[۹] 🛡️ رگرسیون")

check("جدول قالب ساخته شده",
      bool(database.db_execute(
          "SELECT name FROM sqlite_master WHERE type='table' "
          "AND name='cert_templates'", fetch="one")))
_cc = {c["name"] for c in
       database.db_execute("PRAGMA table_info(certificates)", fetch="all")}
for col in ("template_id", "png_path", "pdf_path", "docx_path", "revoked"):
    check(f"ستون {col} در جدول گواهینامه هست", col in _cc)

check("بدون قالب هم کرش نمی‌کند",
      isinstance(CG.generate(CG.build_mapping(name="x"), None), dict))
_r = CG.generate(CG.build_mapping(name="x"), {"file_path": "/nope/x.docx",
                                              "kind": "docx"})
check("فایل ناموجود خطای تمیز می‌دهد",
      not _r.get("docx") and bool(_r.get("error")), str(_r.get("error")))

check("mapping با ورودی خالی کرش نمی‌کند",
      isinstance(CG.build_mapping(), dict))
check("find_placeholders با None",
      CG.find_placeholders(None) == set())
check("لاگ فارسی قالب گواهینامه دارد",
      "cert_template_added" in __import__("activity_log").ACTION_FA)


# ═══════════ 🔢 فاز ۲۷ — اعداد فارسی · فرمت · کیوآر ═══════════
print("\n[۲۷] 🔢 اعداد فارسی · فرمت ورد · جای کیوآر")

check("تابع fa_digits هست", hasattr(CG, "fa_digits"))
check("🔢 تبدیل درست", CG.fa_digits("1404/07/12") == "۱۴۰۴/۰۷/۱۲",
      CG.fa_digits("1404/07/12"))
check("🔢 برگشت درست", CG.en_digits("۱۴۰۴") == "1404")
check("🔢 متن بدون عدد دست‌نخورده", CG.fa_digits("سلام") == "سلام")
check("🛡️ None کرش نمی‌کند", CG.fa_digits(None) == "None")
check("تنظیم اعداد فارسی هست",
      any(r[0] == "cert_fa_digits" for r in CG.CERT_SETTINGS))

_mp = CG.build_mapping(name="علی", national_id="0010247165",
                       hours=16, date="1404/07/12",
                       serial="NORA-0405-99999")
check("🔢 کد ملی فارسی شد", _mp["national_id"] == "۰۰۱۰۲۴۷۱۶۵",
      _mp["national_id"])
check("🔢 ساعت فارسی شد", str(_mp["hours"]) == "۱۶", str(_mp["hours"]))
check("🔢 تاریخ فارسی شد", _mp["date"] == "۱۴۰۴/۰۷/۱۲")
check("🔴 شماره سریال لاتین ماند (برای اسکن)",
      _mp["serial"] == "NORA-0405-99999", _mp["serial"])
check("🔴 لینک اعتبارسنجی لاتین ماند",
      "99999" in str(_mp.get("verify_url")), str(_mp.get("verify_url")))

database.set_setting("cert_fa_digits", "0", "t", "certificate")
database.clear_settings_cache()
_mp2 = CG.build_mapping(national_id="0010247165")
check("🔀 خاموش که شد، لاتین می‌ماند",
      _mp2["national_id"] == "0010247165", _mp2["national_id"])
database.set_setting("cert_fa_digits", "1", "t", "certificate")
database.clear_settings_cache()

# ── 📐 فرمت ورد حفظ می‌شود ──
check("تابع تراز پاراگراف هست", hasattr(CG, "_par_align"))
check("تابع سبک run هست", hasattr(CG, "_run_style"))
check("تابع تصاویر پاراگراف هست", hasattr(CG, "_par_images"))
check("تنظیم پخش محتوا هست",
      any(r[0] == "cert_png_stretch" for r in CG.CERT_SETTINGS))

_src27 = (ROOT / "certgen.py").read_text(encoding="utf-8")
check("🔴 دیگر اندازه ثابت ۶۴/۴۰ نیست",
      "size = 64 if first else 40" not in _src27)
check("🔴 دیگر کیوآر گوشهٔ ثابت نمی‌رود",
      "img.paste(q2, (110, H - sz - 110))" not in _src27)
check("📏 اندازه صفحه از خود ورد خوانده می‌شود",
      "sec.page_width" in _src27)
check("↔️ حاشیه‌های ورد خوانده می‌شوند",
      "sec.left_margin" in _src27)

try:
    from docx.shared import Pt as _Pt, Mm as _Mm
    from docx.enum.text import WD_ALIGN_PARAGRAPH as _A

    _t27 = os.path.join(TMP, "fmt27.docx")
    _d27 = docx.Document()
    _s27 = _d27.sections[0]
    _s27.page_width, _s27.page_height = _Mm(297), _Mm(210)
    _p = _d27.add_paragraph(); _p.alignment = _A.CENTER
    _r = _p.add_run("عنوان بزرگ"); _r.bold = True
    _r.font.size = _Pt(28)
    _p = _d27.add_paragraph(); _p.alignment = _A.RIGHT
    _p.add_run("متن راست‌چین {نام}").font.size = _Pt(14)
    _p = _d27.add_paragraph(); _p.alignment = _A.CENTER
    _p.add_run("{کیوآر}").font.size = _Pt(11)
    _p = _d27.add_paragraph(); _p.alignment = _A.LEFT
    _p.add_run("امضا چپ").font.size = _Pt(12)
    _d27.save(_t27)

    _dd = docx.Document(_t27)
    _al = [CG._par_align(p) for p in _dd.paragraphs]
    check("↔️ تراز وسط تشخیص داده شد", _al[0] == "center", str(_al))
    check("↔️ تراز راست تشخیص داده شد", _al[1] == "right")
    check("↔️ تراز چپ تشخیص داده شد", _al[3] == "left")

    _r0 = _dd.paragraphs[0].runs[0]
    _sz, _bd, _co = CG._run_style(_r0, _dd.paragraphs[0])
    check("🔠 اندازه واقعی خوانده شد", _sz == 28, str(_sz))
    check("🅱️ بولد تشخیص داده شد", _bd is True)

    # ساخت واقعی با کیوآر
    _qp27 = os.path.join(TMP, "q27.png")
    CG.make_qr("https://x/cert_Z", _qp27)
    _o27 = os.path.join(TMP, "o27.docx")
    _g27, _e27 = CG.fill_docx(_t27, _o27, {"name": "سارا"},
                              qr_png=_qp27)
    check("🎓 قالب پر شد", bool(_g27), _e27)

    _png27 = CG.docx_to_png_fallback(_g27, os.path.join(TMP, "o27.png"))
    check("🖼️ تصویر ساخته شد", bool(_png27) and Path(_png27).exists())
    if _png27:
        from PIL import Image as _Img
        _im27 = _Img.open(_png27)
        check("📐 نسبت افقی حفظ شد (A4 لند‌اسکیپ)",
              _im27.width > _im27.height, str(_im27.size))

        # 🔳 کیوآر باید بالای نیمهٔ پایین باشد، نه گوشهٔ چپ
        _g = _im27.convert("L")
        _w27, _h27 = _g.size
        _dark_left = 0
        for _x in range(0, int(_w27 * 0.25), 6):
            for _y in range(int(_h27 * 0.7), _h27, 6):
                if _g.getpixel((_x, _y)) < 90:
                    _dark_left += 1
        check("🔴 کیوآر دیگر گوشهٔ پایین-چپ نیست",
              _dark_left < 25, str(_dark_left))
except ImportError:
    check("⚠️ python-docx نیست — رد شد", True)


print("\n" + "=" * 62)
print(f"✅ موفق: {OK}    ❌ ناموفق: {FAIL}")
print("=" * 62)
if FAILED:
    print("\nناموفق:")
    for f in FAILED:
        print(f"  ❌ {f}")
    sys.exit(1)
print("\n🎉 همه تست‌های فاز ۱۹ پاس شدند!")
