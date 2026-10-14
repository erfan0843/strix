#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
🧪 تست فاز ۳۴ — کیوآر شناور · شناسهٔ دولایه · اعتبار زمانی

درخواست‌های کاربر:
  ۱. «همچنان کیوآرکد جاش نیست»
  ۲. «مبنای استعلام چیه؟ هر رویداد یدونه شماره ثبت داره،
      چطوری شخصی که استعلام میزنه اطلاعات اون گواهی میاد بالا؟»
  ۳. «اعتبار گواهینامه‌ها پیش‌فرض ۲۴ ماه باشه و قابل تنظیم»
"""
import sys
import os
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

OK = 0
FAIL = 0
ERRORS = []


def check(label, cond, detail=""):
    global OK, FAIL
    if cond:
        OK += 1
    else:
        FAIL += 1
        ERRORS.append(f"{label} {detail}")


def section(t):
    print(f"\n{'─' * 62}\n{t}\n{'─' * 62}")


import database
import certgen as cg
import certverify as cv

HAVE_DOCX = True
try:
    import docx
    from docx.shared import Mm, Emu
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn
except ImportError:
    HAVE_DOCX = False

HAVE_QR = True
try:
    import qrcode
except ImportError:
    HAVE_QR = False

TMP = Path("/tmp/_p34")
TMP.mkdir(exist_ok=True)


# ══════════════════════════════════════════════════════════
section("۱) 🔳 ترتیب OOXML — ریشهٔ نامرئی بودن کیوآر")
# ══════════════════════════════════════════════════════════
if not (HAVE_DOCX and HAVE_QR):
    print("   ⏭️ python-docx یا qrcode نصب نیست")
else:
    from PIL import Image
    Image.new("RGB", (1492, 1054), (20, 50, 120)).save(str(TMP / "bg.png"))
    qrcode.make("TEST").save(str(TMP / "q.png"))

    d = docx.Document()
    s = d.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Mm(297), Mm(210)
    d.add_paragraph().add_run().add_picture(str(TMP / "bg.png"),
                                            width=Mm(296))
    d.add_paragraph("{نام}")
    d.save(str(TMP / "t.docx"))

    got, err = cg.fill_docx(str(TMP / "t.docx"), str(TMP / "o.docx"),
                            {"name": "علی"}, qr_png=str(TMP / "q.png"))
    check("قالب پر شد", bool(got), err)

    if got:
        d2 = docx.Document(str(TMP / "o.docx"))
        # 🔴 ترتیب اجباری استاندارد OOXML
        STD = ["simplePos", "positionH", "positionV", "extent",
               "effectExtent", "wrapNone", "wrapSquare", "docPr",
               "cNvGraphicFramePr", "graphic"]
        found = False
        for anc in d2.element.body.iter(qn("wp:anchor")):
            tags = [c.tag.split("}")[-1] for c in anc]
            idx = [STD.index(t) for t in tags if t in STD]
            check("ترتیب OOXML درست", idx == sorted(idx), str(tags))
            check("wrapNone بعد از extent",
                  tags.index("wrapNone") > tags.index("extent"),
                  str(tags))
            check("docPr بعد از wrapNone",
                  tags.index("docPr") > tags.index("wrapNone"),
                  str(tags))
            found = True
            break
        check("anchor ساخته شد", found)
        check('behindDoc="0"', 'behindDoc="0"' in d2.element.xml)
        check('relativeFrom="page"',
              'relativeFrom="page"' in d2.element.xml)

        # مختصات داخل صفحه
        for anc in d2.element.body.iter(qn("wp:anchor")):
            for tag, lim in (("wp:positionH", Emu(Mm(297))),
                             ("wp:positionV", Emu(Mm(210)))):
                pos = anc.find(qn(tag))
                if pos is None:
                    continue
                off = pos.find(qn("wp:posOffset"))
                if off is not None and off.text:
                    v = int(off.text)
                    check(f"{tag} داخل صفحه", 0 <= v <= int(lim),
                          f"{v}")
            break

        # پس‌زمینه دقیقاً اندازهٔ صفحه
        sec = d2.sections[0]
        pw, ph = int(sec.page_width), int(sec.page_height)
        for inl in d2.element.body.iter(qn("wp:inline")):
            e = inl.find(qn("wp:extent"))
            if e is not None and int(e.get("cx")) > pw * 0.8:
                check("پس‌زمینه = عرض صفحه", int(e.get("cx")) == pw)
                check("پس‌زمینه = ارتفاع صفحه", int(e.get("cy")) == ph)
                break


# ══════════════════════════════════════════════════════════
section("۲) 🔑 شناسهٔ دولایه")
# ══════════════════════════════════════════════════════════
cv.init_tables()
for fn in ["new_person_code", "build_full_code", "normalize_code",
           "stamp", "find", "status_of", "report", "verify_link"]:
    check(f"cv.{fn}", hasattr(cv, fn))

# کد فردی
codes = [cv.new_person_code() for _ in range(6)]
check("کدها یکتا", len(set(codes)) == 6)
check("طول درست", all(len(c) == 6 for c in codes), str(codes[0]))
check("بدون O/0/I/1",
      not any(ch in "O0I1" for c in codes for ch in c), str(codes))

# ترکیب
check("ترکیب کامل",
      cv.build_full_code("1421/5472/د", "A7K2") == "1421/5472/د-A7K2")
check("بدون شمارهٔ نامه",
      cv.build_full_code("", "A7K2") == "A7K2")
check("بدون کد فردی",
      cv.build_full_code("1421/د", "") == "1421/د")

# پاکسازی ورودی
for raw, want in [("a7k2m9", "A7K2M9"), ("A7 K2 M9", "A7K2M9"),
                  ("  a7k2m9  ", "A7K2M9"), ("A7-K2", "A7-K2")]:
    check(f"پاکسازی «{raw}»", cv.normalize_code(raw) == want,
          cv.normalize_code(raw))


# ══════════════════════════════════════════════════════════
section("۳) 📅 اعتبار زمانی")
# ══════════════════════════════════════════════════════════
check("۶ گزینهٔ اعتبار", len(cv.VALIDITY) == 6)
check("پیش‌فرض ۲۴ ماه", cv.default_months() == 24,
      str(cv.default_months()))

# افزودن ماه — حالت‌های سخت
for d, m, want in [
    ("1404/07/29", 24, "1406/07/29"),
    ("1404/01/31", 1, "1404/02/31"),      # هر دو ۳۱ روزه
    ("1404/06/31", 1, "1404/07/30"),      # ۳۱ → ۳۰
    ("1403/12/30", 12, "1404/12/29"),     # کبیسه → غیرکبیسه
    ("1404/12/29", 12, "1405/12/29"),
]:
    got = cv.add_months(d, m)
    check(f"{d} + {m} ماه", got == want, f"{got} ≠ {want}")
check("بدون ماه", cv.add_months("1404/01/01", 0) == "")
check("تاریخ خالی", cv.add_months("", 24) == "")

# انقضا
check("آینده معتبر", not cv.is_expired({"expires_on": "1499/01/01"}))
check("گذشته منقضی", cv.is_expired({"expires_on": "1390/01/01"}))
check("بدون تاریخ معتبر", not cv.is_expired({"expires_on": ""}))

# روز باقی‌مانده
t = cv.today_jalali()
check("امروز شمسی", bool(t) and len(t) == 10, t)
d0 = cv.days_left({"expires_on": t})
check("امروز = صفر روز", d0 == 0, str(d0))
d2y = cv.days_left({"expires_on": cv.add_months(t, 24)})
check("۲ سال ≈ ۷۳۰ روز", d2y and 725 <= d2y <= 735, str(d2y))
check("بدون تاریخ None", cv.days_left({"expires_on": ""}) is None)


# ══════════════════════════════════════════════════════════
section("۴) 🎯 سناریوی واقعی — یک رویداد، چند نفر")
# ══════════════════════════════════════════════════════════
import events as EV
import attendance as ATT
EV.init_event_tables()
ATT.init_attendance_tables()
cv.init_tables()

EVID = 943401
DOC = "1421/5472/5000/د"
database.db_execute("DELETE FROM certificates WHERE event_id=?", (EVID,))
database.db_execute("DELETE FROM events WHERE id=?", (EVID,))
database.db_execute(
    "INSERT INTO events (id,title,event_type,status,cert_doc_no,"
    "cert_valid_months,created_by) VALUES (?,?,?,?,?,?,?)",
    (EVID, "همایش مشترک", "workshop", "finished", DOC, 24, 1))

check("شمارهٔ نامهٔ رویداد", cv.event_doc_no(EVID) == DOC)
check("اعتبار رویداد", cv.default_months(EVID) == 24)

people = []
for u in (6101, 6102, 6103):
    cid_ = database.db_execute(
        "INSERT INTO certificates (event_id,bale_user_id,serial) "
        "VALUES (?,?,?)", (EVID, u, f"OLD-{u}"))
    info = cv.stamp(cid_, EVID)
    people.append((cid_, info))

check("۳ گواهینامه", len(people) == 3)
check("🔴 شمارهٔ نامه مشترک",
      len({i["doc_no"] for _c, i in people}) == 1,
      str([i["doc_no"] for _c, i in people]))
check("🔑 کد فردی یکتا",
      len({i["person_code"] for _c, i in people}) == 3)
check("🔗 شناسهٔ کامل یکتا",
      len({i["full_code"] for _c, i in people}) == 3)
check("شناسه شامل شمارهٔ نامه",
      all(i["full_code"].startswith(DOC) for _c, i in people))
check("تاریخ صدور ثبت شد",
      all(i["issued_on"] for _c, i in people))
check("تاریخ انقضا ۲ سال بعد",
      all(i["expires_on"] == cv.add_months(i["issued_on"], 24)
          for _c, i in people))

# 🔍 هر شناسه دقیقاً همان نفر را برمی‌گرداند
for cid_, info in people:
    r, why = cv.find(code=info["full_code"])
    check(f"استعلام {info['person_code']}",
          r and r["id"] == cid_, why)

# فقط کد فردی
cid0, info0 = people[0]
r, _ = cv.find(code=info0["person_code"])
check("کد فردی تنها هم کار می‌کند", r and r["id"] == cid0)

# شمارهٔ نامه + کد فردی
r, why = cv.find(doc_no=DOC, national_id=info0["person_code"])
check("نامه + کد فردی", r and r["id"] == cid0, why)

# ورودی نامرتب
r, _ = cv.find(code=info0["person_code"].lower())
check("حروف کوچک هم پذیرفته می‌شود", r and r["id"] == cid0)

# کد اشتباه
r, why = cv.find(code="NOTEXIST99")
check("کد نامعتبر رد شد", r is None and bool(why), why)

# 🔴 مهم‌ترین: شمارهٔ نامهٔ تنها نباید یک نفر خاص برگرداند
r, _ = cv.find(code=DOC)
check("🔴 شمارهٔ نامهٔ تنها کافی نیست", r is None,
      "نباید کسی را برگرداند — کد فردی لازم است")


# ══════════════════════════════════════════════════════════
section("۵) 📋 گزارش استعلام")
# ══════════════════════════════════════════════════════════
row = database.db_execute("SELECT * FROM certificates WHERE id=?",
                          (cid0,), fetch="one")
k, icon, lbl = cv.status_of(row)
check("وضعیت معتبر", k == "valid", k)
rep = cv.report(row)
check("گزارش ساخته شد", bool(rep))
check("گزارش: شمارهٔ نامه", DOC in rep)
check("گزارش: کد فردی", info0["person_code"] in rep)
check("گزارش: تاریخ صدور", "تاریخ صدور" in rep)
check("گزارش: اعتبار", "اعتبار تا" in rep)
check("گزارش: عنوان رویداد", "همایش مشترک" in rep)

# منقضی
database.db_execute("UPDATE certificates SET expires_on='1390/01/01' "
                    "WHERE id=?", (cid0,))
row2 = database.db_execute("SELECT * FROM certificates WHERE id=?",
                           (cid0,), fetch="one")
k2, _i, _l = cv.status_of(row2)
check("وضعیت منقضی", k2 == "expired", k2)
check("گزارش منقضی", "تمام رسیده" in cv.report(row2)
      or "پایان رسیده" in cv.report(row2))

# باطل
database.db_execute("UPDATE certificates SET revoked=1 WHERE id=?",
                    (cid0,))
row3 = database.db_execute("SELECT * FROM certificates WHERE id=?",
                           (cid0,), fetch="one")
check("وضعیت باطل", cv.status_of(row3)[0] == "revoked")
check("باطل مقدم بر انقضا", "باطل" in cv.report(row3))
database.db_execute("UPDATE certificates SET revoked=0, "
                    "expires_on=? WHERE id=?",
                    (info0["expires_on"], cid0))

# پوشاندن کد ملی
m = cv._mask_nid("0010247165")
check("کد ملی پوشانده شد", "•" in m and "0010247165" not in m, m)
check("کد ملی کوتاه امن", cv._mask_nid("123") == "—")

# آمار استعلام
cv.log_verify(row, 999, "qr", "valid")
row4 = database.db_execute("SELECT * FROM certificates WHERE id=?",
                           (cid0,), fetch="one")
check("شمارندهٔ استعلام", (row4.get("verify_count") or 0) >= 1)


# ══════════════════════════════════════════════════════════
section("۶) ⚙️ تنظیمات و پارامترها")
# ══════════════════════════════════════════════════════════
_ST = getattr(cg, "CERT_SETTINGS", [])
keys = [r[0] for r in _ST]
for k in ("cert_valid_months", "cert_doc_no", "cert_code_len",
          "cert_show_nid", "cert_expiry_warn"):
    check(f"تنظیم {k}", k in keys)
dflt = {r[0]: r[4] for r in _ST}
check("پیش‌فرض ۲۴ ماه", dflt.get("cert_valid_months") == "24")

# پارامترهای قالب
for ph, want in [("شماره نامه", "doc_no"), ("تاریخ صدور", "issued_on"),
                 ("تاریخ انقضا", "expires_on"),
                 ("اعتبار تا", "expires_on"),
                 ("کد اعتبارسنجی", "verify_code")]:
    check(f"پارامتر {{{ph}}}", cg.resolve_key(ph) == want,
          str(cg.resolve_key(ph)))

check("doc_no لاتین می‌ماند", "doc_no" in cg.NO_FA_KEYS)

# ستون‌های رویداد
evc = database.db_execute("PRAGMA table_info(events)", fetch="all") or []
en = [c["name"] for c in evc]
check("ستون cert_doc_no", "cert_doc_no" in en)
check("ستون cert_valid_months", "cert_valid_months" in en)

cc = database.db_execute("PRAGMA table_info(certificates)",
                         fetch="all") or []
cn = [c["name"] for c in cc]
for c in ("doc_no", "person_code", "full_code", "issued_on",
          "expires_on", "valid_months", "verify_count"):
    check(f"ستون {c}", c in cn)


# ══════════════════════════════════════════════════════════
section("۷) روتر و اتصال‌ها")
# ══════════════════════════════════════════════════════════
mn = Path("main.py").read_text(encoding="utf-8")
check("import certverify", "import certverify" in mn)
check("init_tables", "cv_mod.init_tables()" in mn)
check("deep-link cv_", 'param.startswith("cv_")' in mn)
check("استعلام یکپارچه", "cv_mod.find(code=serial)" in mn)
check("cron هشدار انقضا", "cv_mod.warn_expiring()" in mn)
for r in ["cv_ask", "cv_bycode", "cv_bydoc", "cve_doc_",
          "cve_val_", "cve_st_", "cvv_"]:
    check(f"مسیر {r}", r in mn)

ci = Path("certissue.py").read_text(encoding="utf-8")
check("certissue شناسه می‌زند", "_cv.stamp" in ci)
ec = Path("eventcert.py").read_text(encoding="utf-8")
check("eventcert شناسه می‌زند", "_cv.stamp" in ec)
check("دکمهٔ استعلام", 'cve_{eid}' in ec)
cb = Path("certbatch.py").read_text(encoding="utf-8")
check("certbatch شناسه می‌زند", "_cv.stamp" in cb)


# ══════════════════════════════════════════════════════════
section("۸) پاکسازی و واژه‌نامه")
# ══════════════════════════════════════════════════════════
database.db_execute("DELETE FROM certificates WHERE event_id=?", (EVID,))
database.db_execute("DELETE FROM events WHERE id=?", (EVID,))
database.db_execute("DELETE FROM cert_verifies WHERE by_user=999")
check("پاکسازی", not database.db_scalar(
    "SELECT COUNT(*) FROM certificates WHERE event_id=?", (EVID,)))

for fn in ["certverify.py"]:
    bad = []
    inq = False
    for i, ln in enumerate(Path(fn).read_text(encoding="utf-8")
                           .splitlines(), 1):
        if "«" in ln and "»" not in ln:
            inq = True
        elif "»" in ln and "«" not in ln:
            inq = False
            continue
        if inq or ln.lstrip().startswith("#"):
            continue
        if "«" in ln and "»" in ln:
            continue
        if re.search(r"گواهی(?!نامه)(?! می)(?!\w)", ln):
            bad.append(str(i))
    check(f"«گواهینامه» در {fn}", not bad, " · ".join(bad[:3]))


print(f"\n{'=' * 62}")
print(f"✅ موفق: {OK}    ❌ ناموفق: {FAIL}")
print("=" * 62)
if ERRORS:
    print("\n❌ خطاها:")
    for e in ERRORS[:30]:
        print(f"   • {e}")
    sys.exit(1)
print("\n🎉 همه تست‌های فاز ۳۴ پاس شدند!")
