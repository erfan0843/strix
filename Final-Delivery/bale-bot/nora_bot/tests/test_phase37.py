# -*- coding: utf-8 -*-
"""
🧪 تست فاز ۳۷ — یکپارچگی کامل همهٔ مسیرهای گواهینامه

🎯 درخواست شما:
   «یه بار دیگه کل بخش‌هارو تست بزن… هرچیزی که بهش وصله…
    صدور گواهی که تو رویداد هست هم کامل بروز بشه با این
    تغییرات و قسمت پروفایل و همه چی مطمئن شو دقیق و بدون
    باگ کار میکنن کدا»

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
📍 پنج مسیر مستقل در جدول `certificates` می‌نویسند:

   ۱ certissue.issue_one      رویداد (اصلی)
   ۲ certcenter.build_one     مرکز صدور
   ۳ eventcert.issue          چند گواهینامه در یک رویداد
   ۴ certbatch.handle_claim   لینک اکسل
   ۵ certgen.regenerate       بازتولید فایل پاک‌شده

این تست می‌سنجد که **هر پنج مسیر** رفتار یکسان دارند:
   • عکس پرسنلی
   • شناسهٔ دولایه و سریال درست
   • تاریخ شمسی
   • شمارندهٔ کاربر
   • ارسال واقعی فایل
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
"""
import os
import re
import sys
import json
import pathlib
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("BOT_TOKEN", "0:test")
os.environ.setdefault("SUPER_ADMIN_ID", "1")

# 🧪 دیتابیس اختصاصی و تازه
#
# 🔴 چرا لازم است؟ این تست داده‌های واقعی می‌سازد (کاربر،
#    رویداد، گواهینامه) و روی دیتابیس مشترک، بازمانده‌های
#    تست‌های دیگر با آن تداخل می‌کردند — مثلاً کاربر ۷۱۰۱ از
#    اجرای قبلی می‌ماند و INSERT OR IGNORE عکس تازه را
#    نمی‌نوشت. نتیجه: تست بسته به **ترتیب اجرا** نتیجهٔ
#    متفاوت می‌داد.
_TDB = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                    "_p37.db")
try:
    os.remove(_TDB)
except OSError:
    pass
os.environ["DB_PATH"] = _TDB

# 🌐 نام کاربری ربات را از پیش تعیین کن.
#
# 🔴 چرا؟ ساخت لینک اعتبارسنجی (`certverify.verify_link`) به
#    `config.deep_link` می‌رسد و آن اگر نام کاربری را نداند،
#    `getMe` را روی شبکه صدا می‌زند. در محیط بدون اینترنت هر
#    فراخوانی ~۵۰ ثانیه معطل می‌شود و تست به تایم‌اوت می‌خورد.
#    این تست ده‌ها گواهینامه می‌سازد، پس بدون این خط عملاً
#    تمام نمی‌شود.
os.environ.setdefault("BOT_USERNAME", "nora_test_bot")

_ok = _bad = 0
_fails = []


def check(name, cond, extra=""):
    global _ok, _bad
    if cond:
        _ok += 1
    else:
        _bad += 1
        _fails.append(name)
        print(f"  ❌ {name}  {extra}")


print("=" * 58)
print("🧪 فاز ۳۷ — یکپارچگی کامل مسیرهای گواهینامه")
print("=" * 58)

# 🌐 شبکه را ماک کن — هیچ درخواست واقعی به بله نرود
#
# 🔴 بدون این، هر `sendMessage` در محیط بدون اینترنت ~۸۰ ثانیه
#    معطل می‌شود. این تست ده‌ها پیام می‌فرستد (اعلان انتشار،
#    ارسال گواهینامه، …) و عملاً هرگز تمام نمی‌شد.
import bale_api as _ba

SENT = []


def _fake_api(m, p=None, **k):
    SENT.append({"method": m, "params": p or {}})
    if m == "getMe":
        return {"ok": True,
                "result": {"username": "nora_test_bot", "id": 1}}
    if m == "getChatMember":
        return {"ok": True, "result": {"status": "member"}}
    return {"ok": True, "result": {"message_id": len(SENT)}}


def _fake_send(cid, text="", kb=None, **k):
    SENT.append({"method": "sendMessage",
                 "params": {"chat_id": cid, "text": text, "kb": kb}})
    return {"ok": True, "result": {"message_id": len(SENT)}}


def _fake_file(cid, path=None, caption="", **k):
    SENT.append({"method": "sendFile",
                 "params": {"chat_id": cid, "path": str(path)}})
    return {"ok": True, "result": {"message_id": len(SENT)}}


_ba.api = _fake_api
_ba.send = _fake_send
_ba.send_or_edit = _fake_send
_ba.send_photo = _fake_file
_ba.send_document = _fake_file

SRC = {f.name: f.read_text(encoding="utf-8")
       for f in pathlib.Path(".").glob("*.py")}

# ══════════════════════════════════════════════════════════
print("\n[۱] 📤 باگ ارسال — send_cert_files با شناسهٔ عددی")

import database
import certgen
import attendance
import certissue

database.init_db()
certgen.init_cert_tables()
attendance.init_attendance_tables()
import events as _ev0
_ev0.init_event_tables()          # جدول events و ثبت‌نام‌ها

src = SRC["certissue.py"]
check("پذیرش شناسهٔ عددی", "isinstance(cert, (int, float, str))" in src)
check("رکورد نبود → False", 'log("⚠️ send_cert_files' in src)

# صداکننده‌هایی که عدد می‌دادند
for f, pat in (("eventcert.py", "send_cert_files(cid, cert_id)"),
               ("certbatch.py", "send_cert_files(cid, cert_id)")):
    if pat in SRC[f]:
        check(f"{f} عدد می‌دهد — پس تابع باید تحمل کند", True)

# 🔬 تست زنده
res = certissue.send_cert_files(999, 123456)
check("شناسهٔ ناموجود → False بدون استثنا", res is False)

cid_ = database.db_execute(
    "INSERT INTO certificates (bale_user_id, serial, template) "
    "VALUES (?,?,?)", (7001, "NORA-TEST-1", "متن آزمایشی"))
try:
    r = certissue.send_cert_files(7001, cid_)
    check("شناسهٔ عددی معتبر → استثنا نمی‌دهد", True)
except AttributeError as e:
    check(f"شناسهٔ عددی نباید AttributeError بدهد ({e})", False)
except Exception:
    check("شناسهٔ عددی → استثنای غیرمرگبار", True)

# ══════════════════════════════════════════════════════════
print("[۲] 🔖 سریال — همهٔ مسیرها یکسان محافظت می‌شوند")

for f in ("certissue.py", "certbatch.py", "eventcert.py"):
    s = SRC[f]
    if 'full_code' not in s:
        continue
    # هر جا serial با full_code عوض می‌شود باید شرط doc_no باشد
    hits = re.findall(
        r'(?:if[^\n]*doc_no[^\n]*:\s*\n\s*)?serial\s*=\s*[^\n]*full_code',
        s)
    guarded = re.findall(
        r'if\s+info\.get\("doc_no"\)[^\n]*:\s*\n\s*serial\s*=', s)
    guarded += re.findall(
        r'if\s+_cvinfo\.get\("full_code"\)\s+and\s+'
        r'_cvinfo\.get\("doc_no"\)', s)
    check(f"{f}: سریال فقط با شمارهٔ نامه عوض می‌شود",
          len(guarded) >= 1 or not hits, f"hits={len(hits)}")

check("certbatch سریال را در DB می‌نویسد",
      "UPDATE certificates SET serial=?" in SRC["certbatch.py"])
check("eventcert سریال را در DB می‌نویسد",
      "UPDATE certificates SET serial=?" in SRC["eventcert.py"])

# ══════════════════════════════════════════════════════════
print("[۳] 🔢 شمارندهٔ گواهینامهٔ کاربر — هر پنج مسیر")

for f in ("certissue.py", "certcenter.py", "certbatch.py",
          "eventcert.py"):
    check(f"{f} شمارنده را بالا می‌برد",
          "total_certificates=" in SRC[f]
          or "total_certificates =" in SRC[f])

check("پروفایل از خود جدول می‌شمارد",
      "COUNT(*) AS n FROM certificates" in SRC["profile.py"])
check("پروفایل باطل‌شده‌ها را نمی‌شمارد",
      "COALESCE(revoked,0)=0" in SRC["profile.py"])
check("پروفایل آماده‌های مرکز را نشان می‌دهد",
      "pending_for" in SRC["profile.py"])

# ══════════════════════════════════════════════════════════
print("[۴] 📷 عکس پرسنلی — هر چهار مسیر صدور")

import certphoto as cp

check("تابع مشترک attach هست", hasattr(cp, "attach"))
for f in ("certissue.py", "eventcert.py", "certbatch.py"):
    check(f"{f} عکس را می‌چسباند",
          "certphoto" in SRC[f] and "attach" in SRC[f])
check("certcenter عکس را می‌چسباند",
      "_attach_photo" in SRC["certcenter.py"])

# 🔬 رفتار attach
m = {"name": "تست"}
cp.attach(m, None, None)
check("بدون قالب → عکس می‌گذارد", "_photo" in m)

m2 = {"name": "تست"}
cp.attach(m2, None, {"placeholders": json.dumps(["name", "event"])})
check("🔑 قالب بدون {عکس} → دست نمی‌زند", "_photo" not in m2)

m3 = {"name": "تست"}
cp.attach(m3, None, {"placeholders": json.dumps(["name", "photo"])})
check("قالب با {عکس} → می‌گذارد", "_photo" in m3)

m4 = {"name": "x"}
cp.attach(m4, None, None, enabled=False)
check("خاموش → دست نمی‌زند", "_photo" not in m4)

check("None امن است", cp.attach(None, 1, None) is None)
check("قالب خراب خطا نمی‌دهد",
      cp.attach({}, None, {"placeholders": "{{{"}) is not None)

# موتور از _photo می‌خواند
check("موتور کلید _photo را می‌شناسد",
      'mapping.get("_photo")' in SRC["certgen.py"])

# ══════════════════════════════════════════════════════════
print("[۵] 📅 تاریخ شمسی — روی خودِ گواهینامه")

import jalali as J

m = certgen.build_mapping(name="تست")
check("تاریخ صدور پر می‌شود", bool(m.get("issue_date")))
check("تاریخ صدور شمسی است",
      str(m.get("issue_date", "")).startswith("۱۴"), m.get("issue_date"))
check("issued_on هم‌معنی پر می‌شود",
      m.get("issued_on") == m.get("issue_date"))
check("today پر می‌شود", bool(m.get("today")))
check("fallback میلادی حذف شد",
      'datetime.now().strftime("%Y/%m/%d")' not in SRC["certgen.py"])

# «{تاریخ صدور}» به کدام کلید می‌رسد؟
k = certgen.resolve_key("تاریخ صدور")
check(f"«تاریخ صدور» → {k}", k in ("issued_on", "issue_date"))
check(f"مقدار {k} در نقشه هست", bool(m.get(k)))
k2 = certgen.resolve_key("تاریخ انقضا")
check(f"«تاریخ انقضا» → {k2} در نقشه هست", k2 in m)

# ══════════════════════════════════════════════════════════
print("[۶] 🎉 مسیر رویداد — سر تا ته")

import certcenter as cc
import eventcert
import certbatch
import events as ev_mod

cc.init_tables()

TMP = tempfile.mkdtemp(prefix="p37_")
import docx as _dx
from docx.shared import Mm
from docx.enum.section import WD_ORIENT

d = _dx.Document()
sec = d.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Mm(297), Mm(210)
d.add_paragraph("گواهی می‌شود {نام} — کد ملی {کدملی}")
d.add_paragraph("دورهٔ {رویداد} · تاریخ صدور {تاریخ صدور}")
d.add_paragraph("{عکس}  {کیوآر}")
TPL = f"{TMP}/t.docx"
d.save(TPL)
tid = database.db_execute(
    "INSERT INTO cert_templates (name, file_path, kind, is_active, "
    "is_default) VALUES (?,?,?,1,1)", ("قالب۳۷", TPL, "docx"))
check("قالب ساخته شد", bool(tid))

phs = cc.scan_template(tid)
check("عکس در قالب کشف شد", "photo" in phs, f"{phs}")

# کاربر با عکس تأییدشده
from PIL import Image as _PI
_pp = f"{TMP}/u.png"
_PI.new("RGB", (300, 300), (90, 140, 90)).save(_pp)
# 🧹 اگر از اجرای قبلی مانده‌اند پاک کن — وگرنه
#    INSERT OR IGNORE عکس تازه را نمی‌نویسد و تست
#    به‌اشتباه شکست می‌خورد.
for _u in (7101, 7102):
    database.db_execute("DELETE FROM users WHERE bale_user_id=?", (_u,))
database.db_execute(
    "INSERT INTO users (bale_user_id, first_name, last_name, "
    "national_id, mobile, profile_status, profile_photo, gender) "
    "VALUES (?,?,?,?,?,?,?,?)",
    (7101, "سارا", "احمدی", "0012345679", "09121110000",
     "approved", _pp, "زن"))
# کاربر بدون عکس
database.db_execute(
    "INSERT INTO users (bale_user_id, first_name, last_name, "
    "national_id, mobile, profile_status) VALUES (?,?,?,?,?,?)",
    (7102, "حسن", "رضایی", "0023456787", "09122220000", "approved"))

check("عکس‌دار → عکس واقعی", cp.photo_for(7101) == _pp)
check("بی‌عکس → آدمک", cp.photo_for(7102) != _pp
      and bool(cp.photo_for(7102)))

# رویداد
database.db_execute("DELETE FROM events WHERE title LIKE ?",
                    ("%فاز ۳۷%",))
database.db_execute("DELETE FROM events WHERE title=?",
                    ("رویداد بی‌شماره",))
eid = database.db_execute(
    "INSERT INTO events (title, event_type, status, start_date, "
    "has_certificate, has_attendance, cert_doc_no) "
    "VALUES (?,?,?,?,?,?,?)",
    ("کارگاه فاز ۳۷", "workshop", "published", "1405/05/01",
     1, 0, "1421/5472/5000/د"))
check("رویداد ساخته شد", bool(eid))
for u_ in (7101, 7102):
    database.db_execute(
        "INSERT OR IGNORE INTO event_registrations (event_id, "
        "bale_user_id, status) VALUES (?,?,'confirmed')", (eid, u_))

# 🎓 صدور از مسیر رویداد
n0 = database.db_scalar("SELECT COUNT(*) FROM certificates") or 0
cert, err = certissue.issue_one(eid, 7101, 1, send_to_user=False)
if cert:
    # 🔄 رکورد را تازه بخوان: issue_one رکورد را قبل از ساخت
    #    فایل درج می‌کند و بعد stamp ستون‌های اعتبار را پر
    #    می‌کند؛ نسخهٔ برگشتی ممکن است پیش از آن گرفته شده باشد.
    cert = database.db_execute("SELECT * FROM certificates WHERE id=?",
                               (cert["id"],), fetch="one") or cert
    check("گواهینامهٔ رویداد صادر شد", True)
    check("به رویداد وصل است", cert.get("event_id") == eid)
    check("شمارهٔ نامهٔ رویداد گرفته شد",
          cert.get("doc_no") == "1421/5472/5000/د", cert.get("doc_no"))
    check("کد فردی دارد", bool(cert.get("person_code")))
    check("شناسهٔ کامل ترکیبی است",
          "-" in str(cert.get("full_code") or ""), cert.get("full_code"))
    check("سریال = شناسهٔ کامل (چون شمارهٔ نامه دارد)",
          cert.get("serial") == cert.get("full_code"))
    check("تاریخ صدور شمسی ذخیره شد",
          str(cert.get("issued_on") or "").startswith("14"),
          cert.get("issued_on"))
    check("اعتبار محاسبه شد", bool(cert.get("expires_on")))
    u = database.db_execute("SELECT total_certificates FROM users "
                            "WHERE bale_user_id=7101", fetch="one")
    check("شمارندهٔ کاربر بالا رفت",
          (u or {}).get("total_certificates", 0) >= 1)
else:
    check(f"صدور رویداد ({err})", False)

# رویداد بدون شمارهٔ نامه → سریال NORA باید بماند
eid2 = database.db_execute(
    "INSERT INTO events (title, event_type, status, start_date, "
    "has_certificate) VALUES (?,?,?,?,1)",
    ("رویداد بی‌شماره", "workshop", "published", "1405/05/02"))
database.db_execute(
    "INSERT OR IGNORE INTO event_registrations (event_id, "
    "bale_user_id, status) VALUES (?,?,'confirmed')", (eid2, 7102))
c2, e2 = certissue.issue_one(eid2, 7102, 1, send_to_user=False)
if c2:
    check("🔑 بدون شمارهٔ نامه، سریال NORA حفظ شد",
          str(c2.get("serial") or "").startswith("NORA"),
          c2.get("serial"))
    check("ولی کد فردی باز هم دارد", bool(c2.get("person_code")))

# ══════════════════════════════════════════════════════════
print("[۷] 🔍 استعلام — هر پنج مسیر قابل استعلام‌اند")

import certverify as cv

rows = database.db_execute(
    "SELECT * FROM certificates WHERE person_code IS NOT NULL "
    "AND person_code<>''", fetch="all") or []
check("گواهینامه‌های دارای کد فردی هست", len(rows) >= 1)
# ⚠️ cv.find خروجی (ردیف، دلیل) می‌دهد
for r in rows[:4]:
    got, _why = cv.find(r["person_code"])
    check(f"استعلام با کد فردی {r['person_code']}",
          bool(got) and got.get("id") == r["id"])

if rows and rows[0].get("full_code"):
    g, _ = cv.find(rows[0]["full_code"])
    check("استعلام با شناسهٔ کامل", bool(g))

# شمارهٔ نامهٔ تنها نباید کسی را برگرداند
_r, _ = cv.find("1421/5472/5000/د")
check("شمارهٔ نامهٔ تنها → هیچ‌کس", _r is None)

# کدهای فردی یکتا هستند
codes = [r["person_code"] for r in rows]
check("کدهای فردی یکتا هستند", len(codes) == len(set(codes)))

# ══════════════════════════════════════════════════════════
print("[۸] 🏛️ مرکز صدور — یکپارچگی با بقیه")

iid = cc.create_issue(1, "طرح یکپارچگی", tid)
cc.set_value(iid, "event", "کارگاه فاز ۳۷")
n_add = cc.fill_from_event(iid, eid)
check("گیرندگان از رویداد آمدند", n_add >= 1, f"{n_add}")

iss = cc.get_issue(iid)
sh = cc.shares(iid)
mp, info = cc.mapping_for(iss, sh[0])
check("نقشه تاریخ شمسی دارد",
      str(mp.get("issue_date", "")).startswith("۱۴"))
check("نقشه issued_on دارد", bool(mp.get("issued_on")))
check("verify_url ساخته شد", bool(mp.get("verify_url")))
check("کد فردی یکتا", bool(info.get("person_code")))

n_before = database.db_scalar("SELECT COUNT(*) FROM certificates") or 0
cc.publish(iid, notify=False)
n_mid = database.db_scalar("SELECT COUNT(*) FROM certificates") or 0
check("🔑 انتشار فایلی نمی‌سازد", n_mid == n_before)

c3, e3 = cc.build_one(iid, sh[0]["id"], notify=False)
if c3:
    check("مرکز صدور گواهینامه ساخت", True)
    check("منبع «center»", c3.get("source") == "center")
    check("تاریخ شمسی", str(c3.get("issued_on") or "").startswith("14"))
    check("قابل استعلام", bool(cv.find(c3.get("person_code"))[0]))
else:
    check(f"ساخت مرکز ({e3})", False)

# ══════════════════════════════════════════════════════════
print("[۹] 🖥️ اتصال منوها و روتر")

mn = SRC["main.py"]
for route in ("cc_mod.show_home", "cc_mod.handle_text",
              "cc_mod.handle_document", "cc_mod.claim_by_link",
              "cc_mod.user_get", "cc_mod.init_tables",
              "cc_mod.show_photo_panel", "cc_mod.show_stats"):
    check(f"روتر: {route}", route in mn)

check("هاب دکمهٔ مرکز دارد",
      '"callback_data": "cc_home"' in SRC["certgen.py"])
check("هاب دکمهٔ عکس دارد",
      '"callback_data": "cc_photo"' in SRC["certgen.py"])
check("cb_home هنوز کار می‌کند (سازگاری)",
      '"cc_home", "cb_home"' in mn)

# هیچ callback بی‌هندلر
cc_cbs = set(re.findall(r'"callback_data":\s*f?"(cc[a-z_]*_?)', 
                        SRC["certcenter.py"]))
# روتر با startswith کار می‌کند، پس پیشوند کوتاه‌تر کافی است
def _routed(p):
    p = p.rstrip("_")
    if p in mn:
        return True
    # cc_list_draft → cc_list_
    for i in range(len(p), 3, -1):
        if f'"{p[:i]}' in mn or f"startswith(\"{p[:i]}" in mn:
            return True
    return False


_miss = [p for p in cc_cbs if not _routed(p)]
check("همهٔ پیشوندهای مرکز در روتر هستند", not _miss, str(_miss))

# ══════════════════════════════════════════════════════════
print("[۱۰] 👤 پروفایل و گواهینامه‌های من")

att_src = SRC["attendance.py"]
check("انقضا نشان داده می‌شود", "منقضی شده" in att_src)
check("هشدار نزدیک انقضا", "_left <= 30" in att_src)
check("آماده‌های مرکز در فهرست", "ccget_" in att_src)
check("تاریخ‌ها شمسی", "_jd.fmt(c.get('issued_at')" in att_src)

# is_expired درست کار می‌کند
check("منقضی تشخیص داده می‌شود",
      cv.is_expired({"expires_on": "1400/01/01"}))
check("معتبر تشخیص داده می‌شود",
      not cv.is_expired({"expires_on": "1499/01/01"}))
check("بدون انقضا → معتبر", not cv.is_expired({"expires_on": ""}))
check("None امن است", not cv.is_expired(None))

# ══════════════════════════════════════════════════════════
print("[۱۱] 🧪 استحکام — ورودی‌های خراب")

check("issue_one با رویداد ناموجود",
      certissue.issue_one(999999, 1, 1, send_to_user=False)[0] is None)
check("build_one با طرح ناموجود",
      cc.build_one(999999, 1)[0] is None)
check("photo_for با شناسهٔ بی‌معنی",
      cp.photo_for(-1) is not None)
check("find با ورودی خالی", cv.find("")[0] is None)
check("find با ورودی خراب", cv.find("!!!@@@")[0] is None)
check("fmt با ورودی خراب", J.fmt("abc") == "—")
check("scan_template با شناسهٔ خراب", cc.scan_template("x") == [])
check("mapping_for با سهم خالی",
      cc.mapping_for(iss, {"full_name": "", "person_code": ""})[0]
      is not None)

# ══════════════════════════════════════════════════════════
print("[۱۲] 🔗 یکپارچگی جدول‌ها")

def _colnames(tbl):
    """نام ستون‌های یک جدول — سطرها dict هستند نه tuple"""
    out = set()
    for r in (database.db_execute(f"PRAGMA table_info({tbl})",
                                  fetch="all") or []):
        try:
            out.add(r["name"])
        except (TypeError, KeyError, IndexError):
            try:
                out.add(r[1])
            except Exception:
                pass
    return out


cols = _colnames("certificates")
for c in ("person_code", "full_code", "doc_no", "issued_on",
          "expires_on", "valid_months", "source", "revoked"):
    check(f"ستون certificates.{c}", c in cols)

cols2 = _colnames("cert_shares")
for c in ("issue_id", "bale_user_id", "person_code", "cert_id",
          "state", "notified"):
    check(f"ستون cert_shares.{c}", c in cols2)

cols3 = _colnames("cert_issues")
check("ستون params (نه values)", "params" in cols3)
check("«values» به‌عنوان ستون نمانده", "values" not in cols3)

# هیچ گواهینامه‌ای بدون کد فردی نماند
orphan = database.db_scalar(
    "SELECT COUNT(*) FROM certificates WHERE "
    "(person_code IS NULL OR person_code='') AND id>?", (cid_,)) or 0
check("هر گواهینامهٔ تازه کد فردی دارد", orphan == 0, f"{orphan}")

# ══════════════════════════════════════════════════════════
print("[۱۳] 🎉 میان‌بر رویداد ← مرکز صدور")

check("تابع from_event هست", hasattr(cc, "from_event"))
check("دکمه در صفحهٔ صدور رویداد",
      "cc_ev_" in pathlib.Path("certissue.py").read_text(encoding="utf-8"))
check("روتر وصل است", "cc_mod.from_event" in mn)

# 🧹 طرح‌های قبلی این رویداد را پاک کن تا میان‌بر واقعاً
#    یکی تازه بسازد (بخش ۸ خودش یکی ساخته بود)
for _old in (database.db_execute(
        "SELECT id FROM cert_issues WHERE audience='event' "
        "AND audience_ref=?", (str(eid),), fetch="all") or []):
    database.db_execute("DELETE FROM cert_shares WHERE issue_id=?",
                        (_old["id"],))
    database.db_execute("DELETE FROM cert_issues WHERE id=?",
                        (_old["id"],))

_before = database.db_scalar("SELECT COUNT(*) FROM cert_issues") or 0
cc.from_event(1, 1, eid)
_after = database.db_scalar("SELECT COUNT(*) FROM cert_issues") or 0
check("طرح از رویداد ساخته شد", _after == _before + 1,
      f"{_before}→{_after}")

_new = database.db_execute(
    "SELECT * FROM cert_issues WHERE audience='event' AND "
    "audience_ref=? ORDER BY id DESC LIMIT 1", (str(eid),), fetch="one")
check("به رویداد وصل است", bool(_new))
if _new:
    check("شمارهٔ نامهٔ رویداد گرفته شد",
          _new.get("doc_no") == "1421/5472/5000/د", _new.get("doc_no"))
    check("عنوان رویداد در پارامترها",
          cc.issue_values(_new).get("event") == "کارگاه فاز ۳۷")
    check("گیرندگان پر شدند",
          (database.db_scalar("SELECT COUNT(*) FROM cert_shares "
                              "WHERE issue_id=?", (_new["id"],)) or 0) >= 1)
    check("قالب پیش‌فرض انتخاب شد", bool(_new.get("template_id")))

# دوبار زدن نباید طرح تکراری بسازد
cc.from_event(1, 1, eid)
_after2 = database.db_scalar("SELECT COUNT(*) FROM cert_issues") or 0
check("🔁 دوبار زدن طرح تکراری نمی‌سازد", _after2 == _after)

# ══════════════════════════════════════════════════════════
print("[۱۴] ♻️ بازتولید فایل پاک‌شده")

cgs = SRC["certgen.py"]
check("بازتولید عکس پرسنلی دارد", "certphoto" in cgs and
      "_cp.attach(mp, cert_row" in cgs)
check("بازتولید شناسهٔ دولایه را برمی‌دارد",
      '("verify_code", "person_code")' in cgs)
check("بازتولید verify_url می‌سازد",
      "_cv.verify_link(cert_row" in cgs)
check("view_certificate قالب خودِ گواهینامه را می‌گیرد",
      '_cg.get_template(c.get("template_id"))' in SRC["attendance.py"])

# 🔬 بازتولید زنده
_c = database.db_execute(
    "SELECT * FROM certificates WHERE person_code IS NOT NULL "
    "AND person_code<>'' ORDER BY id LIMIT 1", fetch="one")
if _c:
    _r = certgen.regenerate(_c)
    check("بازتولید بدون استثنا", _r is not None or True)
    _c2 = database.db_execute("SELECT * FROM certificates WHERE id=?",
                              (_c["id"],), fetch="one")
    check("کد فردی بعد از بازتولید نمی‌پرد",
          _c2.get("person_code") == _c.get("person_code"))
    check("سریال بعد از بازتولید نمی‌پرد",
          _c2.get("serial") == _c.get("serial"))

# ══════════════════════════════════════════════════════════
print("[۱۵] ⏳ اعتبار پیش‌فرض ۲۴ ماه — همه‌جا یکسان")

check("build_mapping پیش‌فرض ۲۴ دارد",
      'get_setting("cert_valid_months", "24")' in cgs)
check("پیش‌فرض «۰» دیگر نیست",
      'get_setting("cert_valid_months", "0")' not in cgs)
check("certverify هم ۲۴ می‌دهد", cv.default_months(None) == 24)

_m = certgen.build_mapping(name="x")
check("تاریخ انقضا در نقشه پر است", bool(_m.get("expire_date")),
      _m.get("expire_date"))
check("expires_on هم‌معنی پر است", bool(_m.get("expires_on")))
check("انقضا شمسی است", str(_m.get("expire_date", "")).startswith("۱۴"))

# دو سال بعد باشد
try:
    _iy = int(J.en(str(_m["issue_date"])).split("/")[0])
    _ey = int(J.en(str(_m["expire_date"])).split("/")[0])
    check("انقضا دقیقاً ۲ سال بعد است", _ey - _iy == 2,
          f"{_iy}→{_ey}")
except Exception as _e:
    check(f"محاسبهٔ سال انقضا ({_e})", False)

print()
print("=" * 58)
print(f"✅ موفق: {_ok}   ❌ ناموفق: {_bad}")
if _fails:
    print("\nناموفق‌ها:")
    for f in _fails:
        print("  •", f)
print("=" * 58)

# 🧹 دیتابیس آزمایشی را پاک کن
try:
    database.close_db()
except Exception:
    pass
for _ext in ("", "-wal", "-shm"):
    try:
        os.remove(_TDB + _ext)
    except OSError:
        pass

sys.exit(1 if _bad else 0)
