#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
🧪 تست فاز ۲۹ — سه مشکل گواهینامه از روی قالب

۱. فضای سفید اضافه  → PixelHeight غلط + برش نشدن حاشیه
۲. طرح جابه‌جا       → همان ریشه
۳. کیوآر درج نمی‌شد  → placeholder نبود، بی‌صدا رد می‌شد
"""
import sys
import os
import io
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

OK = 0
FAIL = 0
ERRORS = []


def check(label, cond, detail=""):
    global OK, FAIL
    if cond:
        OK += 1
    else:
        FAIL += 1
        ERRORS.append(f"{label} {detail}")


def section(t):
    print(f"\n{'─' * 62}\n{t}\n{'─' * 62}")


import certgen as cg

HAVE_DOCX = True
try:
    import docx
    from docx.shared import Mm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT
except ImportError:
    HAVE_DOCX = False

HAVE_QR = True
try:
    import qrcode
except ImportError:
    HAVE_QR = False

TMP = Path("/tmp/_p29")
TMP.mkdir(exist_ok=True)


def _make_qr(path, fill="black", back="white", border=4, size=None):
    q = qrcode.QRCode(border=border)
    q.add_data("TEST")
    q.make()
    im = q.make_image(fill_color=fill, back_color=back).convert("RGB")
    if size:
        im = im.resize(size)
    im.save(path)
    return path


def _make_tpl(path, landscape=True, with_qr_img=False,
              with_qr_ph=False, page=(297, 210)):
    d = docx.Document()
    sec = d.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(page[0]), Mm(page[1])
    sec.left_margin = sec.right_margin = Mm(20)

    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("گواهینامه حضور").font.size = Pt(30)

    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("{نام} با کدملی {کدملی}").font.size = Pt(20)

    if with_qr_ph:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("{کیوآر}")

    if with_qr_img and HAVE_QR:
        s = str(TMP / "s.png")
        _make_qr(s)
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(s, width=Mm(25))

    d.save(str(path))
    return path


# ══════════════════════════════════════════════════════════
section("۱) توابع تازه موجودند")
# ══════════════════════════════════════════════════════════
check("trim_page_margin", hasattr(cg, "trim_page_margin"))
check("_append_orphan_images", hasattr(cg, "_append_orphan_images"))

_ST = getattr(cg, "CERT_SETTINGS", None) or getattr(cg, "SETTINGS", [])
keys = [r[0] for r in _ST]
check("تنظیم cert_trim_margin", "cert_trim_margin" in keys)
check("تنظیم cert_force_images", "cert_force_images" in keys)

dflt = {r[0]: r[4] for r in _ST} if keys else {}
check("برش پیش‌فرض روشن", dflt.get("cert_trim_margin") == "1")
check("qr در تصویرهای همیشه‌درج",
      "qr" in (dflt.get("cert_force_images") or ""))


# ══════════════════════════════════════════════════════════
section("۲) برش حاشیهٔ سفید")
# ══════════════════════════════════════════════════════════
try:
    from PIL import Image

    def _canvas(dw, dh, fw, fh, pos="tl"):
        """طرح رنگی در بوم سفید بزرگ‌تر"""
        design = Image.new("RGB", (dw, dh), (250, 240, 200))
        for x in range(0, dw, 40):
            for y in range(0, dh, 40):
                design.paste((30, 60, 120), (x, y, min(x + 20, dw),
                                             min(y + 20, dh)))
        cw, ch = int(dw * fw), int(dh * fh)
        c = Image.new("RGB", (cw, ch), "white")
        xy = {"tl": (0, 0), "mid": ((cw - dw) // 2, (ch - dh) // 2)}[pos]
        c.paste(design, xy)
        p = str(TMP / "c.png")
        c.save(p)
        return p, (dw, dh)

    # نوار پهن → باید بریده شود
    for name, fw, fh in [("راست ۵۰٪", 1.5, 1.0),
                         ("پایین ۴۰٪", 1.0, 1.4),
                         ("هر دو ۳۰٪", 1.3, 1.3)]:
        p, (dw, dh) = _canvas(600, 420, fw, fh)
        cg.trim_page_margin(p)
        a = Image.open(p)
        check(f"نوار {name} بریده شد",
              abs(a.width - dw) < 12 and abs(a.height - dh) < 12,
              f"{a.size} ≠ ({dw},{dh})")

    # بدون نوار → دست نخورد
    p, (dw, dh) = _canvas(600, 420, 1.0, 1.0)
    before = Image.open(p).size
    cg.trim_page_margin(p)
    check("بدون نوار دست‌نخورده", Image.open(p).size == before)

    # نوار باریک ۲٪ → دست نخورد (بخشی از طرح است)
    p, _ = _canvas(600, 420, 1.02, 1.02)
    before = Image.open(p).size
    cg.trim_page_margin(p)
    check("نوار باریک دست‌نخورده", Image.open(p).size == before,
          f"{before} → {Image.open(p).size}")

    # تصویر تماماً سفید → نباید بترکد
    Image.new("RGB", (400, 300), "white").save(str(TMP / "w.png"))
    r = cg.trim_page_margin(str(TMP / "w.png"))
    check("تصویر سفید خطا نمی‌دهد", r is not None)

    # تصویر خیلی کوچک
    Image.new("RGB", (20, 20), "white").save(str(TMP / "s2.png"))
    r = cg.trim_page_margin(str(TMP / "s2.png"))
    check("تصویر کوچک خطا نمی‌دهد", r is not None)

    # مسیر نامعتبر
    r = cg.trim_page_margin("/tmp/_nope_29.png")
    check("مسیر نامعتبر خطا نمی‌دهد", r is not None)
except Exception as e:
    check("تست برش", False, str(e))


# ══════════════════════════════════════════════════════════
section("۳) کیوآر همیشه درج می‌شود")
# ══════════════════════════════════════════════════════════
if not (HAVE_DOCX and HAVE_QR):
    print("   ⏭️ python-docx یا qrcode نصب نیست")
else:
    qr = str(TMP / "qr.png")
    _make_qr(qr)
    mp = {"name": "علی رضایی", "national_id": "۱۲۳"}

    def _imgs(path):
        dd = docx.Document(str(path))
        return sum(1 for r in dd.part.rels.values()
                   if "image" in r.reltype)

    # حالت الف: قالب دارای {کیوآر}
    t = _make_tpl(TMP / "t_ph.docx", with_qr_ph=True)
    got, err = cg.fill_docx(t, TMP / "o_ph.docx", mp, qr_png=qr)
    check("قالب با {کیوآر} پر شد", bool(got), err)
    if got:
        check("کیوآر درج شد (پارامتر)", _imgs(TMP / "o_ph.docx") >= 1)

    # حالت ب: قالب دارای تصویر نمونه
    t = _make_tpl(TMP / "t_img.docx", with_qr_img=True)
    got, err = cg.fill_docx(t, TMP / "o_img.docx", mp, qr_png=qr)
    check("قالب با تصویر نمونه پر شد", bool(got), err)
    if got:
        check("کیوآر درج شد (تصویر)", _imgs(TMP / "o_img.docx") >= 1)

    # 🔴 حالت ج: قالب بدون هیچ جای کیوآر — باگ اصلی
    t = _make_tpl(TMP / "t_none.docx")
    got, err = cg.fill_docx(t, TMP / "o_none.docx", mp, qr_png=qr)
    check("قالب بدون جای کیوآر پر شد", bool(got), err)
    if got:
        n = _imgs(TMP / "o_none.docx")
        check("🔳 کیوآر ته صفحه اضافه شد", n >= 1,
              f"{n} تصویر — باید ≥۱ باشد")

    # متن‌ها هم باید جایگزین شده باشند
    if got:
        dd = docx.Document(str(TMP / "o_none.docx"))
        txt = "\n".join(p.text for p in dd.paragraphs)
        check("نام جایگزین شد", "علی رضایی" in txt)
        check("پارامتر خام نماند", "{نام}" not in txt)


# ══════════════════════════════════════════════════════════
section("۴) ابعاد صفحه در تشخیص")
# ══════════════════════════════════════════════════════════
if not HAVE_DOCX:
    print("   ⏭️ python-docx نصب نیست")
else:
    src = str(cg.__file__)
    body = Path(src).read_text(encoding="utf-8")
    check("pw_mm در خروجی", '"pw_mm"' in body)
    check("orient در خروجی", '"orient"' in body)
    check("تشخیص A4 افقی", '"A4 افقی"' in body)
    check("هشدار اندازهٔ دلخواه", "اندازهٔ صفحه استاندارد نیست" in body)


# ══════════════════════════════════════════════════════════
section("۵) فرمول ابعاد PNG")
# ══════════════════════════════════════════════════════════
body = Path(cg.__file__).read_text(encoding="utf-8")

# 🔴 فرمول غلط قدیمی نباید باقی مانده باشد
check("فرمول غلط px*1.414 حذف شد",
      'int(px * 1.414)' not in body,
      "PixelHeight نباید همیشه ×۱.۴۱۴ باشد")
check("نسبت از سند خوانده می‌شود", "ratio = " in body)
check("page_height استفاده می‌شود", "_sec.page_height" in body)

# شبیه‌سازی محاسبه
for pw_mm, ph_mm, label in [(297, 210, "A4 افقی"),
                            (210, 297, "A4 عمودی"),
                            (210, 148, "A5 افقی")]:
    pw_emu = pw_mm * 36000
    ph_emu = ph_mm * 36000
    ratio = ph_emu / pw_emu
    inches = pw_emu / 914400.0
    px = max(600, min(4000, int(round(inches * 150))))
    py = max(600, min(4000, int(round(px * ratio))))
    want = pw_mm / ph_mm
    got_r = px / py
    check(f"نسبت {label}", abs(got_r - want) < 0.02,
          f"{got_r:.3f} ≠ {want:.3f}")


# ══════════════════════════════════════════════════════════
section("۶) تشخیص کیوآر — حالت‌های واقعی")
# ══════════════════════════════════════════════════════════
if not HAVE_QR:
    print("   ⏭️ qrcode نصب نیست")
else:
    from PIL import Image, ImageDraw

    def _blob(im):
        b = io.BytesIO()
        im.save(b, "PNG")
        return b.getvalue()

    q = qrcode.QRCode(border=2)
    q.add_data("X")
    q.make()
    base = q.make_image(fill_color="#1B3A6B",
                        back_color="white").convert("RGB")

    cases = {}
    qb = qrcode.QRCode(border=8)
    qb.add_data("X")
    qb.make()
    cases["حاشیه پهن"] = _blob(qb.make_image().convert("RGB"))
    cases["رنگی آبی"] = _blob(base)

    lg = base.copy()
    w, h = lg.size
    dr = ImageDraw.Draw(lg)
    s2 = w // 5
    dr.rectangle([w // 2 - s2 // 2, h // 2 - s2 // 2,
                  w // 2 + s2 // 2, h // 2 + s2 // 2], fill="#C9A84C")
    cases["با لوگو وسط"] = _blob(lg)

    cv = Image.new("RGB", (900, 900), "white")
    cv.paste(base.resize((200, 200)), (350, 350))
    cases["کوچک وسط بوم"] = _blob(cv)

    for name, blob in cases.items():
        check(f"تشخیص «{name}»", cg.looks_like_qr(blob))

    # چیزهایی که نباید QR تشخیص داده شوند
    photo = Image.new("RGB", (400, 300), (120, 150, 190))
    check("عکس ساده QR نیست", not cg.looks_like_qr(_blob(photo)))

    line = Image.new("RGB", (600, 60), "white")
    check("نوار باریک QR نیست", not cg.looks_like_qr(_blob(line)))


# ══════════════════════════════════════════════════════════
section("۷) واژه‌نامه")
# ══════════════════════════════════════════════════════════
import re
for fn in ["certgen.py", "certissue.py", "ticket.py"]:
    bad = []
    for i, ln in enumerate(Path(fn).read_text(encoding="utf-8")
                           .splitlines(), 1):
        st = ln.lstrip()
        if st.startswith("#"):
            continue
        # نقل‌قول مستقیم کاربر (بین « ») دست نمی‌خورد
        if "«" in ln and "»" in ln:
            continue
        if re.search(r"گواهی(?!نامه)(?! می)(?!\w)", ln):
            bad.append(f"خط {i}")
    check(f"«گواهینامه» در {fn}", not bad, " · ".join(bad[:3]))


print(f"\n{'=' * 62}")
print(f"✅ موفق: {OK}    ❌ ناموفق: {FAIL}")
print("=" * 62)
if ERRORS:
    print("\n❌ خطاها:")
    for e in ERRORS[:30]:
        print(f"   • {e}")
    sys.exit(1)
print("\n🎉 همه تست‌های فاز ۲۹ پاس شدند!")
