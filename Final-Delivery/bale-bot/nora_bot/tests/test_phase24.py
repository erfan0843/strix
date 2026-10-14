"""
تست فاز ۲۴ — 🔳 کیوآر · 💳 ثبت‌نام نیمه‌تمام · 🧱 هزینه زیرساخت

درخواست کاربر (عیناً):
  «کیو آر کد رو به صورت تصویری میزارم تو ورد تشخیص نمیده، متنی هم
   میسازم نمیندازه تو گواهی. چرا کاربر وقتی پرداختش کنسل میکنه کلا
   ثبت نام واسش بسته میشه؟ کلا رویداد با هزینه زیر ۲۰ هزار رو موقع
   ثبت نام به کاربر کلا رایگان نشون بده، مبالغ زیر ۲۰ هزار تو
   رویدادها هزینه زیرساخت و خدمات معرفی میشه... منو مالی پرداخت هم
   جامع کن که بهم ریخته نباشه»
"""
import io
import os
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

TMP = tempfile.mkdtemp(prefix="nora24_")
os.environ.update({
    "BOT_TOKEN": "1095412158:T", "SUPER_ADMIN_ID": "100000001",
    "SAFIR_API_KEY": "k", "SAFIR_BOT_ID": "317041514",
    "PAYMENT_PROVIDER_TOKEN": "WALLET-TEST-1111111111111111",
    "DB_PATH": f"{TMP}/t.db", "BACKUP_DIR": f"{TMP}/b",
    "IMPORT_DIR": f"{TMP}/i", "FILES_DIR": f"{TMP}/f",
    "PROFILES_DIR": f"{TMP}/f/p", "EVENTS_DIR": f"{TMP}/f/e",
    "LOG_DIR": f"{TMP}/l", "BROADCAST_RATE": "999",
})

import bale_api  # noqa: E402

SENT = []


def _api(m, p=None, **k):
    SENT.append({"m": m, "p": p or {}})
    if m == "getMe":
        return {"ok": True, "result": {"username": "nora_bot"}}
    return {"ok": True, "result": {"message_id": len(SENT) + 100}}


def _send(cid, text, kb=None, reply_to=None):
    SENT.append({"m": "sendMessage",
                 "p": {"cid": cid, "text": text, "kb": kb}})
    return {"ok": True, "result": {"message_id": len(SENT) + 100}}


def _soe(cid, text, kb=None, mid=None, force_new=False):
    SENT.append({"m": "editMessageText",
                 "p": {"cid": cid, "text": text, "kb": kb}})
    return {"ok": True, "result": {"message_id": mid or 1}}


def _file(method, key, cid, f, caption=None, kb=None):
    SENT.append({"m": method, "p": {"cid": cid, "text": caption or "",
                                    "kb": kb}})
    return {"ok": True, "result": {"message_id": len(SENT) + 100}}


bale_api.api = _api
bale_api.api_async = _api
bale_api.send = _send
bale_api.send_or_edit = _soe
bale_api._send_file = _file
bale_api.send_photo = lambda c, p, cap=None, kb=None: _file(
    "sendPhoto", "photo", c, p, cap, kb)
bale_api.send_document = lambda c, d, cap=None, kb=None: _file(
    "sendDocument", "document", c, d, cap, kb)
bale_api.send_invoice = lambda c, **kw: (
    SENT.append({"m": "sendInvoice",
                 "p": {"cid": c, "text": str(kw.get("prices"))}}),
    {"ok": True, "result": {"message_id": 1}})[1]
bale_api.answer_callback = lambda *a, **k: {"ok": True}

import config  # noqa: E402
import database  # noqa: E402
import auth  # noqa: E402
import certgen as CG  # noqa: E402
import events as EV  # noqa: E402
import payments as P  # noqa: E402
import permissions as PM  # noqa: E402
import infra_fee as INF  # noqa: E402
import forms as F  # noqa: E402
import form_fill as FF  # noqa: E402
import main as M  # noqa: E402

for mod in (CG, EV, P, INF, F, FF, M):
    for n in ("send", "send_or_edit", "send_photo", "send_document"):
        if hasattr(mod, n):
            setattr(mod, n, getattr(bale_api, n))

OK = FAIL = 0
FAILED = []


def check(name, cond, extra=""):
    global OK, FAIL
    if cond:
        OK += 1
        print(f"  ✅ {name}")
    else:
        FAIL += 1
        FAILED.append(name)
        print(f"  ❌ {name}   {extra}")


def texts():
    return [s["p"].get("text") or "" for s in SENT]


def last():
    t = [x for x in texts() if x]
    return t[-1] if t else ""


def alltext():
    return "\n".join(texts())


def allcb():
    out = []
    for s in SENT:
        kb = s["p"].get("kb") or {}
        for r in kb.get("inline_keyboard", []):
            for b in r:
                if isinstance(b, dict) and b.get("callback_data"):
                    out.append(b["callback_data"])
    return out


def msg(uid, t):
    return {"message_id": 1, "from": {"id": uid, "first_name": "ت"},
            "chat": {"id": uid, "type": "private"}, "text": t}


def cb(uid, d):
    return {"id": "c", "from": {"id": uid, "first_name": "ت"}, "data": d,
            "message": {"message_id": 1,
                        "chat": {"id": uid, "type": "private"}}}


print("=" * 62)
print("🧪 تست فاز ۲۴ — کیوآر · ثبت‌نام نیمه‌تمام · هزینه زیرساخت")
print("=" * 62)

database.init_db()
EV.init_event_tables()
P.init_payment_tables()
PM.init_permission_tables()
CG.init_cert_tables()
F.init_form_tables()
database.seed_default_settings()

S = 100000001
U1, U2 = 2410001, 2410002
auth.create_auth_user({"id": S, "first_name": "مدیر"}, "09120000000",
                      "phone")
auth.ensure_super_admin(S)
for i, u in enumerate((U1, U2)):
    auth.create_auth_user({"id": u, "first_name": f"کاربر{i+1}"},
                          f"0912444000{i}", "phone")
    database.db_execute(
        "UPDATE users SET is_authenticated=1, is_active=1, "
        "complete_profile_percent=100, profile_status='approved' "
        "WHERE bale_user_id=?", (u,))


# ═══════════ ۱) 🔳 تشخیص تصویر کیوآر ═══════════
print("\n[۱] 🔳 تشخیص تصویری کیوآر")

_qr_ok = True
try:
    import qrcode
    from PIL import Image
except ImportError:
    _qr_ok = False

if _qr_ok:
    def _blob(im, fmt="PNG", **kw):
        b = io.BytesIO()
        im.save(b, fmt, **kw)
        return b.getvalue()

    base = qrcode.make("https://ble.ir/x?start=cert_A").convert("RGB")

    check("🔳 PNG خام", CG.looks_like_qr(_blob(base)))
    check("🔳 JPEG کیفیت پایین",
          CG.looks_like_qr(_blob(base, "JPEG", quality=60)))
    check("🔳 کوچک‌شده (اسکرین‌شات)",
          CG.looks_like_qr(_blob(base.resize((150, 150), Image.LANCZOS))))

    # 🔴 حاشیه سفید — قبلاً رد می‌شد
    pad = Image.new("RGB", (base.width + 90, base.height + 40), "white")
    pad.paste(base, (45, 20))
    check("🔴 با حاشیه سفید (قبلاً رد می‌شد)",
          CG.looks_like_qr(_blob(pad)))

    # 🔴 QR رنگی — قبلاً رد می‌شد
    col = base.copy()
    px = col.load()
    for y in range(col.height):
        for x in range(col.width):
            if px[x, y][0] < 128:
                px[x, y] = (10, 40, 160)
    check("🔴 QR آبی رنگی (قبلاً رد می‌شد)",
          CG.looks_like_qr(_blob(col)))

    # 🔴 کوچک وسط بوم بزرگ — قبلاً رد می‌شد
    big = Image.new("RGB", (900, 900), "white")
    big.paste(base.resize((380, 380)), (260, 260))
    check("🔴 کوچک وسط بوم بزرگ (قبلاً رد می‌شد)",
          CG.looks_like_qr(_blob(big)))

    # پس‌زمینه کرم
    cream = Image.new("RGB", base.size, (250, 245, 230))
    bpx, cpx = base.load(), cream.load()
    for y in range(base.height):
        for x in range(base.width):
            if bpx[x, y][0] < 128:
                cpx[x, y] = (0, 0, 0)
    check("🔳 پس‌زمینه کرم", CG.looks_like_qr(_blob(cream)))

    print("\n  --- نباید اشتباه تشخیص داده شوند ---")
    logo = Image.new("RGB", (200, 200), (200, 30, 30))
    check("🛡️ لوگوی تک‌رنگ رد می‌شود",
          not CG.looks_like_qr(_blob(logo)))
    sign = Image.new("RGB", (300, 110), "white")
    check("🛡️ امضای سفید رد می‌شود",
          not CG.looks_like_qr(_blob(sign)))
    ph = Image.new("RGB", (400, 400))
    p2 = ph.load()
    for y in range(400):
        for x in range(400):
            p2[x, y] = ((x * 7) % 256, (y * 5) % 256, ((x + y) * 3) % 256)
    check("🛡️ عکس رنگی رد می‌شود", not CG.looks_like_qr(_blob(ph)))

    check("تابع بریدن حاشیه هست", hasattr(CG, "_trim_border"))
    _f = CG._image_fingerprint(_blob(base))
    for k in ("two_tone", "ink", "edges"):
        check(f"اثر انگشت «{k}» دارد", k in _f)
else:
    check("⚠️ qrcode/PIL نیست — رد شد", True)


# ═══════════ ۲) 🔤 نام‌های متنی کیوآر ═══════════
print("\n[۲] 🔤 هرجور بنویسد، بشناسد")

for name in ("کیوآر", "کیو آر", "کیوآرکد", "کیو آر کد", "کد QR",
             "qr", "QR", "qrcode", "qr_code", "qr code",
             "کیوارکد", "تصویر کیوآر", "کد دوبعدی"):
    check(f"🔤 «{{{name}}}» → qr", CG.resolve_key(name) == "qr",
          str(CG.resolve_key(name)))

for name in ("بارکد", "بار کد", "barcode"):
    check(f"🔤 «{{{name}}}» → barcode",
          CG.resolve_key(name) == "barcode")

check("🛡️ نام بی‌ربط None می‌دهد",
      CG.resolve_key("نام خانوادگی مادر") is None)
check("🛡️ خالی کرش نمی‌کند", CG.resolve_key("") is None)
check("🛡️ None کرش نمی‌کند", CG.resolve_key(None) is None)


# ═══════════ ۳) 🎓 درج واقعی در گواهینامه ═══════════
print("\n[۳] 🎓 درج کیوآر در فایل ورد")

_docx_ok = True
try:
    import docx
    from docx.shared import Mm
except ImportError:
    _docx_ok = False

if _docx_ok and _qr_ok:
    qp = Path(TMP) / "q.png"
    CG.make_qr("https://ble.ir/x?start=cert_T1", qp)

    # الف) قالب با متن {کیوآرکد}
    t1 = Path(TMP) / "t1.docx"
    d = docx.Document()
    d.add_paragraph("گواهی می‌شود {نام} شرکت نمود.")
    d.add_paragraph("اعتبارسنجی: {کیوآرکد}")
    d.save(str(t1))

    o1 = Path(TMP) / "o1.docx"
    got, err = CG.fill_docx(str(t1), o1, {"name": "علی رضایی"},
                            qr_png=str(qp))
    check("🎓 قالب متنی پر شد", bool(got), err)
    if got:
        dd = docx.Document(got)
        body = "\n".join(p.text for p in dd.paragraphs)
        n_img = sum(1 for r in dd.part.rels.values()
                    if "image" in r.reltype)
        check("🔴 تصویر کیوآر درج شد (باگ رفع شد)", n_img >= 1,
              str(n_img))
        check("🔴 آکولاد {کیوآرکد} پاک شد", "کیوآرکد" not in body,
              body[:80])
        check("نام هم جایگزین شد", "علی رضایی" in body)

    # ب) قالب با تصویر نمونه رنگی
    col2 = qrcode.make("نمونه").convert("RGB")
    cp = col2.load()
    for y in range(col2.height):
        for x in range(col2.width):
            if cp[x, y][0] < 128:
                cp[x, y] = (20, 90, 40)
    sample = Path(TMP) / "sample_colored.png"
    col2.save(str(sample))

    t2 = Path(TMP) / "t2.docx"
    d = docx.Document()
    d.add_paragraph("گواهی می‌شود {نام} شرکت نمود.")
    d.add_paragraph().add_run().add_picture(str(sample), width=Mm(30))
    d.save(str(t2))

    d2 = docx.Document(str(t2))
    check("🔴 تصویر نمونه رنگی شناسایی شد",
          CG.find_qr_placeholder(d2) is not None)

    o2 = Path(TMP) / "o2.docx"
    got2, err2 = CG.fill_docx(str(t2), o2, {"name": "سارا"},
                              qr_png=str(qp))
    check("🎓 قالب تصویری پر شد", bool(got2), err2)
    if got2:
        dd = docx.Document(got2)
        newb = qp.read_bytes()
        swapped = any(r.target_part.blob == newb
                      for r in dd.part.rels.values()
                      if "image" in r.reltype)
        check("🔴 تصویر نمونه با کیوآر واقعی عوض شد", swapped)
else:
    check("⚠️ python-docx نیست — رد شد", True)


# ═══════════ ۴) 💳 ثبت‌نام نیمه‌تمام قفل نشود ═══════════
print("\n[۴] 💳 کنسل پرداخت، ثبت‌نام را نبندد")

EID = database.db_execute(
    "INSERT INTO events (title,event_type,status,capacity,created_by,"
    "is_paid,price,payment_method,start_date,require_profile) "
    "VALUES (?,?,?,?,?,?,?,?,?,?)",
    ("کارگاه پولی", "workshop", "published", 50, S, 1, 500000,
     "bale_wallet", "1404/07/01", 0))

SENT.clear()
M.handle_callback(cb(U1, f"ureg_{EID}"))
_r = database.db_execute(
    "SELECT * FROM event_registrations WHERE event_id=? AND bale_user_id=?",
    (EID, U1), fetch="one")
check("ثبت‌نام اول ساخته شد", bool(_r))
check("وضعیت در انتظار پرداخت", _r and _r["status"] == "pending")

auth.clear_state(U1)
SENT.clear()
M.handle_callback(cb(U1, f"ureg_{EID}"))
_t = last()
check("🔴 قفل نمی‌شود (باگ اصلی رفع شد)",
      "قبلاً ثبت‌نام کرده‌اید" not in _t, _t[:70])
check("پیام مهربان «نیمه‌تمام» می‌دهد", "نیمه‌تمام" in _t, _t[:70])
check("دکمه ادامه پرداخت دارد", f"urepay_{EID}" in allcb())
check("دکمه شروع دوباره دارد", f"uregnew_{EID}" in allcb())
check("دکمه انصراف دارد", f"ucancel_{EID}" in allcb())

SENT.clear()
M.handle_callback(cb(U1, f"urepay_{EID}"))
check("💳 ادامه پرداخت، فاکتور می‌فرستد",
      any(s["m"] == "sendInvoice" for s in SENT))

SENT.clear()
M.handle_callback(cb(U1, f"uregnew_{EID}"))
_r2 = database.db_execute(
    "SELECT * FROM event_registrations WHERE event_id=? AND bale_user_id=?",
    (EID, U1), fetch="one")
check("🔄 شروع دوباره، ثبت‌نام تازه می‌سازد", bool(_r2))
check("رکورد قبلی پاک شد (id عوض شد)",
      not _r2 or _r2["id"] != _r["id"], f"{_r['id']} → {_r2['id']}")

# 🛡️ اگر پولی پرداخت شده، پاک نشود
database.db_execute("UPDATE event_registrations SET amount_paid=200000 "
                    "WHERE id=?", (_r2["id"],))
SENT.clear()
M.handle_callback(cb(U1, f"uregnew_{EID}"))
check("🛡️ با پرداخت جزئی، ثبت‌نام پاک نمی‌شود",
      bool(database.db_execute(
          "SELECT 1 FROM event_registrations WHERE id=?",
          (_r2["id"],), fetch="one")))
check("پیام راهنما می‌دهد", "پشتیبانی" in last(), last()[:60])

# تأییدشده نباید این مسیر را ببیند
database.db_execute("UPDATE event_registrations SET status='confirmed', "
                    "amount_paid=500000 WHERE id=?", (_r2["id"],))
SENT.clear()
M.handle_callback(cb(U1, f"ureg_{EID}"))
check("✅ ثبت‌نام تأییدشده پیام درست می‌گیرد",
      "قبلاً ثبت‌نام" in last(), last()[:60])


# ═══════════ ۵) 🧱 هزینه زیرساخت ═══════════
print("\n[۵] 🧱 هزینه زیرساخت و خدمات")

check("پیش‌فرض روشن است", INF.enabled())
check("سقف پیش‌فرض ۲۰٬۰۰۰", INF.threshold() == 20000,
      str(INF.threshold()))
check("عنوان پیش‌فرض درست", "زیرساخت" in INF.title())

for a, want in ((5000, True), (10000, True), (20000, True),
                (20001, False), (500000, False), (0, False)):
    check(f"{a:,} → {'زیرساخت' if want else 'عادی'}",
          INF.is_infra(a) is want)

_lbl, _is = INF.price_label(10000)
check("برچسب «رایگان» می‌سازد", "رایگان" in _lbl, _lbl)
check("عدد را هم می‌گوید", "10" in _lbl)
check("عنوان زیرساخت در برچسب", "زیرساخت" in _lbl)
check("سطر فاکتور عنوان زیرساخت می‌دهد",
      "زیرساخت" in INF.invoice_line(10000))
check("مبلغ بالا سطر عادی می‌دهد", INF.invoice_line(500000) == "")

# نمای کاربر
E_LOW = database.db_execute(
    "INSERT INTO events (title,event_type,status,capacity,created_by,"
    "is_paid,price,payment_method,start_date,require_profile,"
    "show_price_publicly) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
    ("کارگاه کم‌هزینه", "workshop", "published", 50, S, 1, 10000,
     "bale_wallet", "1404/07/01", 0, 1))
SENT.clear()
M.handle_callback(cb(U2, f"uev_{E_LOW}"))
_t = last()
check("🔴 کاربر «رایگان» می‌بیند", "رایگان" in _t, _t[:90])
check("توضیح زیرساخت هم می‌آید", "زیرساخت" in _t)
check("دکمه ثبت‌نام «رایگان» است",
      any("رایگان" in b.get("text", "")
          for s in SENT
          for r in (s["p"].get("kb") or {}).get("inline_keyboard", [])
          for b in r))

E_HI = database.db_execute(
    "INSERT INTO events (title,event_type,status,capacity,created_by,"
    "is_paid,price,payment_method,start_date,require_profile,"
    "show_price_publicly) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
    ("دوره گران", "workshop", "published", 50, S, 1, 500000,
     "bale_wallet", "1404/07/01", 0, 1))
SENT.clear()
M.handle_callback(cb(U2, f"uev_{E_HI}"))
_t = last()
check("💳 مبلغ بالا عادی می‌ماند",
      "500,000" in _t or "500٬000" in _t, _t[:90])
check("«رایگان» نمی‌گوید", "🎁 رایگان" not in _t)

# مبلغ واقعی دست‌نخورده
SENT.clear()
M.handle_callback(cb(U2, f"ureg_{E_LOW}"))
_rl = database.db_execute(
    "SELECT * FROM event_registrations WHERE event_id=? AND bale_user_id=?",
    (E_LOW, U2), fetch="one")
_tx = database.db_execute(
    "SELECT * FROM transactions WHERE registration_id=?",
    (_rl["id"],), fetch="one") if _rl else None
check("🔴 مبلغ دیتابیس دست‌نخورده", not _tx or _tx["amount"] == 10000,
      str(_tx["amount"] if _tx else "—"))
check("🧾 عنوان فاکتور «زیرساخت» شد",
      any("زیرساخت" in (s["p"].get("text") or "")
          for s in SENT if s["m"] == "sendInvoice"),
      str([s["p"].get("text") for s in SENT
           if s["m"] == "sendInvoice"])[:90])


# ═══════════ ۶) 🖥️ پنل زیرساخت ═══════════
print("\n[۶] 🖥️ پنل ادمین زیرساخت")

SENT.clear()
M.handle_callback(cb(S, "pay_infra"))
check("پنل باز شد", "زیرساخت" in last())
check("پیش‌نمایش نشان می‌دهد", "کاربر این را می‌بیند" in last())
for key, _i, _n, _k, _d, _h in INF.INFRA_SETTINGS:
    pre = "infb_" if _k == "bool" else "infe_"
    check(f"تنظیم «{_n}» دکمه دارد", f"{pre}{key}" in allcb())

SENT.clear()
M.handle_callback(cb(S, "inf_prev"))
check("👁️ پیش‌نمایش باز شد", "پیش‌نمایش" in last())

M.handle_callback(cb(S, "infe_infra_max"))
M.process_message(msg(S, "۵۰۰۰۰"))
check("🔢 سقف با عدد فارسی عوض شد", INF.threshold() == 50000,
      str(INF.threshold()))
check("حالا ۳۰٬۰۰۰ هم زیرساخت است", INF.is_infra(30000))

M.handle_callback(cb(S, "infe_infra_title"))
M.process_message(msg(S, "هزینه خدمات"))
check("🏷️ عنوان عوض شد", INF.title() == "هزینه خدمات", INF.title())
M.handle_callback(cb(S, "infe_infra_title"))
M.process_message(msg(S, "-"))
check("♻️ با «-» به پیش‌فرض برمی‌گردد",
      "زیرساخت" in INF.title())

M.handle_callback(cb(S, "infb_infra_show_amount"))
_l2, _ = INF.price_label(10000)
check("🔢 خاموش کردن نمایش عدد", "10" not in _l2, _l2)
M.handle_callback(cb(S, "infb_infra_show_amount"))

M.handle_callback(cb(S, "infb_infra_on"))
check("🔴 خاموش شد", not INF.enabled())
SENT.clear()
M.handle_callback(cb(U2, f"uev_{E_LOW}"))
check("خاموش که شد، عدد برمی‌گردد",
      "10,000" in last() or "10٬000" in last(), last()[:80])
M.handle_callback(cb(S, "infb_infra_on"))
M.handle_callback(cb(S, "infe_infra_max"))
M.process_message(msg(S, "20000"))
check("♻️ به حالت اول برگشت",
      INF.enabled() and INF.threshold() == 20000)


# ═══════════ ۷) 💳 منوی مالی مرتب ═══════════
print("\n[۷] 💳 منوی پرداخت و مالی")

SENT.clear()
M.handle_callback(cb(S, "set_payment"))
_kb = None
for s in reversed(SENT):
    if (s["p"].get("kb") or {}).get("inline_keyboard"):
        _kb = s["p"]["kb"]["inline_keyboard"]
        break
_labels = [b.get("text", "") for r in (_kb or []) for b in r]
_txt = " ".join(_labels)
check("منو باز شد", bool(_kb))
check("🗂️ گروه گزارش مالی دارد", "گزارش مالی" in _txt)
check("🗂️ گروه راه‌های دریافت دارد", "راه‌های دریافت" in _txt)
check("🗂️ گروه تنظیمات دارد", "⚙️ تنظیمات" in _txt)
check("🧱 هزینه زیرساخت در منو هست", "pay_infra" in allcb())
check("داشبورد درآمد هست", "fin_dash" in allcb())
check("تراکنش‌ها هست", "fin_tx_0_all" in allcb())
check("کارت‌ها هست", "pay_cards" in allcb())
check("جداکننده‌ها noop هستند",
      allcb().count("noop") >= 3, str(allcb().count("noop")))


# ═══════════ ۸) 🔒 امنیت ═══════════
print("\n[۸] 🔒 دسترسی")

for d in ("pay_infra", "inf_prev", "infb_infra_on",
          "infe_infra_max"):
    SENT.clear()
    M.handle_callback(cb(U1, d))
    check(f"🔒 کاربر عادی: {d}",
          "سوپرادمین" in last() or "دسترسی" in last(), last()[:40])


# ═══════════ ۹) 🛡️ ورودی خراب ═══════════
print("\n[۹] 🛡️ ورودی‌های خراب")

for d in ("urepay_abc", "uregnew_99999", "urepay_99999",
          "infb_nokey", "infe_nokey", "infb_infra_max"):
    SENT.clear()
    try:
        M.handle_callback(cb(S, d))
        check(f"🛡️ «{d}» کرش نمی‌کند", True)
    except Exception as e:
        check(f"🛡️ «{d}» کرش نمی‌کند", False, str(e)[:70])

check("is_infra با None", INF.is_infra(None) is False)
check("is_infra با متن", INF.is_infra("abc") is False)
check("price_label با صفر", INF.price_label(0)[1] is False)
check("looks_like_qr با بایت خراب",
      CG.looks_like_qr(b"notanimage") is False)
check("_image_fingerprint با خراب",
      CG._image_fingerprint(b"xx") is None)


# ═══════════ ۹.۵) 🩺 عیب‌یابی قالب (فاز ۲۵) ═══════════
print("\n[۹.۵] 🩺 کالبدشکافی قالب — «نمیتونم فایل ورد بفرستم»")

check("تابع diagnose_template هست", hasattr(CG, "diagnose_template"))
check("تابع show_diagnose هست", hasattr(CG, "show_diagnose"))
check("تابع test_build هست", hasattr(CG, "test_build"))
check("تابع pdf_to_png_lo هست", hasattr(CG, "pdf_to_png_lo"))

_src = (ROOT / "certgen.py").read_text(encoding="utf-8")
check("🔴 PDF بدون تگ صادر می‌شود (رفع MuPDF)",
      "UseTaggedPDF" in _src)
check("🔴 فیلتر درست Writer (نه Impress)",
      "writer_png_Export" in _src)
check("🔇 خطاهای داخلی MuPDF خاموش می‌شوند",
      "mupdf_display_errors" in _src)
check("🔴 خطای PyMuPDF دیگر بی‌صدا نیست",
      "⚠️ PyMuPDF" in _src)
check("گزارش نهایی شکست تبدیل هست",
      "هیچ موتوری نتوانست" in _src)

if _docx_ok and _qr_ok:
    _t3 = Path(TMP) / "diag.docx"
    _d3 = docx.Document()
    _d3.add_paragraph("گواهی می‌شود {نام} با کدملی {کدملی}")
    _d3.add_paragraph("شماره: {شماره سریال}")     # نام اشتباه
    _qr3 = qrcode.make("x").convert("RGB")
    _p3 = _qr3.load()
    for _y in range(_qr3.height):
        for _x in range(_qr3.width):
            if _p3[_x, _y][0] < 128:
                _p3[_x, _y] = (0, 70, 140)
    _qp3 = Path(TMP) / "cq.png"
    _qr3.save(str(_qp3))
    _d3.add_paragraph().add_run().add_picture(str(_qp3), width=Mm(28))
    _d3.save(str(_t3))

    _tid3 = database.db_execute(
        "INSERT INTO cert_templates (name,kind,file_path,placeholders,"
        "is_active,is_default,created_by) VALUES (?,?,?,?,?,?,?)",
        ("قالب تست", "docx", str(_t3), "[]", 1, 0, S))

    _dg = CG.diagnose_template(_tid3)
    check("🩺 گزارش ساخته شد", _dg["ok"], _dg.get("err"))
    check("🔴 تصویر کیوآر رنگی شناسایی شد", bool(_dg["qr_img"]),
          str(_dg["images"])[:100])
    check("پارامترهای درست شمرده شدند",
          "name" in _dg["phs"] and "national_id" in _dg["phs"])
    check("🔴 پارامتر اشتباه گزارش شد",
          any("شماره سریال" in u for u in _dg["unknown"]),
          str(_dg["unknown"]))
    check("تعداد تصاویر درست", len(_dg["images"]) >= 1)

    SENT.clear()
    M.handle_callback(cb(S, f"ctdiag_{_tid3}"))
    _t = last()
    check("🖥️ صفحه بررسی باز شد", "بررسی قالب" in _t, _t[:60])
    check("وضعیت کیوآر را می‌گوید", "کیوآرکد" in _t)
    check("پارامتر ناشناس را هشدار می‌دهد", "شناخته نشد" in _t)
    check("دکمه تست ساخت دارد", f"cttest_{_tid3}" in allcb())

    SENT.clear()
    M.handle_callback(cb(S, f"cttest_{_tid3}"))
    _t = last()
    check("🧪 تست ساخت اجرا شد", "نتیجه تست" in _t, _t[:60])
    check("سه مرحله گزارش می‌شود",
          "ورد" in _t and "PDF" in _t and "تصویر" in _t)
    check("موتور را می‌گوید", "موتور" in _t)

    SENT.clear()
    M.handle_callback(cb(S, f"ctv_{_tid3}"))
    check("دکمه بررسی در کارت قالب", f"ctdiag_{_tid3}" in allcb())

check("🛡️ diagnose با شناسه بد کرش نمی‌کند",
      CG.diagnose_template(99999)["ok"] is False)
SENT.clear()
M.handle_callback(cb(S, "ctdiag_abc"))
check("🛡️ ctdiag_abc کرش نمی‌کند", True)
M.handle_callback(cb(S, "cttest_99999"))
check("🛡️ cttest_99999 کرش نمی‌کند", True)


# ═══════════ ۱۰) 🛡️ رگرسیون ═══════════
print("\n[۱۰] 🛡️ رگرسیون")

SENT.clear()
M.handle_callback(cb(U2, f"ureg_{E_HI}"))
check("♻️ ثبت‌نام رویداد گران سالم",
      any(s["m"] == "sendInvoice" for s in SENT))

SENT.clear()
M.handle_callback(cb(S, "fin_dash"))
check("♻️ داشبورد درآمد سالم", "درآمد" in last() or "مالی" in last())

SENT.clear()
M.handle_callback(cb(S, "ct_hub"))
check("♻️ هاب گواهینامه سالم", "گواهینامه" in last())

SENT.clear()
M.handle_callback(cb(S, f"ev_view_{EID}"))
check("♻️ کارت رویداد سالم", "کارگاه پولی" in last())

check("♻️ ph_label کیوآر درست", CG.ph_label("qr") == "کیوآر")
check("♻️ find_placeholders کار می‌کند",
      CG.find_placeholders("{نام} و {کیوآر}") == {"name", "qr"})

import guard as G  # noqa: E402
_f = G.check_new_features()
check("🛡️ همه قابلیت‌ها سالم", all(_f.values()),
      str([k for k, v in _f.items() if not v]))


print("\n" + "=" * 62)
print(f"✅ موفق: {OK}    ❌ ناموفق: {FAIL}")
print("=" * 62)
if FAILED:
    print("\nناموفق:")
    for f in FAILED:
        print(f"  ❌ {f}")
    sys.exit(1)
print("\n🎉 همه تست‌های فاز ۲۴ پاس شدند!")
