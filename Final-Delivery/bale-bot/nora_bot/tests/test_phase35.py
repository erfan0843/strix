# -*- coding: utf-8 -*-
"""
🧪 تست فاز ۳۵ — کیوآر سر جای خودش، در صفحهٔ اول

🎯 گزارش کاربر:
   «کیوارکد صادر میکنه ولی نمیزاره سرجاش برده ص دوم
    چرا نمیتونه بزاره جایی که تعیین کردم بدون بهم ریختن فرم؟»

🔴 سه باگ ریشه‌ای که این تست‌ها می‌پوشانند:

   ۱) `_add_floating_image` با `d.add_paragraph()` یک پاراگرافِ
      تازه **ته سند** می‌ساخت. قالب گواهینامه دقیقاً یک صفحه
      است، پس آن پاراگراف سرریز می‌کرد و صفحهٔ دوم باز می‌شد.
      لنگر تصویر هم با آن به صفحهٔ ۲ می‌رفت.
      ➜ اثبات در PDF کاربر:  ص۱ = Im4/Im7/Im8 · ص۲ = Im20 تنها

   ۲) `_place_image` فقط `d.paragraphs` را می‌گشت و
      تکست‌باکس‌ها را نمی‌دید. قالب‌های حرفه‌ای همه‌چیز را
      داخل تکست‌باکس می‌گذارند، پس {کیوآر} هرگز پیدا نمی‌شد.

   ۳) مختصات کادر **بعد از** پاک‌شدن متن‌ها خوانده می‌شد؛
      آن موقع دیگر «{کیوآر}» وجود نداشت و بی‌صدا به گوشهٔ
      پیش‌فرض برمی‌گشت.
"""
import os
import re
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

os.environ.setdefault("BOT_TOKEN", "0:test")
os.environ.setdefault("SUPER_ADMIN_ID", "1")

_ok = _bad = 0
_fails = []


def check(name, cond):
    global _ok, _bad
    if cond:
        _ok += 1
    else:
        _bad += 1
        _fails.append(name)
        print(f"  ❌ {name}")


# ══════════════════════════════════════════════════════════
# 🏗️ ساخت قالب‌های آزمایشی
# ══════════════════════════════════════════════════════════
import tempfile

TMP = tempfile.mkdtemp(prefix="p35_")
MM = 36000                       # EMU در هر میلی‌متر

W_MM, H_MM = 297.0, 210.0        # A4 افقی
BOX_X, BOX_Y, BOX_WH = 200.0, 150.0, 30.0     # کادر کیوآر


def _new_doc():
    import docx
    from docx.shared import Mm
    from docx.enum.section import WD_ORIENT
    d = docx.Document()
    s = d.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Mm(W_MM), Mm(H_MM)
    s.left_margin = s.right_margin = Mm(0)
    s.top_margin = s.bottom_margin = Mm(0)
    return d


def _bg(path):
    from PIL import Image
    Image.new("RGB", (1492, 1054), (240, 230, 200)).save(path)
    return path


def build_plain():
    """قالب بدون {کیوآر} — تصویر یتیم می‌شود"""
    from docx.shared import Mm
    d = _new_doc()
    d.add_paragraph().add_run().add_picture(
        _bg(f"{TMP}/bg.png"), width=Mm(W_MM))
    d.add_paragraph("گواهی می‌شود {نام}")
    p = f"{TMP}/plain.docx"
    d.save(p)
    return p


def build_dml():
    """تکست‌باکس مدرن (DrawingML) با {کیوآر} داخلش"""
    from docx.shared import Mm
    from lxml import etree
    d = _new_doc()
    d.add_paragraph().add_run().add_picture(
        _bg(f"{TMP}/bg.png"), width=Mm(W_MM))
    d.add_paragraph("گواهی می‌شود {نام}")
    nw = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    nwp = ("http://schemas.openxmlformats.org/drawingml/2006/"
           "wordprocessingDrawing")
    na = "http://schemas.openxmlformats.org/drawingml/2006/main"
    nwps = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
    xml = (
        f'<w:r xmlns:w="{nw}"><w:drawing xmlns:wp="{nwp}">'
        f'<wp:anchor distT="0" distB="0" distL="0" distR="0" '
        f'simplePos="0" relativeHeight="251" behindDoc="0" locked="0" '
        f'layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page">'
        f'<wp:posOffset>{int(BOX_X * MM)}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page">'
        f'<wp:posOffset>{int(BOX_Y * MM)}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{int(BOX_WH * MM)}" cy="{int(BOX_WH * MM)}"/>'
        f'<wp:docPr id="99" name="QRBox"/>'
        f'<a:graphic xmlns:a="{na}"><a:graphicData uri="{nwps}">'
        f'<wps:wsp xmlns:wps="{nwps}"><wps:txbx><w:txbxContent><w:p>'
        f'<w:r><w:t>{{کیوآر}}</w:t></w:r>'
        f'</w:p></w:txbxContent></wps:txbx>'
        f'</wps:wsp></a:graphicData></a:graphic>'
        f'</wp:anchor></w:drawing></w:r>')
    d.add_paragraph()._element.append(etree.fromstring(xml))
    p = f"{TMP}/dml.docx"
    d.save(p)
    return p


def build_vml():
    """تکست‌باکس قدیمی (VML) — ورد ۲۰۰۳ و مبدل‌های آنلاین"""
    from docx.shared import Mm
    from lxml import etree
    d = _new_doc()
    d.add_paragraph().add_run().add_picture(
        _bg(f"{TMP}/bg.png"), width=Mm(W_MM))
    pt = 72.0 / 25.4                      # میلی‌متر → پوینت
    xml = (
        '<w:pict xmlns:w="http://schemas.openxmlformats.org/'
        'wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml">'
        f'<v:shape style="position:absolute;left:{BOX_X * pt:.1f}pt;'
        f'top:{BOX_Y * pt:.1f}pt;width:{BOX_WH * pt:.1f}pt;'
        f'height:{BOX_WH * pt:.1f}pt">'
        '<v:textbox><w:txbxContent><w:p><w:r>'
        '<w:t>{کیوآر}</w:t>'
        '</w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>')
    d.add_paragraph().add_run()._element.append(etree.fromstring(xml))
    p = f"{TMP}/vml.docx"
    d.save(p)
    return p


def build_inline():
    """{کیوآر} در پاراگراف معمولی — مسیر قدیمی، نباید بشکند"""
    from docx.shared import Mm
    d = _new_doc()
    d.add_paragraph().add_run().add_picture(
        _bg(f"{TMP}/bg.png"), width=Mm(W_MM))
    d.add_paragraph("گواهی می‌شود {نام}")
    d.add_paragraph("{کیوآر}")
    p = f"{TMP}/inline.docx"
    d.save(p)
    return p


def build_trailing_blanks():
    """قالبی با ۵ خط خالی ته سند — عامل کلاسیک صفحهٔ دوم"""
    from docx.shared import Mm
    d = _new_doc()
    d.add_paragraph().add_run().add_picture(
        _bg(f"{TMP}/bg.png"), width=Mm(W_MM))
    d.add_paragraph("گواهی می‌شود {نام}")
    for _ in range(5):
        d.add_paragraph("")
    p = f"{TMP}/blanks.docx"
    d.save(p)
    return p


def make_qr():
    """
    🔳 تصویر آزمایشی کیوآر.

    ⚠️ اگر کتابخانهٔ qrcode نصب نباشد، یک PNG ساده می‌سازیم.
       این تست دربارهٔ **جای** تصویر است نه محتوای کیوآر، پس
       هر تصویری کار را راه می‌اندازد و تست بی‌دلیل نمی‌شکند.
    """
    p = f"{TMP}/qr.png"
    try:
        import qrcode
        qrcode.make("https://ble.ir/nora?start=cv_TESTCODE").save(p)
    except ImportError:
        from PIL import Image, ImageDraw
        im = Image.new("RGB", (370, 370), "white")
        d = ImageDraw.Draw(im)
        for i in range(0, 370, 20):
            for j in range(0, 370, 20):
                if (i // 20 + j // 20) % 2 == 0:
                    d.rectangle([i, j, i + 19, j + 19], fill="black")
        im.save(p)
    return p


# ══════════════════════════════════════════════════════════
# 🔬 ابزار بازرسی خروجی
# ══════════════════════════════════════════════════════════
def inspect(path):
    """docx خروجی را می‌کاود"""
    x = zipfile.ZipFile(path).read("word/document.xml").decode("utf8")
    body = x.split("<w:body>")[1] if "<w:body>" in x else x
    # 🔴 با regex نشمار: پاراگراف خالی به شکل خودبستهٔ <w:p/>
    #    ذخیره می‌شود و «<w:p>» پیدایش نمی‌کند.
    try:
        import docx as _d
        n_par = len(_d.Document(path).paragraphs)
    except Exception:
        n_par = body.count("<w:p>") + body.count("<w:p ")
    pics = []
    for m in re.finditer(r"<wp:anchor.*?</wp:anchor>", x, re.S):
        blk = m.group(0)
        if "a:blip" not in blk:
            continue
        off = re.findall(r"<wp:posOffset>(-?\d+)</wp:posOffset>", blk)
        if len(off) >= 2:
            pics.append((int(off[0]) / MM, int(off[1]) / MM))
    return {
        "xml": x,
        "paragraphs": n_par,
        "page_breaks": x.count('w:type="page"'),
        "anchors": x.count("<wp:anchor"),
        "float_pics": pics,
        "ph_left": "{کیوآر}" in x,
    }


print("=" * 58)
print("🧪 تست فاز ۳۵ — جای کیوآر و صفحهٔ دوم")
print("=" * 58)

import certgen as cg

QR = make_qr()

# ── ۱) توابع تازه وجود دارند ──────────────────────────────
print("\n[۱] توابع تازه")
for fn in ("_first_body_paragraph", "_drop_trailing_blanks",
           "_textbox_anchor", "_strip_placeholder_text"):
    check(f"تابع {fn}", hasattr(cg, fn))

# ── ۲) دیگر add_paragraph ته سند نیست ─────────────────────
print("[۲] لنگر تصویر ته سند ساخته نمی‌شود")
import inspect as _ins
src = _ins.getsource(cg._add_floating_image)
# 🔎 فقط خطوط کد — کامنت‌ها «add_paragraph» را توضیح می‌دهند
_code = "\n".join(l for l in src.split("\n")
                  if not l.strip().startswith("#"))
check("_add_floating_image دیگر add_paragraph نمی‌کند",
      "d.add_paragraph()" not in _code)
check("از _first_body_paragraph استفاده می‌کند",
      "_first_body_paragraph" in src)
check("relativeHeight قطعی است (hash نیست)",
      "hash(key)" not in src)

# ── ۳) قالب بدون {کیوآر} → یتیم، ولی صفحهٔ اول ────────────
print("[۳] قالب بدون {کیوآر} — تصویر یتیم")
out = f"{TMP}/o_plain.docx"
r, e = cg.fill_docx(build_plain(), out, {"name": "نمونه"}, qr_png=QR)
check("خروجی ساخته شد", bool(r))
if r:
    d = inspect(out)
    check("کیوآر شناور درج شد", len(d["float_pics"]) >= 1)
    check("هیچ شکست صفحه‌ای نیست", d["page_breaks"] == 0)
    # لنگر باید روی پاراگراف پس‌زمینه باشد (اولی)، نه پاراگراف تازه
    before = d["xml"][:d["xml"].find("<wp:anchor")]
    n_before = before.count("<w:p>") + before.count("<w:p ")
    check("لنگر در پاراگراف اول است (نه ته سند)", n_before <= 1)

# ── ۴) تکست‌باکس DrawingML → دقیقاً سر جای کادر ───────────
print("[۴] تکست‌باکس مدرن — کیوآر دقیقاً سر جای کادر")
out = f"{TMP}/o_dml.docx"
r, e = cg.fill_docx(build_dml(), out, {"name": "نمونه"}, qr_png=QR)
check("خروجی ساخته شد", bool(r))
if r:
    d = inspect(out)
    check("متن خام {کیوآر} پاک شد", not d["ph_left"])
    check("هیچ شکست صفحه‌ای نیست", d["page_breaks"] == 0)
    check("تصویر شناور هست", len(d["float_pics"]) >= 1)
    if d["float_pics"]:
        # نزدیک‌ترین تصویر به کادر
        cx, cy = BOX_X + BOX_WH / 2, BOX_Y + BOX_WH / 2
        best = min(d["float_pics"],
                   key=lambda t: abs(t[0] - BOX_X) + abs(t[1] - BOX_Y))
        check(f"X نزدیک کادر است ({best[0]:.0f}mm ≈ {BOX_X:.0f})",
              abs(best[0] - BOX_X) <= 15)
        check(f"Y نزدیک کادر است ({best[1]:.0f}mm ≈ {BOX_Y:.0f})",
              abs(best[1] - BOX_Y) <= 15)
        check("داخل صفحه است", 0 <= best[0] <= W_MM and 0 <= best[1] <= H_MM)

# ── ۵) تکست‌باکس VML قدیمی ────────────────────────────────
print("[۵] تکست‌باکس قدیمی VML")
out = f"{TMP}/o_vml.docx"
r, e = cg.fill_docx(build_vml(), out, {"name": "نمونه"}, qr_png=QR)
check("خروجی ساخته شد", bool(r))
if r:
    d = inspect(out)
    check("متن خام {کیوآر} پاک شد", not d["ph_left"])
    check("هیچ شکست صفحه‌ای نیست", d["page_breaks"] == 0)
    if d["float_pics"]:
        best = min(d["float_pics"],
                   key=lambda t: abs(t[0] - BOX_X) + abs(t[1] - BOX_Y))
        check(f"X از style خوانده شد ({best[0]:.0f}mm)",
              abs(best[0] - BOX_X) <= 15)
        check(f"Y از style خوانده شد ({best[1]:.0f}mm)",
              abs(best[1] - BOX_Y) <= 15)

# ── ۶) {کیوآر} در پاراگراف معمولی — مسیر قدیمی سالم ───────
print("[۶] {کیوآر} در پاراگراف معمولی (سازگاری عقب‌رو)")
out = f"{TMP}/o_inline.docx"
r, e = cg.fill_docx(build_inline(), out, {"name": "نمونه"}, qr_png=QR)
check("خروجی ساخته شد", bool(r))
if r:
    d = inspect(out)
    check("متن {کیوآر} جایگزین شد", not d["ph_left"])
    check("هیچ شکست صفحه‌ای نیست", d["page_breaks"] == 0)
    check("تصویری درج شده", "a:blip" in d["xml"])

# ── ۷) نگهبان تک‌صفحه ─────────────────────────────────────
print("[۷] نگهبان تک‌صفحه — حذف خط‌های خالی ته سند")
tpl = build_trailing_blanks()
before_n = inspect(tpl)["paragraphs"]          # ۷ پاراگراف (۵ خالی)
out = f"{TMP}/o_blanks.docx"
r, e = cg.fill_docx(tpl, out, {"name": "نمونه"}, qr_png=QR)
check("خروجی ساخته شد", bool(r))
if r:
    d = inspect(out)
    check(f"پاراگراف‌های خالی حذف شدند ({before_n} → {d['paragraphs']})",
          d["paragraphs"] < before_n)
    check("هیچ شکست صفحه‌ای نیست", d["page_breaks"] == 0)

# پاراگرافی که تصویر دارد نباید حذف شود
print("[۷ب] نگهبان به تصویر دست نمی‌زند")
import docx as _dx
from docx.shared import Mm as _Mm
_d = _new_doc()
_d.add_paragraph().add_run().add_picture(_bg(f"{TMP}/bg.png"),
                                         width=_Mm(W_MM))
_before = len(_d.paragraphs)
cg._drop_trailing_blanks(_d)
check("پاراگراف دارای تصویر حذف نشد", len(_d.paragraphs) == _before)

# ── ۸) کلید تکراری تنظیمات ────────────────────────────────
print("[۸] تنظیمات")
from collections import Counter
c = Counter(row[0] for row in cg.CERT_SETTINGS)
dups = {k: v for k, v in c.items() if v > 1}
check(f"هیچ کلید تکراری در CERT_SETTINGS نیست ({dups})", not dups)
c2 = Counter(row[0] for row in cg.TICKET_SETTINGS)
check("هیچ کلید تکراری در TICKET_SETTINGS نیست",
      not {k: v for k, v in c2.items() if v > 1})

keys = {row[0] for row in cg.CERT_SETTINGS}
check("تنظیم cert_one_page هست", "cert_one_page" in keys)
check("تنظیم cert_pos_from_tpl هست", "cert_pos_from_tpl" in keys)
check("اعتبار پیش‌فرض ۲۴ ماه است",
      next(r[4] for r in cg.CERT_SETTINGS
           if r[0] == "cert_valid_months") == "24")

# همهٔ تنظیم‌های تازه در پنل قابل تغییرند
for k in ("cert_one_page", "cert_pos_from_tpl"):
    row = next((r for r in cg.CERT_SETTINGS if r[0] == k), None)
    check(f"{k} نوع bool دارد (دکمهٔ روشن/خاموش)",
          row is not None and row[3] == "bool")
    check(f"{k} توضیح فارسی دارد", row is not None and bool(row[5]))

# ── ۹) خاموش‌کردن تنظیم واقعاً اثر دارد ───────────────────
print("[۹] خاموش‌کردن «جای تصویر از خود قالب»")
_orig = cg.get_setting


def _fake(key, dflt=""):
    if key == "cert_pos_from_tpl":
        return "0"
    return _orig(key, dflt)


cg.get_setting = _fake
try:
    out = f"{TMP}/o_off.docx"
    r, e = cg.fill_docx(build_dml(), out, {"name": "نمونه"}, qr_png=QR)
    if r:
        d = inspect(out)
        if d["float_pics"]:
            best = d["float_pics"][0]
            # باید به گوشهٔ پیش‌فرض (۸۲٪،۶۲٪) برود، نه کادر
            check("با خاموش‌بودن، جای پیش‌فرض استفاده شد",
                  abs(best[0] - BOX_X) > 5 or abs(best[1] - BOX_Y) > 5)
finally:
    cg.get_setting = _orig

# ── ۱۰) استحکام: ورودی‌های عجیب ───────────────────────────
print("[۱۰] استحکام")
_d2 = _new_doc()
try:
    got = cg._textbox_anchor(_d2, "qr")
    check("سند بدون کادر → None برمی‌گرداند", got is None)
except Exception as ex:
    check(f"سند بدون کادر نباید خطا بدهد ({ex})", False)

try:
    p = cg._first_body_paragraph(_d2)
    check("سند خالی → پاراگراف می‌سازد", p is not None)
except Exception as ex:
    check(f"_first_body_paragraph نباید خطا بدهد ({ex})", False)

try:
    n = cg._strip_placeholder_text(_d2, "qr")
    check("_strip روی سند خالی امن است", n == 0)
except Exception as ex:
    check(f"_strip نباید خطا بدهد ({ex})", False)

# ── ۱۱) چند تصویر همزمان ──────────────────────────────────
print("[۱۱] چند تصویر همزمان (کیوآر + مهر + امضا)")
from PIL import Image as _PI
for nm in ("stamp", "sign"):
    _PI.new("RGB", (200, 200), (200, 60, 60)).save(f"{TMP}/{nm}.png")
out = f"{TMP}/o_multi.docx"
_orig2 = cg.get_setting


def _fake2(key, dflt=""):
    if key == "cert_force_images":
        return "qr,stamp,sign_img"
    return _orig2(key, dflt)


cg.get_setting = _fake2
try:
    r, e = cg.fill_docx(build_plain(), out, {"name": "نمونه"},
                        qr_png=QR,
                        images={"stamp": f"{TMP}/stamp.png",
                                "sign_img": f"{TMP}/sign.png"})
    check("خروجی چندتصویری ساخته شد", bool(r))
    if r:
        d = inspect(out)
        check(f"هر ۳ تصویر شناور شدند ({len(d['float_pics'])})",
              len(d["float_pics"]) >= 3)
        check("باز هم بدون شکست صفحه", d["page_breaks"] == 0)
        # هیچ‌کدام بیرون صفحه نیفتاده باشند
        check("همه داخل صفحه‌اند",
              all(0 <= x <= W_MM and 0 <= y <= H_MM
                  for x, y in d["float_pics"]))
finally:
    cg.get_setting = _orig2

# ── ۱۲) ترتیب OOXML (نگهبان فاز ۳۴) ───────────────────────
print("[۱۲] ترتیب عناصر OOXML حفظ شده")
out = f"{TMP}/o_order.docx"
r, e = cg.fill_docx(build_plain(), out, {"name": "نمونه"}, qr_png=QR)
if r:
    x = zipfile.ZipFile(out).read("word/document.xml").decode("utf8")
    m = re.search(r"<wp:anchor.*?</wp:anchor>", x, re.S)
    if m:
        blk = m.group(0)
        want = ["wp:simplePos", "wp:positionH", "wp:positionV",
                "wp:extent", "wp:wrapNone", "wp:docPr", "a:graphic"]
        pos = [blk.find("<" + t) for t in want]
        seen = [(t, p) for t, p in zip(want, pos) if p >= 0]
        ordered = all(seen[i][1] < seen[i + 1][1]
                      for i in range(len(seen) - 1))
        check("ترتیب OOXML درست است (لیبره سخت‌گیر است)", ordered)
        check("wrapNone بعد از extent است",
              blk.find("<wp:wrapNone") > blk.find("<wp:extent"))
        check("behindDoc=0 (جلوی پس‌زمینه)", 'behindDoc="0"' in blk)

# ══════════════════════════════════════════════════════════
print()
print("=" * 58)
print(f"✅ موفق: {_ok}   ❌ ناموفق: {_bad}")
if _fails:
    print("\nناموفق‌ها:")
    for f in _fails:
        print("  •", f)
print("=" * 58)

sys.exit(1 if _bad else 0)
