"""
📊 پاسخ‌ها، گزارش تحلیلی و خروجی فرم‌ها — فاز ۹

طبق معماری:
  • لیست پاسخ‌ها با صفحه‌بندی + مشاهده تک‌تک
  • برچسب‌گذاری پاسخ‌ها + یادداشت ادمین
  • تحلیل تک‌تک فیلدها (جدول/درصد/نمودار متنی)
  • خروجی Excel و متنی
  • نتایج نظرسنجی و انتخابات با نمودار
"""
import json
from datetime import datetime

from config import CURRENCY_LABEL, FILES_DIR
from database import db_execute, db_scalar, log
from auth import set_state, get_state_data, clear_state, is_admin as _is_admin
from bale_api import (
    send, send_or_edit, send_document, kb_cancel, kb_remove, BTN_BACK,
)
from forms import (
    FORM_TYPES, get_form, get_fields, get_field, field_meta, field_label,
    parse_options, form_can_manage, fmt_amount, get_tags_for, jnow,
    STATE_FB_EDIT,
)
import jalali as _jd

RESP_PER_PAGE = 6


def _guard(cid, uid, fid):
    form = get_form(fid)
    if not form:
        send(cid, "❌ فرم یافت نشد")
        return None
    if not form_can_manage(uid, form):
        send(cid, "❌ دسترسی ندارید")
        return None
    return form


# ==================== لیست پاسخ‌ها ====================
def show_responses(cid, uid, fid, page=0, flt="all"):
    form = _guard(cid, uid, fid)
    if not form:
        return

    where = "form_id=?"
    params = [fid]
    if flt == "completed":
        where += " AND status IN ('completed','confirmed')"
    elif flt == "pending":
        where += " AND status='pending_payment'"
    elif flt == "draft":
        where += " AND status='draft'"
    elif flt == "paid":
        where += " AND payment_status='paid'"
    else:
        where += " AND status<>'draft'"

    rows = db_execute(f"SELECT * FROM form_responses WHERE {where} ORDER BY id DESC",
                      tuple(params), fetch="all") or []

    t = FORM_TYPES.get(form.get("form_type") or "questionnaire",
                       FORM_TYPES["questionnaire"])

    if not rows:
        return send_or_edit(cid, (
            f"📊 {t['result_btn']} — {form['title'][:30]}\n"
            "━━━━━━━━━━━━━━\n\n📭 پاسخی ثبت نشده."
        ), {"inline_keyboard": [
            [{"text": "🎛️ همه پاسخ‌ها", "callback_data": f"fbr_{fid}"}],
            [{"text": "🔙 بازگشت", "callback_data": f"fbv_{fid}"}]]})

    per = RESP_PER_PAGE
    pages = max(1, (len(rows) + per - 1) // per)
    page = max(0, min(int(page or 0), pages - 1))
    chunk = rows[page * per:(page + 1) * per]

    lines = [f"📊 {t['result_btn']} — {form['title'][:28]}", "━━━━━━━━━━━━━━",
             f"📥 مجموع: {len(rows)}", ""]
    btns = []
    for r in chunk:
        u = db_execute("SELECT first_name, last_name, mobile FROM users "
                       "WHERE bale_user_id=?", (r.get("bale_user_id"),), fetch="one") or {}
        nm = f"{u.get('first_name') or ''} {u.get('last_name') or ''}".strip()
        if form.get("anonymous"):
            nm = "ناشناس"
        elif not nm:
            nm = str(r.get("bale_user_id") or "—")

        icon = {"completed": "✅", "confirmed": "🎫",
                "pending_payment": "⏳", "draft": "📝"}.get(r.get("status"), "▫️")
        tags = ""
        if r.get("tags"):
            try:
                tl = json.loads(r["tags"])
                tags = " " + " ".join(str(x) for x in tl[:3])
            except (ValueError, TypeError):
                pass
        lines.append(f"{icon} {nm[:26]}{tags}")
        lines.append(f"   🔖 {r.get('tracking_code') or '—'} | "
                     f"{(r.get('completed_at') or r.get('started_at') or '')[:16]}")
        if r.get("max_score"):
            lines.append(f"   🎯 نمره: {r.get('score')}/{r.get('max_score')}")
        btns.append([{"text": f"{icon} {nm[:28]}", "callback_data": f"fbrv_{r['id']}"}])

    if pages > 1:
        nav = []
        if page > 0:
            nav.append({"text": "⬅️", "callback_data": f"fbrp_{page-1}_{fid}"})
        nav.append({"text": f"{page+1}/{pages}", "callback_data": "noop"})
        if page < pages - 1:
            nav.append({"text": "➡️", "callback_data": f"fbrp_{page+1}_{fid}"})
        btns.append(nav)

    btns.append([{"text": "🎛️ فیلتر", "callback_data": f"fbrf_{fid}"},
                 {"text": "📈 گزارش", "callback_data": f"fbrep_{fid}"}])
    btns.append([{"text": "📄 خروجی Excel", "callback_data": f"fbx_{fid}"},
                 {"text": "📑 خروجی متنی", "callback_data": f"fbpdf_{fid}"}])
    btns.append([{"text": "🔙 بازگشت", "callback_data": f"fbv_{fid}"}])
    send_or_edit(cid, "\n".join(lines), {"inline_keyboard": btns})


def show_resp_filter(cid, uid, fid):
    form = _guard(cid, uid, fid)
    if not form:
        return
    send_or_edit(cid, "🎛️ فیلتر پاسخ‌ها\n━━━━━━━━━━━━━━\n\nیک فیلتر انتخاب کنید:",
                 {"inline_keyboard": [
                     [{"text": "📋 همه", "callback_data": f"fbrfs_all_{fid}"}],
                     [{"text": "✅ تکمیل‌شده", "callback_data": f"fbrfs_completed_{fid}"}],
                     [{"text": "⏳ در انتظار پرداخت", "callback_data": f"fbrfs_pending_{fid}"}],
                     [{"text": "💰 پرداخت‌شده", "callback_data": f"fbrfs_paid_{fid}"}],
                     [{"text": "📝 نیمه‌کاره", "callback_data": f"fbrfs_draft_{fid}"}],
                     [{"text": "🔙 بازگشت", "callback_data": f"fbr_{fid}"}]]})


# ==================== مشاهده یک پاسخ ====================
def view_response(cid, uid, rid):
    r = db_execute("SELECT * FROM form_responses WHERE id=?", (int(rid),), fetch="one")
    if not r:
        return send(cid, "❌ یافت نشد")
    form = _guard(cid, uid, r["form_id"])
    if not form:
        return

    u = db_execute("SELECT * FROM users WHERE bale_user_id=?",
                   (r.get("bale_user_id"),), fetch="one") or {}
    nm = f"{u.get('first_name') or ''} {u.get('last_name') or ''}".strip() or "—"
    if form.get("anonymous"):
        nm = "🕶️ ناشناس"

    lines = [f"📄 پاسخ #{r['id']}", "━━━━━━━━━━━━━━", "",
             f"👤 {nm}"]
    if not form.get("anonymous"):
        if u.get("mobile"):
            lines.append(f"📱 {u['mobile']}")
        lines.append(f"🆔 {r.get('bale_user_id') or '—'}")
    lines += [f"🔖 {r.get('tracking_code') or '—'}",
              f"📅 {(r.get('completed_at') or _jd.fmt(r.get('started_at'), 'datetime'))}",
              f"📊 وضعیت: {_status_name(r.get('status'))}"]

    if r.get("amount"):
        pay_st = {"paid": "✅ پرداخت‌شده", "pending": "⏳ در انتظار",
                  "failed": "❌ ناموفق"}.get(r.get("payment_status"), "—")
        lines.append(f"💰 {fmt_amount(r['amount'])} {CURRENCY_LABEL} — {pay_st}")
    if r.get("max_score"):
        pct = int((r.get("score") or 0) * 100 / r["max_score"]) if r["max_score"] else 0
        lines.append(f"🎯 نمره: {r.get('score')}/{r['max_score']} ({pct}%)")
    if r.get("tags"):
        try:
            tl = json.loads(r["tags"])
            if tl:
                lines.append("🏷️ " + " ".join(str(x) for x in tl))
        except (ValueError, TypeError):
            pass
    if r.get("admin_note"):
        lines.append(f"📝 یادداشت: {r['admin_note']}")

    lines += ["", "━━━ پاسخ‌ها ━━━", ""]
    rows = db_execute("""SELECT a.*, f.label, f.field_type ft, f.field_order
        FROM form_answers a LEFT JOIN form_fields f ON f.id=a.field_id
        WHERE a.response_id=? ORDER BY f.field_order, a.id""", (rid,), fetch="all") or []

    media = []
    for a in rows:
        fl = a.get("label") or field_meta(a.get("ft")).get("name", "سوال")
        v = a.get("value_text") or ""
        if a.get("file_id"):
            media.append((fl, a["file_id"], a.get("file_type")))
            v = v or f"📎 {a.get('file_type') or 'فایل'}"
        lines.append(f"▫️ {fl[:44]}:\n   {v or '—'}")

    btns = [[{"text": "🏷️ برچسب‌گذاری", "callback_data": f"fbrt_{rid}"},
             {"text": "📝 یادداشت", "callback_data": f"fbrn_{rid}"}]]
    if media:
        btns.append([{"text": f"📎 مشاهده فایل‌ها ({len(media)})",
                      "callback_data": f"fbrm_{rid}"}])
    if r.get("status") == "pending_payment":
        btns.append([{"text": "✅ تأیید دستی پرداخت", "callback_data": f"fbrok_{rid}"}])
    if form.get("ticket_enabled") and r.get("status") in ("completed", "confirmed"):
        btns.append([{"text": "🎫 ارسال بلیت به کاربر", "callback_data": f"fbrtk_{rid}"}])
    btns.append([{"text": "💬 پیام به کاربر", "callback_data": f"fbrmsg_{rid}"}])
    btns.append([{"text": "🗑️ حذف پاسخ", "callback_data": f"fbrdel_{rid}"}])
    btns.append([{"text": "🔙 لیست پاسخ‌ها", "callback_data": f"fbr_{r['form_id']}"}])
    send_or_edit(cid, "\n".join(lines), {"inline_keyboard": btns})


def _status_name(s):
    return {"draft": "📝 نیمه‌کاره", "completed": "✅ تکمیل‌شده",
            "confirmed": "🎫 تأییدشده", "pending_payment": "⏳ در انتظار پرداخت",
            "rejected": "❌ ردشده"}.get(s, s or "—")


def show_response_media(cid, uid, rid):
    r = db_execute("SELECT * FROM form_responses WHERE id=?", (int(rid),), fetch="one")
    if not r:
        return send(cid, "❌ یافت نشد")
    if not _guard(cid, uid, r["form_id"]):
        return
    rows = db_execute("""SELECT a.*, f.label FROM form_answers a
        LEFT JOIN form_fields f ON f.id=a.field_id
        WHERE a.response_id=? AND a.file_id IS NOT NULL""", (rid,), fetch="all") or []
    if not rows:
        return send(cid, "📭 فایلی ثبت نشده")

    from bale_api import send_photo, send_video, send_voice
    sent = 0
    for a in rows:
        cap = a.get("label") or "فایل"
        ft = a.get("file_type")
        try:
            if ft == "photo":
                send_photo(cid, a["file_id"], cap)
            elif ft == "video":
                send_video(cid, a["file_id"], cap)
            elif ft in ("voice", "audio"):
                send_voice(cid, a["file_id"], cap)
            else:
                send_document(cid, a["file_id"], cap)
            sent += 1
        except Exception as e:
            log(f"⚠️ ارسال فایل پاسخ: {e}")
        # چند فایلی
        if a.get("value_options"):
            try:
                arr = json.loads(a["value_options"])
                for x in arr[1:] if isinstance(arr, list) else []:
                    if isinstance(x, dict) and x.get("file_id"):
                        if x.get("file_type") == "photo":
                            send_photo(cid, x["file_id"], cap)
                        else:
                            send_document(cid, x["file_id"], cap)
                        sent += 1
            except (ValueError, TypeError):
                pass

    return send(cid, f"📎 {sent} فایل ارسال شد", {"inline_keyboard": [
        [{"text": "🔙 بازگشت", "callback_data": f"fbrv_{rid}"}]]})


# ==================== برچسب و یادداشت ====================
def show_tag_picker(cid, uid, rid):
    r = db_execute("SELECT * FROM form_responses WHERE id=?", (int(rid),), fetch="one")
    if not r:
        return send(cid, "❌ یافت نشد")
    if not _guard(cid, uid, r["form_id"]):
        return

    cur = []
    if r.get("tags"):
        try:
            cur = json.loads(r["tags"]) or []
        except (ValueError, TypeError):
            cur = []

    tags = get_tags_for(r["form_id"])
    if not tags:
        return send_or_edit(cid, (
            "🏷️ برچسب‌گذاری\n━━━━━━━━━━━━━━\n\n"
            "📭 هنوز برچسبی تعریف نشده.\n\n"
            "از «تنظیمات فرم‌ساز ← برچسب‌های عمومی» اضافه کنید."
        ), {"inline_keyboard": [
            [{"text": "➕ افزودن برچسب", "callback_data": "fb_tagadd"}],
            [{"text": "🔙 بازگشت", "callback_data": f"fbrv_{rid}"}]]})

    btns = []
    for t in tags:
        label = f"{t.get('color') or '⚪'} {t['name']}"
        mark = "✅" if label in cur else "⬜"
        btns.append([{"text": f"{mark} {label}", "callback_data": f"fbrts_{t['id']}_{rid}"}])
    btns.append([{"text": "🔙 بازگشت", "callback_data": f"fbrv_{rid}"}])
    send_or_edit(cid, (
        f"🏷️ برچسب‌های این پاسخ\n━━━━━━━━━━━━━━\n\n"
        f"فعلی: {' '.join(cur) if cur else '—'}\n\n"
        f"روی برچسب بزنید تا اضافه/حذف شود:"
    ), {"inline_keyboard": btns})


def toggle_tag(cid, uid, tag_id, rid):
    r = db_execute("SELECT * FROM form_responses WHERE id=?", (int(rid),), fetch="one")
    if not r:
        return send(cid, "❌ یافت نشد")
    if not _guard(cid, uid, r["form_id"]):
        return
    t = db_execute("SELECT * FROM form_tags WHERE id=?", (int(tag_id),), fetch="one")
    if not t:
        return send(cid, "❌ برچسب یافت نشد")

    label = f"{t.get('color') or '⚪'} {t['name']}"
    cur = []
    if r.get("tags"):
        try:
            cur = json.loads(r["tags"]) or []
        except (ValueError, TypeError):
            cur = []
    if label in cur:
        cur.remove(label)
    else:
        cur.append(label)
    db_execute("UPDATE form_responses SET tags=? WHERE id=?",
               (json.dumps(cur, ensure_ascii=False), rid))
    return show_tag_picker(cid, uid, rid)


def start_note(cid, uid, rid):
    r = db_execute("SELECT * FROM form_responses WHERE id=?", (int(rid),), fetch="one")
    if not r:
        return send(cid, "❌ یافت نشد")
    if not _guard(cid, uid, r["form_id"]):
        return
    set_state(uid, STATE_FB_EDIT, {"key": "resp_note", "rid": rid}, merge=False)
    send(cid, (
        f"📝 یادداشت ادمین\n━━━━━━━━━━━━━━\n\n"
        f"📌 فعلی: {r.get('admin_note') or '—'}\n\n"
        f"یادداشت جدید را بنویسید:\n\n💡 «-» = حذف"
    ), kb_cancel())


def start_msg_user(cid, uid, rid):
    r = db_execute("SELECT * FROM form_responses WHERE id=?", (int(rid),), fetch="one")
    if not r:
        return send(cid, "❌ یافت نشد")
    if not _guard(cid, uid, r["form_id"]):
        return
    if not r.get("bale_user_id"):
        return send(cid, "⚠️ کاربر شناسایی نشد")
    set_state(uid, STATE_FB_EDIT, {"key": "resp_msg", "rid": rid}, merge=False)
    send(cid, (
        "💬 پیام به کاربر\n━━━━━━━━━━━━━━\n\n"
        "متنی که برای این کاربر ارسال شود را بنویسید:"
    ), kb_cancel())


def handle_resp_input(message):
    """یادداشت یا پیام به کاربر"""
    uid = message["from"]["id"]
    cid = message["chat"]["id"]
    text = (message.get("text") or "").strip()
    data = get_state_data(uid)
    rid = data.get("rid")
    key = data.get("key")
    clear_state(uid)

    if text in (BTN_BACK, "🔙 بازگشت"):
        send(cid, "❌ لغو شد", kb_remove())
        return view_response(cid, uid, rid)

    r = db_execute("SELECT * FROM form_responses WHERE id=?", (int(rid),), fetch="one")
    if not r:
        return send(cid, "❌ یافت نشد", kb_remove())

    if key == "resp_note":
        val = "" if text == "-" else text[:1000]
        db_execute("UPDATE form_responses SET admin_note=?, reviewed_by=?, "
                   "reviewed_at=CURRENT_TIMESTAMP WHERE id=?", (val, uid, rid))
        send(cid, "✅ یادداشت ذخیره شد", kb_remove())
        return view_response(cid, uid, rid)

    if key == "resp_msg":
        form = get_form(r["form_id"])
        target = r.get("bale_user_id")
        if not target:
            return send(cid, "⚠️ کاربر شناسایی نشد", kb_remove())
        send(target, (
            f"💬 پیام از مدیریت\n━━━━━━━━━━━━━━\n\n"
            f"📋 درباره فرم: {form['title'] if form else ''}\n"
            f"🔖 {r.get('tracking_code') or ''}\n\n{text}"
        ))
        send(cid, "✅ پیام ارسال شد", kb_remove())
        return view_response(cid, uid, rid)

    return send(cid, "⚠️ نامعتبر", kb_remove())


def approve_response(cid, uid, rid):
    """تأیید دستی پاسخ در انتظار پرداخت"""
    r = db_execute("SELECT * FROM form_responses WHERE id=?", (int(rid),), fetch="one")
    if not r:
        return send(cid, "❌ یافت نشد")
    form = _guard(cid, uid, r["form_id"])
    if not form:
        return
    from forms import gen_tracking, gen_qr_code, recount_responses
    tracking = r.get("tracking_code") or gen_tracking()
    qr = r.get("qr_code") or gen_qr_code()
    db_execute("""UPDATE form_responses SET status='confirmed', payment_status='paid',
        tracking_code=?, qr_code=?, completed_at=COALESCE(completed_at,CURRENT_TIMESTAMP),
        reviewed_by=?, reviewed_at=CURRENT_TIMESTAMP WHERE id=?""",
        (tracking, qr, uid, rid))
    recount_responses(form["id"])

    # 👤 اطلاعات کاربر را در پروفایلش هم ذخیره کن (درخواست کاربر)
    filled = 0
    if form.get("fill_profile", 1):
        try:
            from forms import sync_profile_from_response
            filled = sync_profile_from_response(rid)
        except Exception as e:
            log(f"⚠️ پرکردن پروفایل: {e}")

    if r.get("bale_user_id"):
        msg = (f"✅ ثبت‌نام شما تأیید شد!\n\n"
               f"📋 {form['title']}\n🔖 کد پیگیری: {tracking}")
        if filled:
            msg += f"\n\n👤 {filled} مورد از اطلاعات شما در پروفایلتان ذخیره شد."
        send(r["bale_user_id"], msg)
        if form.get("ticket_enabled"):
            try:
                import form_fill
                form_fill.send_ticket(r["bale_user_id"], r["bale_user_id"], rid)
            except Exception as e:
                log(f"⚠️ بلیت: {e}")
    send(cid, "✅ تأیید شد")
    return view_response(cid, uid, rid)


def send_ticket_to_user(cid, uid, rid):
    r = db_execute("SELECT * FROM form_responses WHERE id=?", (int(rid),), fetch="one")
    if not r:
        return send(cid, "❌ یافت نشد")
    if not _guard(cid, uid, r["form_id"]):
        return
    target = r.get("bale_user_id")
    if not target:
        return send(cid, "⚠️ کاربر شناسایی نشد")
    try:
        import form_fill
        form_fill.send_ticket(target, target, rid)
        send(cid, "🎫 بلیت برای کاربر ارسال شد")
    except Exception as e:
        log(f"⚠️ ارسال بلیت: {e}")
        send(cid, "⚠️ خطا در ارسال بلیت")
    return view_response(cid, uid, rid)


def delete_response(cid, uid, rid, confirmed=False):
    r = db_execute("SELECT * FROM form_responses WHERE id=?", (int(rid),), fetch="one")
    if not r:
        return send(cid, "❌ یافت نشد")
    form = _guard(cid, uid, r["form_id"])
    if not form:
        return
    if not confirmed:
        return send_or_edit(cid, (
            f"🗑️ حذف پاسخ #{rid}\n━━━━━━━━━━━━━━\n\n"
            f"🔖 {r.get('tracking_code') or '—'}\n\n"
            f"⚠️ این عمل بازگشت‌پذیر نیست."
        ), {"inline_keyboard": [
            [{"text": "🗑️ بله، حذف کن", "callback_data": f"fbrdy_{rid}"}],
            [{"text": "❌ انصراف", "callback_data": f"fbrv_{rid}"}]]})
    fid = r["form_id"]
    db_execute("DELETE FROM form_answers WHERE response_id=?", (rid,))
    db_execute("DELETE FROM form_responses WHERE id=?", (rid,))
    from forms import recount_responses
    recount_responses(fid)
    send(cid, "🗑️ پاسخ حذف شد")
    return show_responses(cid, uid, fid, 0)


# ==================== گزارش تحلیلی ====================
def show_report(cid, uid, fid):
    """📈 گزارش و تحلیل — آمار کلی + تحلیل هر فیلد"""
    form = _guard(cid, uid, fid)
    if not form:
        return

    total = db_scalar("SELECT COUNT(*) FROM form_responses WHERE form_id=? "
                      "AND status IN ('completed','confirmed')", (fid,))
    pending = db_scalar("SELECT COUNT(*) FROM form_responses WHERE form_id=? "
                        "AND status='pending_payment'", (fid,))
    draft = db_scalar("SELECT COUNT(*) FROM form_responses WHERE form_id=? "
                      "AND status='draft'", (fid,))
    paid = db_scalar("SELECT COALESCE(SUM(amount),0) FROM form_responses "
                     "WHERE form_id=? AND payment_status='paid'", (fid,))
    views = form.get("views_count") or 0
    conv = (total * 100 / views) if views else 0

    lines = [
        f"📈 گزارش فرم «{form['title'][:28]}»",
        "━━━━━━━━━━━━━━",
        "",
        "📊 آمار کلی:",
        f"├── 👁️ بازدید: {views}",
        f"├── ✅ تکمیل‌شده: {total}",
        f"├── ⏳ در انتظار پرداخت: {pending}",
        f"├── 📝 نیمه‌کاره: {draft}",
        f"├── 📈 نرخ تبدیل: {conv:.1f}%",
        f"└── 💰 مجموع پرداختی: {fmt_amount(paid)} {CURRENCY_LABEL}",
    ]

    if form.get("limit_enabled") and form.get("max_responses"):
        left = int(form["max_responses"]) - total
        pct = int(total * 100 / int(form["max_responses"]))
        blocks = max(0, min(10, pct // 10))
        lines += ["", f"👥 ظرفیت: {total}/{form['max_responses']} ({pct}%)",
                  "▓" * blocks + "░" * (10 - blocks),
                  f"🪑 باقیمانده: {max(0, left)}"]

    # نمره آزمون
    if form.get("form_type") == "exam" and total:
        avg = db_scalar("SELECT COALESCE(AVG(score),0) FROM form_responses "
                        "WHERE form_id=? AND status IN ('completed','confirmed')", (fid,))
        mx = db_scalar("SELECT COALESCE(MAX(score),0) FROM form_responses "
                       "WHERE form_id=?", (fid,))
        mn = db_scalar("SELECT COALESCE(MIN(score),0) FROM form_responses "
                       "WHERE form_id=? AND status IN ('completed','confirmed')", (fid,))
        lines += ["", "🎯 نمرات:", f"├── میانگین: {float(avg):.1f}",
                  f"├── بیشترین: {mx}", f"└── کمترین: {mn}"]

    lines += ["", "💡 برای تحلیل هر سوال، دکمه‌های زیر را بزنید."]

    btns = [[{"text": "📋 تحلیل تک‌تک سوالات", "callback_data": f"fbfa2_{fid}"}]]
    if form.get("form_type") in ("survey", "election"):
        btns.append([{"text": "🥧 نمایش نتایج", "callback_data": f"fbres_{fid}"}])
    btns.append([{"text": "📄 خروجی Excel", "callback_data": f"fbx_{fid}"},
                 {"text": "📑 خروجی متنی", "callback_data": f"fbpdf_{fid}"}])
    btns.append([{"text": "📊 لیست پاسخ‌ها", "callback_data": f"fbr_{fid}"}])
    btns.append([{"text": "🔙 بازگشت", "callback_data": f"fbv_{fid}"}])
    send_or_edit(cid, "\n".join(lines), {"inline_keyboard": btns})


def show_field_analysis_menu(cid, uid, fid):
    form = _guard(cid, uid, fid)
    if not form:
        return
    fields = get_fields(fid)
    if not fields:
        return send(cid, "📭 سوالی نیست")
    btns = []
    for i, f in enumerate(fields, 1):
        m = field_meta(f.get("field_type"))
        btns.append([{"text": f"{i}. {m['icon']} {field_label(f)[:30]}",
                      "callback_data": f"fban_{f['id']}"}])
    btns.append([{"text": "🔙 بازگشت", "callback_data": f"fbrep_{fid}"}])
    send_or_edit(cid, (
        "📋 تحلیل سوالات\n━━━━━━━━━━━━━━\n\n"
        "سوال مورد نظر را انتخاب کنید:"
    ), {"inline_keyboard": btns})


def analyze_field(cid, uid, field_id):
    """تحلیل یک فیلد — جدول/درصد/نمودار متنی"""
    f = get_field(field_id)
    if not f:
        return send(cid, "❌ یافت نشد")
    form = _guard(cid, uid, f["form_id"])
    if not form:
        return

    m = field_meta(f.get("field_type"))
    kind = m.get("kind")

    rows = db_execute("""SELECT a.value_text, a.value_number, a.value_options
        FROM form_answers a JOIN form_responses r ON r.id=a.response_id
        WHERE a.field_id=? AND r.status IN ('completed','confirmed')""",
        (field_id,), fetch="all") or []

    answered = [r for r in rows if (r.get("value_text") or "").strip()
                or r.get("value_number") is not None]
    total_resp = db_scalar("SELECT COUNT(*) FROM form_responses WHERE form_id=? "
                           "AND status IN ('completed','confirmed')", (f["form_id"],))

    lines = [f"📊 تحلیل: {field_label(f)[:40]}", "━━━━━━━━━━━━━━", "",
             f"{m['icon']} نوع: {m['name']}",
             f"📥 پاسخ داده: {len(answered)} از {total_resp}"]
    if total_resp:
        lines.append(f"📈 نرخ پاسخ: {len(answered)*100/total_resp:.1f}%")
    lines.append("")

    if not answered:
        lines.append("📭 هنوز پاسخی ثبت نشده")
        return send_or_edit(cid, "\n".join(lines), {"inline_keyboard": [
            [{"text": "🔙 بازگشت", "callback_data": f"fbfa2_{f['form_id']}"}]]})

    # ---- فیلدهای گزینه‌ای / امتیازی ----
    if kind in ("choice", "dropdown", "quiz", "multi", "rating", "confirm",
                "education", "province_city", "item"):
        counts = {}
        for r in answered:
            vals = []
            if r.get("value_options"):
                try:
                    arr = json.loads(r["value_options"])
                    if isinstance(arr, list):
                        vals = [str(x) for x in arr if not isinstance(x, dict)]
                except (ValueError, TypeError):
                    pass
            if not vals:
                vals = [str(r.get("value_text") or "").strip()]
            for v in vals:
                if v:
                    counts[v] = counts.get(v, 0) + 1

        tot = sum(counts.values()) or 1
        ordered = sorted(counts.items(), key=lambda x: -x[1])
        lines.append("┌─── نتایج ───")
        for i, (k, n) in enumerate(ordered[:15], 1):
            pct = n * 100 / tot
            bar = "█" * max(1, int(pct / 5))
            lines.append(f"│ {i}. {k[:26]}")
            lines.append(f"│    {bar} {n} ({pct:.1f}%)")
        lines.append("└──────────")

        if kind == "rating":
            nums = [float(r["value_number"]) for r in answered
                    if r.get("value_number") is not None]
            if nums:
                lines += ["", f"⭐ میانگین امتیاز: {sum(nums)/len(nums):.2f}",
                          f"📈 بیشترین: {int(max(nums))} | کمترین: {int(min(nums))}"]
        if kind == "quiz" and f.get("correct_option"):
            opts = parse_options(f)
            ci = int(f["correct_option"]) if str(f["correct_option"]).isdigit() else 0
            if 1 <= ci <= len(opts):
                right = counts.get(opts[ci - 1], 0)
                lines += ["", f"✅ پاسخ صحیح: {opts[ci-1]}",
                          f"🎯 درصد صحیح: {right*100/tot:.1f}%"]

    # ---- عددی ----
    elif kind in ("number", "amount", "duration"):
        nums = []
        for r in answered:
            if r.get("value_number") is not None:
                nums.append(float(r["value_number"]))
            else:
                s = "".join(ch for ch in str(r.get("value_text") or "")
                            if ch.isdigit() or ch == ".")
                if s:
                    try:
                        nums.append(float(s))
                    except ValueError:
                        pass
        if nums:
            nums.sort()
            mid = len(nums) // 2
            median = nums[mid] if len(nums) % 2 else (nums[mid - 1] + nums[mid]) / 2
            lines += ["📐 آمار عددی:",
                      f"├── میانگین: {sum(nums)/len(nums):,.1f}",
                      f"├── میانه: {median:,.1f}",
                      f"├── بیشترین: {max(nums):,.0f}",
                      f"├── کمترین: {min(nums):,.0f}",
                      f"└── مجموع: {sum(nums):,.0f}"]
        else:
            lines.append("⚠️ داده عددی معتبری نیست")

    # ---- تاریخ ----
    elif kind in ("date", "datetime", "time"):
        vals = sorted(str(r.get("value_text") or "") for r in answered
                      if r.get("value_text"))
        if vals:
            lines += ["📅 بازه:", f"├── قدیمی‌ترین: {vals[0]}",
                      f"└── جدیدترین: {vals[-1]}", "", "🔢 پرتکرارها:"]
            counts = {}
            for v in vals:
                counts[v] = counts.get(v, 0) + 1
            for k, n in sorted(counts.items(), key=lambda x: -x[1])[:5]:
                lines.append(f"   • {k} — {n} بار")

    # ---- فایل ----
    elif kind in ("photo", "file", "video", "voice"):
        nfiles = db_scalar("""SELECT COUNT(*) FROM form_answers a
            JOIN form_responses r ON r.id=a.response_id
            WHERE a.field_id=? AND a.file_id IS NOT NULL
              AND r.status IN ('completed','confirmed')""", (field_id,))
        lines += [f"📎 فایل‌های دریافتی: {nfiles}"]

    # ---- موقعیت ----
    elif kind == "location":
        lines += [f"📍 موقعیت‌های ثبت‌شده: {len(answered)}", "", "نمونه:"]
        for r in answered[:5]:
            lines.append(f"   • {r.get('value_text')}")

    # ---- متنی ----
    else:
        texts = [str(r.get("value_text") or "").strip() for r in answered]
        texts = [t for t in texts if t]
        if texts:
            avg_len = sum(len(t) for t in texts) / len(texts)
            lines += [f"✍️ میانگین طول پاسخ: {avg_len:.0f} کاراکتر", "",
                      "📝 نمونه پاسخ‌ها:"]
            for t in texts[:8]:
                lines.append(f"   • {t[:56]}")
            if len(texts) > 8:
                lines.append(f"   … و {len(texts)-8} پاسخ دیگر")
            counts = {}
            for t in texts:
                counts[t] = counts.get(t, 0) + 1
            rep = [(k, n) for k, n in counts.items() if n > 1]
            if rep:
                lines += ["", "🔁 پاسخ‌های تکراری:"]
                for k, n in sorted(rep, key=lambda x: -x[1])[:5]:
                    lines.append(f"   • {k[:36]} — {n} بار")

    fields = get_fields(f["form_id"])
    idx = next((i for i, x in enumerate(fields) if x["id"] == f["id"]), 0)
    btns = []
    nav = []
    if idx > 0:
        nav.append({"text": "⬅️ سوال قبل", "callback_data": f"fban_{fields[idx-1]['id']}"})
    if idx < len(fields) - 1:
        nav.append({"text": "سوال بعد ➡️", "callback_data": f"fban_{fields[idx+1]['id']}"})
    if nav:
        btns.append(nav)
    btns.append([{"text": "📋 همه سوالات", "callback_data": f"fbfa2_{f['form_id']}"}])
    btns.append([{"text": "🔙 گزارش", "callback_data": f"fbrep_{f['form_id']}"}])
    send_or_edit(cid, "\n".join(lines), {"inline_keyboard": btns})


def show_results(cid, uid, fid):
    """🥧 نمایش نتایج — برای نظرسنجی و انتخابات"""
    form = _guard(cid, uid, fid)
    if not form:
        return
    fields = get_fields(fid)
    choice_fields = [f for f in fields
                     if field_meta(f.get("field_type")).get("kind")
                     in ("choice", "dropdown", "multi", "rating", "item")]

    total = db_scalar("SELECT COUNT(*) FROM form_responses WHERE form_id=? "
                      "AND status IN ('completed','confirmed')", (fid,))
    lines = [f"🥧 نتایج «{form['title'][:30]}»", "━━━━━━━━━━━━━━",
             f"👥 مشارکت‌کنندگان: {total}", ""]

    if not choice_fields:
        lines.append("⚠️ این فرم سوال گزینه‌ای ندارد.")
    for f in choice_fields[:6]:
        m = field_meta(f.get("field_type"))
        lines.append(f"\n{m['icon']} {field_label(f)[:44]}")
        rows = db_execute("""SELECT a.value_text, a.value_options FROM form_answers a
            JOIN form_responses r ON r.id=a.response_id
            WHERE a.field_id=? AND r.status IN ('completed','confirmed')""",
            (f["id"],), fetch="all") or []
        counts = {}
        for r in rows:
            vals = []
            if r.get("value_options"):
                try:
                    arr = json.loads(r["value_options"])
                    if isinstance(arr, list):
                        vals = [str(x) for x in arr if not isinstance(x, dict)]
                except (ValueError, TypeError):
                    pass
            if not vals:
                v = str(r.get("value_text") or "").strip()
                vals = [v] if v else []
            for v in vals:
                counts[v] = counts.get(v, 0) + 1
        tot = sum(counts.values()) or 1
        for k, n in sorted(counts.items(), key=lambda x: -x[1])[:8]:
            pct = n * 100 / tot
            bar = "█" * max(1, int(pct / 5))
            lines.append(f"   {bar} {k[:22]} — {n} ({pct:.1f}%)")
        if not counts:
            lines.append("   📭 بدون پاسخ")

    send_or_edit(cid, "\n".join(lines), {"inline_keyboard": [
        [{"text": "📋 تحلیل سوالات", "callback_data": f"fbfa2_{fid}"}],
        [{"text": "🔙 گزارش", "callback_data": f"fbrep_{fid}"}]]})


# ==================== خروجی Excel ====================
def export_excel(cid, uid, fid):
    form = _guard(cid, uid, fid)
    if not form:
        return
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        return send(cid, "❌ openpyxl نصب نیست:\npy -m pip install openpyxl")

    fields = get_fields(fid)
    rows = db_execute("""SELECT * FROM form_responses WHERE form_id=?
        AND status<>'draft' ORDER BY id""", (fid,), fetch="all") or []
    if not rows:
        return send(cid, "📭 پاسخی برای خروجی نیست")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "پاسخ‌ها"
    ws.sheet_view.rightToLeft = True

    head = ["ردیف", "کد پیگیری", "نام", "موبایل", "وضعیت",
            "تاریخ ثبت", "مبلغ", "وضعیت پرداخت", "نمره", "برچسب‌ها", "یادداشت"]
    head += [field_label(f)[:40] for f in fields]
    ws.append(head)

    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="4472C4")
    for i in range(1, len(head) + 1):
        c = ws.cell(row=1, column=i)
        c.font = bold
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    for n, r in enumerate(rows, 1):
        u = db_execute("SELECT first_name, last_name, mobile FROM users "
                       "WHERE bale_user_id=?", (r.get("bale_user_id"),), fetch="one") or {}
        nm = f"{u.get('first_name') or ''} {u.get('last_name') or ''}".strip()
        if form.get("anonymous"):
            nm, mob = "ناشناس", ""
        else:
            mob = u.get("mobile") or ""
        tags = ""
        if r.get("tags"):
            try:
                tags = " ".join(json.loads(r["tags"]))
            except (ValueError, TypeError):
                pass

        line = [n, r.get("tracking_code") or "", nm, mob,
                _status_name(r.get("status")),
                (r.get("completed_at") or r.get("started_at") or "")[:19],
                r.get("amount") or 0, r.get("payment_status") or "",
                (f"{r.get('score')}/{r.get('max_score')}" if r.get("max_score") else ""),
                tags, r.get("admin_note") or ""]

        amap = {}
        for a in db_execute("SELECT * FROM form_answers WHERE response_id=?",
                            (r["id"],), fetch="all") or []:
            v = a.get("value_text") or ""
            if not v and a.get("value_options"):
                try:
                    arr = json.loads(a["value_options"])
                    v = "، ".join(str(x) for x in arr if not isinstance(x, dict))
                except (ValueError, TypeError):
                    v = ""
            if not v and a.get("file_id"):
                v = f"[{a.get('file_type') or 'فایل'}]"
            amap[a.get("field_id")] = v
        line += [amap.get(f["id"], "") for f in fields]
        ws.append(line)

    for i, h in enumerate(head, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = \
            max(10, min(32, len(str(h)) + 6))
    ws.freeze_panes = "A2"

    # شیت آمار
    ws2 = wb.create_sheet("آمار")
    ws2.sheet_view.rightToLeft = True
    ws2.append(["سوال", "نوع", "تعداد پاسخ", "پرتکرارترین", "تعداد"])
    for c in range(1, 6):
        ws2.cell(row=1, column=c).font = bold
        ws2.cell(row=1, column=c).fill = fill
    for f in fields:
        m = field_meta(f.get("field_type"))
        ans = db_execute("""SELECT a.value_text FROM form_answers a
            JOIN form_responses r ON r.id=a.response_id
            WHERE a.field_id=? AND r.status IN ('completed','confirmed')
              AND a.value_text IS NOT NULL AND a.value_text<>''""",
            (f["id"],), fetch="all") or []
        counts = {}
        for a in ans:
            v = a["value_text"]
            counts[v] = counts.get(v, 0) + 1
        top = max(counts.items(), key=lambda x: x[1]) if counts else ("—", 0)
        ws2.append([field_label(f)[:50], m["name"], len(ans), str(top[0])[:40], top[1]])
    for i in range(1, 6):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 26

    from pathlib import Path
    out = Path(FILES_DIR) / "forms" / "exports"
    out.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    fp = out / f"form_{fid}_{stamp}.xlsx"
    try:
        wb.save(fp)
    except Exception as e:
        log(f"❌ ذخیره اکسل: {e}")
        return send(cid, "❌ خطا در ساخت فایل")

    res = send_document(cid, str(fp), (
        f"📄 خروجی فرم «{form['title'][:40]}»\n\n"
        f"📊 {len(rows)} پاسخ | {len(fields)} سوال\n"
        f"📅 {jnow()}"
    ))
    if not res.get("ok"):
        send(cid, f"⚠️ ارسال فایل ناموفق بود.\n📁 مسیر: {fp}")
    return send(cid, "✅ خروجی آماده شد", {"inline_keyboard": [
        [{"text": "🔙 بازگشت", "callback_data": f"fbv_{fid}"}]]})


def export_text(cid, uid, fid):
    """📑 خروجی متنی کامل (قابل چاپ)"""
    form = _guard(cid, uid, fid)
    if not form:
        return
    fields = get_fields(fid)
    rows = db_execute("""SELECT * FROM form_responses WHERE form_id=?
        AND status IN ('completed','confirmed') ORDER BY id""", (fid,), fetch="all") or []
    if not rows:
        return send(cid, "📭 پاسخی برای خروجی نیست")

    out = [f"گزارش فرم: {form['title']}",
           "=" * 50,
           f"تاریخ تهیه: {jnow()}",
           f"تعداد پاسخ: {len(rows)}",
           f"تعداد سوال: {len(fields)}",
           "=" * 50, ""]

    for n, r in enumerate(rows, 1):
        u = db_execute("SELECT first_name, last_name, mobile FROM users "
                       "WHERE bale_user_id=?", (r.get("bale_user_id"),), fetch="one") or {}
        nm = f"{u.get('first_name') or ''} {u.get('last_name') or ''}".strip() or "—"
        if form.get("anonymous"):
            nm = "ناشناس"
        out += [f"[{n}] {nm}",
                f"    کد پیگیری: {r.get('tracking_code') or '—'}",
                f"    تاریخ: {(r.get('completed_at') or '')[:19]}"]
        if r.get("amount"):
            out.append(f"    مبلغ: {fmt_amount(r['amount'])} {CURRENCY_LABEL} "
                       f"({r.get('payment_status')})")
        if r.get("max_score"):
            out.append(f"    نمره: {r.get('score')}/{r.get('max_score')}")

        amap = {}
        for a in db_execute("SELECT * FROM form_answers WHERE response_id=?",
                            (r["id"],), fetch="all") or []:
            v = a.get("value_text") or ""
            if not v and a.get("file_id"):
                v = f"[{a.get('file_type') or 'فایل'}]"
            amap[a.get("field_id")] = v
        for f in fields:
            out.append(f"    - {field_label(f)[:44]}: {amap.get(f['id'], '—')}")
        out += ["-" * 50, ""]

    from pathlib import Path
    d = Path(FILES_DIR) / "forms" / "exports"
    d.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    fp = d / f"form_{fid}_{stamp}.txt"
    try:
        fp.write_text("\n".join(out), encoding="utf-8")
    except Exception as e:
        log(f"❌ خروجی متنی: {e}")
        return send(cid, "❌ خطا در ساخت فایل")

    res = send_document(cid, str(fp), f"📑 گزارش متنی «{form['title'][:40]}»")
    if not res.get("ok"):
        send(cid, f"⚠️ ارسال ناموفق.\n📁 {fp}")
    return send(cid, "✅ خروجی آماده شد", {"inline_keyboard": [
        [{"text": "🔙 بازگشت", "callback_data": f"fbv_{fid}"}]]})


# ==================== پیش‌نمایش فرم ====================
def preview_form(cid, uid, fid):
    """👁️ پیش‌نمایش فرم آن‌طور که کاربر می‌بیند"""
    form = _guard(cid, uid, fid)
    if not form:
        return
    fields = get_fields(fid)
    if not fields:
        return send(cid, "📭 این فرم سوالی ندارد")

    t = FORM_TYPES.get(form.get("form_type") or "questionnaire",
                       FORM_TYPES["questionnaire"])
    lines = ["👁️ پیش‌نمایش فرم", "━━━━━━━━━━━━━━", "",
             f"{t['icon']} {form['title']}"]
    if form.get("description"):
        lines.append(f"\n{form['description']}")
    lines.append("\n━━━ سوالات ━━━\n")

    for i, f in enumerate(fields, 1):
        m = field_meta(f.get("field_type"))
        star = " *" if f.get("is_required") else ""
        lines.append(f"{i}. {m['icon']} {field_label(f)}{star}")
        if f.get("description"):
            lines.append(f"    📄 {f['description']}")
        if f.get("help_text"):
            lines.append(f"    ⚠️ {f['help_text']}")
        opts = parse_options(f)
        if opts:
            for o in opts[:6]:
                sym = "☑️" if m.get("kind") == "multi" else "⚪"
                lines.append(f"    {sym} {o}")
            if len(opts) > 6:
                lines.append(f"    … و {len(opts)-6} گزینه دیگر")
        if f.get("condition_field_id"):
            lines.append("    ◇ (نمایش مشروط)")
        lines.append("")

    lines.append("━━━━━━━━━━━━━━")
    if form.get("end_text"):
        lines.append(f"\n💚 پیام پایانی:\n{form['end_text']}")
    if (form.get("payment_method") or "none") != "none":
        lines.append(f"\n💳 هزینه: {fmt_amount(form.get('price'))} {CURRENCY_LABEL}")

    send_or_edit(cid, "\n".join(lines), {"inline_keyboard": [
        [{"text": "🧪 تست واقعی (پر کردن فرم)", "callback_data": f"ffopen_{fid}"}],
        [{"text": "📝 ویرایش سوالات", "callback_data": f"fbf_{fid}"}],
        [{"text": "🔙 بازگشت", "callback_data": f"fbv_{fid}"}]]})


# ==================== ارسال همگانی لینک فرم ====================
def broadcast_form(cid, uid, fid):
    """📢 ارسال لینک فرم به همه کاربران"""
    form = _guard(cid, uid, fid)
    if not form:
        return
    if form.get("status") != "active":
        return send_or_edit(cid, (
            "⚠️ ابتدا فرم را منتشر کنید!"
        ), {"inline_keyboard": [
            [{"text": "📌 انتشار فرم", "callback_data": f"fbpub_{fid}"}],
            [{"text": "🔙 بازگشت", "callback_data": f"fblink_{fid}"}]]})

    n = db_scalar("SELECT COUNT(*) FROM users WHERE bale_user_id IS NOT NULL "
                  "AND (is_blocked IS NULL OR is_blocked=0)")
    from forms import form_link
    return send_or_edit(cid, (
        f"📢 ارسال همگانی فرم\n━━━━━━━━━━━━━━\n\n"
        f"📋 {form['title']}\n"
        f"👥 گیرندگان: {n} کاربر\n"
        f"🔗 {form_link(form)}\n\n"
        f"⚠️ پیام برای همه ارسال می‌شود. مطمئنید؟"
    ), {"inline_keyboard": [
        [{"text": "📢 بله، ارسال کن", "callback_data": f"fbcastgo_{fid}"}],
        [{"text": "❌ انصراف", "callback_data": f"fblink_{fid}"}]]})


def do_broadcast_form(cid, uid, fid):
    form = _guard(cid, uid, fid)
    if not form:
        return
    import time as _t
    from config import BROADCAST_RATE
    from forms import form_link, FORM_TYPES as FT

    t = FT.get(form.get("form_type") or "questionnaire", FT["questionnaire"])
    targets = db_execute("SELECT bale_user_id FROM users WHERE bale_user_id IS NOT NULL "
                         "AND (is_blocked IS NULL OR is_blocked=0)", fetch="all") or []
    if not targets:
        return send(cid, "📭 کاربری نیست")

    body = [f"{t['icon']} {form['title']}", "━━━━━━━━━━━━━━", ""]
    if form.get("description"):
        body.append(form["description"][:300])
    if (form.get("payment_method") or "none") != "none" and form.get("price"):
        body.append(f"\n💳 هزینه: {fmt_amount(form['price'])} {CURRENCY_LABEL}")
    else:
        body.append("\n🎁 رایگان")
    if form.get("end_at"):
        body.append(f"⛔ مهلت: {form['end_at']}")
    text = "\n".join(body)
    kb = {"inline_keyboard": [
        [{"text": "▶️ شرکت در فرم", "callback_data": f"ffopen_{fid}"}]]}

    send(cid, f"📤 در حال ارسال به {len(targets)} کاربر...")
    ok = fail = 0
    delay = 1.0 / max(1, BROADCAST_RATE)
    for i, r in enumerate(targets, 1):
        try:
            res = send(r["bale_user_id"], text, kb)
            if res.get("ok"):
                ok += 1
            else:
                fail += 1
        except Exception:
            fail += 1
        if i % 50 == 0:
            send(cid, f"📊 {i}/{len(targets)} — ✅ {ok} | ❌ {fail}")
        _t.sleep(delay)

    return send(cid, (
        f"✅ ارسال تمام شد\n━━━━━━━━━━━━━━\n\n"
        f"📤 موفق: {ok}\n❌ ناموفق: {fail}"
    ), {"inline_keyboard": [
        [{"text": "🔙 بازگشت", "callback_data": f"fbv_{fid}"}]]})


# ══════════════════════════════════════════════════════════
# 🔀 فاز ۱۷ — گردش کار: صف تأیید پاسخ‌ها
#
# پاسخ‌هایی که فرمشان needs_approval دارد اینجا بررسی می‌شوند.
# ══════════════════════════════════════════════════════════
def show_approvals(cid, uid, page=0):
    """🔀 صف تأیید پاسخ‌های فرم"""
    from auth import is_admin as _ia
    if not _ia(uid):
        return send(cid, "❌ دسترسی ندارید")
    import permissions as perm
    if not perm.has_any(uid, "fm_responses", "fm_edit"):
        return send(cid, "🔒 دسترسی به پاسخ‌ها ندارید")

    from config import SUPER_ADMIN_ID as _S
    if uid == _S:
        rows = db_execute("""SELECT r.*, f.title ftitle, f.owner_id,
            u.first_name, u.last_name, u.mobile
            FROM form_responses r
            JOIN forms f ON f.id=r.form_id
            LEFT JOIN users u ON u.bale_user_id=r.bale_user_id
            WHERE r.status='pending_review'
            ORDER BY r.id DESC LIMIT 50""", fetch="all") or []
    else:
        rows = db_execute("""SELECT r.*, f.title ftitle, f.owner_id,
            u.first_name, u.last_name, u.mobile
            FROM form_responses r
            JOIN forms f ON f.id=r.form_id
            LEFT JOIN users u ON u.bale_user_id=r.bale_user_id
            WHERE r.status='pending_review' AND f.owner_id=?
            ORDER BY r.id DESC LIMIT 50""", (uid,), fetch="all") or []

    lines = ["🔀 صف تأیید پاسخ‌ها", "━━━━━━━━━━━━━━",
             f"⏳ {len(rows)} پاسخ منتظر بررسی شماست", ""]
    btns = []
    if not rows:
        lines.append("✅ پاسخ منتظری نیست")
    for r in rows[:10]:
        nm = f"{r.get('first_name') or ''} {r.get('last_name') or ''}".strip()
        lines += [f"⏳ {str(r.get('ftitle'))[:28]}",
                  f"   👤 {nm or '—'}  📱 {r.get('mobile') or '—'}",
                  f"   🔖 {r.get('tracking_code') or '—'}"
                  f"   🕐 {_jd.fmt(r.get('created_at'), 'datetime')}", ""]
        btns.append([
            {"text": f"👁️ {nm or r['id']}"[:20],
             "callback_data": f"fbrv_{r['id']}"},
            {"text": "✅", "callback_data": f"fwok_{r['id']}"},
            {"text": "❌", "callback_data": f"fwno_{r['id']}"},
        ])
    btns.append([{"text": "🔙 فرم‌ساز", "callback_data": "fb_panel"}])
    send_or_edit(cid, "\n".join(lines), {"inline_keyboard": btns})


def review_response(cid, uid, rid, approve=True):
    """✅/❌ تأیید یا رد یک پاسخ"""
    from auth import is_admin as _ia, log_action as _log
    if not _ia(uid):
        return send(cid, "❌ دسترسی ندارید")
    try:
        rid = int(rid)
    except (TypeError, ValueError):
        return send(cid, "❌ شناسه نامعتبر")

    r = db_execute("""SELECT r.*, f.title ftitle, f.owner_id, f.event_id
        FROM form_responses r JOIN forms f ON f.id=r.form_id
        WHERE r.id=?""", (rid,), fetch="one")
    if not r:
        return send(cid, "❌ پاسخ یافت نشد")

    from config import SUPER_ADMIN_ID as _S
    if uid != _S and r.get("owner_id") != uid:
        import permissions as perm
        if not perm.has_perm(uid, "fm_responses"):
            return send(cid, "🔒 این فرم مال شما نیست")

    if r.get("status") != "pending_review":
        return send(cid, f"ℹ️ قبلاً بررسی شده ({r.get('status')})")

    new = "completed" if approve else "rejected"
    # ⚠️ form_responses ستون updated_at ندارد — reviewed_by/at دارد
    db_execute("""UPDATE form_responses SET status=?, reviewed_by=?,
        reviewed_at=CURRENT_TIMESTAMP WHERE id=?""", (new, uid, rid))
    _log(uid, "form_reviewed", {"rid": rid, "ok": approve})

    target = r.get("bale_user_id")
    if target:
        try:
            if approve:
                send(target, (
                    "✅ پاسخ شما تأیید شد\n━━━━━━━━━━━━━━\n\n"
                    f"📋 {r.get('ftitle')}\n"
                    f"🔖 {r.get('tracking_code') or '—'}\n\n"
                    "💚 سپاس از شما."))
            else:
                send(target, (
                    "❌ پاسخ شما تأیید نشد\n━━━━━━━━━━━━━━\n\n"
                    f"📋 {r.get('ftitle')}\n\n"
                    "💡 برای اطلاعات بیشتر با پشتیبانی تماس بگیرید."))
        except Exception as e:
            log(f"⚠️ اعلام نتیجه بررسی: {e}")

    # 🔗 اگر فرم به رویدادی وصل است و تأیید شد → ثبت‌نام ادامه یابد
    if approve and r.get("event_id") and target:
        try:
            import events as _ev
            ev = _ev.get_event(r["event_id"])
            if ev and ev.get("status") in ("published", "closed"):
                exists = db_execute(
                    "SELECT id FROM event_registrations "
                    "WHERE event_id=? AND bale_user_id=?",
                    (ev["id"], target), fetch="one")
                if not exists:
                    _ev.register(target, target, ev["id"])
        except Exception as e:
            log(f"⚠️ ثبت‌نام پس از تأیید: {e}")

    # 🔗 وبهوک
    try:
        from form_fill import _fire_webhook
        form = db_execute("SELECT * FROM forms WHERE id=?",
                          (r["form_id"],), fetch="one")
        _fire_webhook(form, r, "approved" if approve else "rejected")
    except Exception as e:
        log(f"⚠️ وبهوک بررسی: {e}")

    send(cid, "✅ تأیید شد" if approve else "❌ رد شد")
    return show_approvals(cid, uid)
