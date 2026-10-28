#!/usr/bin/env python3
"""فاز ۹.۵ §۷.۴ — آزمون تکرارپذیر «PDF واقعی گواهی» روی مسیر واقعی ربات (آفلاین، بدون شبکه).

چه می‌کند
---------
* یک پایگاه موقت، یک قالب وردِ فارسی با جای‌گذارهای رسمی مرکز صدور
  (``{نام}`` ``{رویداد}`` ``{شماره}`` ``{تاریخ صدور}`` ``{کیوآر}``) می‌سازد؛
* از **همان مسیر واقعی** ``certcenter`` (طرح → نوع «پایان محتوای ضبط‌شده» → گیرندهٔ
  دستی → انتشار → ``build_one``) گواهی می‌سازد — ``certgen.generate`` جایگزین نمی‌شود؛
* فایل PDF خروجی را ساختاری بررسی می‌کند (سرآیند، ``%%EOF``، شمار صفحه، فونت/تصویر)،
  جای‌گذاری فارسی را در DOCX و (در صورت وجود ``pdftotext``) در متن PDF می‌سنجد؛
* صدور دوباره (بدون ردیف تکراری) و ابطال/بازگردانی با همگام‌سازی آینه را اجرا می‌کند؛
* نسخهٔ ابزارها را گزارش می‌کند.

نتیجه
-----
خروجی JSON روی stdout. کد خروج:
  0 = موتور LibreOffice موجود بود و همهٔ بررسی‌ها گذشت (PDF واقعی با لایهٔ متن)؛
  2 = LibreOffice در این محیط نیست؛ مسیر و PDF جایگزین (تصویری) آزموده شد ولی
      «PDF با لایهٔ متن» هنوز **نیازمند staging** است؛
  1 = خطای واقعی (ساخت ناموفق، فایل نامعتبر، ردیف تکراری، ابطال بی‌اثر…).

اجرا::

    cd Final-Delivery/bale-bot && python tools/certificate-pdf-check.py [--keep] [--json]

هیچ پیام بله فرستاده نمی‌شود (``bale_api`` در حافظه جایگزین می‌شود) و هیچ چیزی در
پایگاه عملیاتی نوشته نمی‌شود؛ همه‌چیز در یک پوشهٔ موقت است.
"""
from __future__ import annotations

import argparse
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

BOT_ROOT = Path(__file__).resolve().parents[1] / "nora_bot"

ADMIN = 7300001
HOLDER = 7301001
HOLDER_NAME = "مریم آزمایشی"
ISSUE_TITLE = "پایان دورهٔ ضبط‌شدهٔ نجوم مقدماتی (آزمون ۹.۵)"


def _prepare_env(workdir: Path) -> None:
    """متغیرهای محیطی حداقلی — پیش از import ماژول‌های ربات."""
    os.environ.update({
        "BOT_TOKEN": "1095412158:TEST_TOKEN_FOR_PDF_CHECK",
        "SUPER_ADMIN_ID": str(ADMIN),
        "DB_PATH": str(workdir / "cert-check.db"),
        "BACKUP_DIR": str(workdir / "backups"),
        "IMPORT_DIR": str(workdir / "imports"),
        "FILES_DIR": str(workdir / "files"),
        "PROFILES_DIR": str(workdir / "files/profiles"),
        "EVENTS_DIR": str(workdir / "files/events"),
        "LOG_DIR": str(workdir / "logs"),
    })
    if str(BOT_ROOT) not in sys.path:
        sys.path.insert(0, str(BOT_ROOT))


def _silence_network(bale_api: Any) -> list[tuple[str, dict[str, Any]]]:
    sent: list[tuple[str, dict[str, Any]]] = []

    def fake_raw_call(method, params=None, **_kw):
        sent.append((str(method), dict(params or {})))
        return {"ok": True, "result": {"message_id": len(sent), "username": "nora_qa_bot", "status": "member"}}

    def capture_send(cid, text, kb=None, *_a, **_k):
        sent.append(("sendMessage", {"chat_id": cid, "text": text}))
        return {"ok": True, "result": {"message_id": len(sent)}}

    bale_api._raw_call = fake_raw_call
    bale_api.api_async = lambda m, p=None: {"ok": True}
    bale_api.send = capture_send
    bale_api.send_or_edit = capture_send
    return sent


def make_template(path: Path) -> None:
    """قالب وردِ فارسی با جای‌گذارهای رسمی (راست‌به‌چپ)."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = docx.Document()
    lines = (
        "گواهی پایان محتوای آموزشی ضبط‌شده",
        "بدین‌وسیله گواهی می‌شود {نام} محتوای آموزشی «{رویداد}» را به پایان رسانده است.",
        "شمارهٔ گواهی: {شماره}",
        "تاریخ صدور: {تاریخ صدور}",
        "{کیوآر}",
    )
    for text in lines:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        bidi = OxmlElement("w:bidi")
        paragraph._p.get_or_add_pPr().append(bidi)
        run = paragraph.add_run(text)
        rtl = OxmlElement("w:rtl")
        run._r.get_or_add_rPr().append(rtl)
        run.font.name = "Vazirmatn"
        rpr = run._r.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        fonts.set(qn("w:cs"), "Vazirmatn")
    document.save(path)


def inspect_pdf(path: Path) -> dict[str, Any]:
    data = path.read_bytes()
    tail = data[-2048:]
    pages = len(re.findall(rb"/Type\s*/Page(?![s/])", data))
    return {
        "bytes": len(data),
        "header": data[:8].decode("latin-1", "replace"),
        "valid_header": data.startswith(b"%PDF-1."),
        "eof_marker": b"%%EOF" in tail,
        "pages": pages,
        "has_font": b"/Font" in data,
        "has_image": b"/Image" in data,
        "has_text_operator": b"BT" in data and b"ET" in data,
    }


def docx_text(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def pdf_text(path: Path) -> str | None:
    exe = shutil.which("pdftotext")
    if not exe:
        return None
    try:
        out = subprocess.run([exe, "-layout", str(path), "-"], capture_output=True, timeout=60, check=False)
    except (OSError, subprocess.SubprocessError):
        return None
    return out.stdout.decode("utf-8", "replace") if out.returncode == 0 else None


def tool_versions(certgen: Any) -> dict[str, Any]:
    versions: dict[str, Any] = {"python": platform.python_version(), "platform": platform.platform()}
    try:
        import docx  # noqa: PLC0415

        versions["python_docx"] = getattr(docx, "__version__", "installed")
    except Exception:  # noqa: BLE001
        versions["python_docx"] = None
    try:
        import PIL  # noqa: PLC0415
        from PIL import features  # noqa: PLC0415

        versions["pillow"] = PIL.__version__
        versions["raqm_text_shaping"] = bool(features.check("raqm"))
    except Exception:  # noqa: BLE001
        versions["pillow"] = None
    lo = certgen.libreoffice_bin()
    versions["libreoffice_path"] = str(lo) if lo else None
    versions["libreoffice_version"] = None
    if lo:
        try:
            out = subprocess.run([str(lo), "--version"], capture_output=True, timeout=60, check=False)
            versions["libreoffice_version"] = out.stdout.decode("utf-8", "replace").strip()[:120] or None
        except (OSError, subprocess.SubprocessError):
            pass
    versions["pdftoppm"] = shutil.which("pdftoppm")
    versions["pdftotext"] = shutil.which("pdftotext")
    return versions


def run(workdir: Path) -> dict[str, Any]:
    """اجرای کامل؛ خروجی گزارش. استثناها به عنوان خطای واقعی بالا می‌روند."""
    _prepare_env(workdir)
    import bale_api  # noqa: PLC0415

    sent = _silence_network(bale_api)
    import attendance  # noqa: PLC0415
    import auth  # noqa: PLC0415
    import certcenter as cc  # noqa: PLC0415
    import certgen  # noqa: PLC0415
    import database  # noqa: PLC0415
    import shared_certificates as sc  # noqa: PLC0415

    database.close_db()
    database.DB_PATH = os.environ["DB_PATH"]
    database._conn = None
    database.init_db()
    attendance.init_attendance_tables()
    certgen.init_cert_tables()
    cc.init_tables()
    database.seed_default_settings()
    database.db_execute(
        "INSERT INTO users (bale_user_id, first_name, last_name, is_authenticated, is_admin) VALUES (?,?,?,1,1)",
        (ADMIN, "مدیر", "آزمایشی"))
    database.db_execute(
        "INSERT INTO users (bale_user_id, first_name, last_name, is_authenticated, national_id) VALUES (?,?,?,1,?)",
        (HOLDER, HOLDER_NAME.split()[0], HOLDER_NAME.split()[1], "0012345678"))
    assert auth.is_admin(ADMIN)

    template_path = workdir / "template-fa.docx"
    make_template(template_path)
    tid = database.db_execute(
        "INSERT INTO cert_templates (name, file_path, kind, is_active, is_default) VALUES (?,?,?,1,1)",
        ("قالب فارسی آزمون ۹.۵", str(template_path), "docx"))
    assert tid

    checks: list[dict[str, Any]] = []

    def check(name: str, ok: bool, **info: Any) -> None:
        checks.append({"check": name, "ok": bool(ok), **info})

    lo = certgen.libreoffice_bin()
    report: dict[str, Any] = {
        "tool": "certificate-pdf-check", "offline": True,
        "libreoffice_available": bool(lo), "versions": tool_versions(certgen), "checks": checks,
    }

    # ── مسیر واقعی مرکز صدور ──
    iid = cc.create_issue(ADMIN, ISSUE_TITLE, tid)
    cc.set_kind(ADMIN, ADMIN, iid, "recording_completion")
    check("issue_kind_recording_completion", cc.kind_of(cc.get_issue(iid)) == "recording_completion")
    # پارامترهای پرسیدنی قالب از خودِ فایل ورد خوانده می‌شوند (scan واقعی)؛ «رویداد» ثابت طرح است.
    ask, auto, images = cc.template_fields(tid)
    check("template_scan_real_docx", "event" in ask and "name" in auto and "qr" in images, ask=ask, auto=auto, images=images)
    cc.set_value(iid, "event", ISSUE_TITLE)
    cc.add_share(iid, uid=HOLDER, name=HOLDER_NAME)
    ready, todo = cc.issue_ready(cc.get_issue(iid))
    check("issue_ready", bool(ready), todo=todo if not ready else "")
    _n, err = cc.publish(iid, notify=False)
    check("publish_silent", err == "", error=err)
    share = cc.shares(iid)[0]

    cert, err = cc.build_one(iid, share["id"])
    check("build_one_real_generate", bool(cert) and err == "", error=err)
    if not cert:
        return report
    report["engine"] = None
    files = {k: cert.get(k) for k in ("png_path", "pdf_path", "docx_path")}
    report["outputs"] = {k: (Path(v).name if v else None) for k, v in files.items()}
    check("certificate_row_source_recording", cert.get("source") == "recording" and cert.get("event_id") is None)
    check("serial_and_codes_present", bool(cert.get("serial")) and bool(cert.get("full_code")))

    # DOCX: جای‌گذاری فارسی واقعی — نام و عنوان حاضر، هیچ آکولادی نمانده
    docx_path = files.get("docx_path")
    if docx_path and Path(docx_path).exists():
        text = docx_text(Path(docx_path))
        check("docx_persian_substitution", HOLDER_NAME in text and ISSUE_TITLE in text and "{" not in text,
              name_present=HOLDER_NAME in text, title_present=ISSUE_TITLE in text, braces_left=text.count("{"))
    else:
        check("docx_persian_substitution", False, reason="docx output missing")

    # PDF: ساختار معتبر
    pdf_path = files.get("pdf_path")
    if pdf_path and Path(pdf_path).exists():
        info = inspect_pdf(Path(pdf_path))
        report["pdf"] = info
        check("pdf_structure_valid", info["valid_header"] and info["eof_marker"] and info["pages"] >= 1, **info)
        engine = "office" if (lo and info["has_font"] and info["has_text_operator"]) else ("builtin_image" if info["has_image"] else "unknown")
        report["engine"] = engine
        if lo:
            check("pdf_has_text_layer_office_engine", info["has_font"] and info["has_text_operator"], engine=engine)
            extracted = pdf_text(Path(pdf_path))
            if extracted is None:
                check("pdf_text_contains_persian_name", False, skipped=True, reason="pdftotext not installed")
            else:
                letters = [ch for ch in HOLDER_NAME if ch.strip()]
                check("pdf_text_contains_persian_name", all(ch in extracted for ch in letters),
                      note="RTL order may differ in pdftotext output; character presence is checked")
        else:
            check("pdf_has_text_layer_office_engine", False, skipped=True,
                  reason="LibreOffice not installed: image-based fallback PDF only → نیازمند staging")
    else:
        report["pdf"] = None
        check("pdf_structure_valid", False, reason="pdf output missing")

    png_path = files.get("png_path")
    if png_path and Path(png_path).exists():
        from PIL import Image  # noqa: PLC0415

        with Image.open(png_path) as image:
            w, h = image.size
            # صفحه‌ای که واقعاً چیزی روی آن رسم شده (نه یک بوم سفید)
            grey = image.convert("L")
            dark = int(sum(1 for px in grey.tobytes() if px < 128))
        check("png_rendered_non_blank", w > 400 and h > 400 and dark > 500, width=w, height=h, dark_pixels=dark)
    else:
        check("png_rendered_non_blank", False, reason="png output missing")

    # صدور دوباره: همان ردیف، نه ردیف دوم؛ فایل‌ها دوباره ساخته نمی‌شوند
    again, err = cc.build_one(iid, share["id"])
    count = database.db_scalar("SELECT COUNT(*) FROM certificates WHERE bale_user_id=?", (HOLDER,))
    check("retry_is_idempotent", err == "" and again and again["id"] == cert["id"] and count == 1, rows=count)

    # ابطال/بازگردانی با همگام‌سازی آینه (آینه در حافظه ثبت می‌شود؛ شبکه‌ای نیست)
    synced: list[int] = []
    original_sync = sc.sync_after_change
    sc.sync_after_change = lambda cid: synced.append(int(cid))
    try:
        cc.revoke_share(ADMIN, ADMIN, share["id"])
        revoked = int(database.db_scalar("SELECT revoked FROM certificates WHERE id=?", (cert["id"],)) or 0)
        run_ = database.db_execute
        status = sc.build_payload(run_, sc.load_certificate(run_, cert["id"]), source="bot", source_key="k")["status"]
        check("revoke_sets_flag_and_syncs_mirror", revoked == 1 and status == "revoked" and synced == [cert["id"]],
              revoked=revoked, mirror_status=status, synced=len(synced))
        cc.revoke_share(ADMIN, ADMIN, share["id"], undo=True)
        restored = int(database.db_scalar("SELECT revoked FROM certificates WHERE id=?", (cert["id"],)) or 0)
        check("restore_clears_flag_and_syncs_mirror", restored == 0 and len(synced) == 2, revoked=restored)
        still = database.db_scalar("SELECT COUNT(*) FROM certificates")
        check("revoke_never_deletes_rows", still == 1, rows=still)
    finally:
        sc.sync_after_change = original_sync

    audit = database.db_execute(
        "SELECT action_type FROM user_actions WHERE action_type IN ('cert_revoked','cert_restored') ORDER BY id", fetch="all") or []
    check("revoke_restore_audited", [r["action_type"] for r in audit] == ["cert_revoked", "cert_restored"])
    # پاسخ‌های پنل به مدیر (در حافظه) مجازند؛ هیچ پیام یا فایلی برای دارنده فرستاده نمی‌شود.
    to_holder = [m for m, p in sent if str(p.get("chat_id")) == str(HOLDER)]
    files_sent = [m for m, _p in sent if m in ("sendDocument", "sendPhoto")]
    check("no_bale_delivery_to_holder", not to_holder and not files_sent, calls=len(sent), to_holder=len(to_holder))
    database.close_db()
    return report


def verdict(report: dict[str, Any]) -> tuple[int, str]:
    failed = [c["check"] for c in report.get("checks", []) if not c["ok"] and not c.get("skipped")]
    if failed:
        return 1, "FAIL: " + ", ".join(failed)
    if not report.get("libreoffice_available"):
        return 2, "STAGING_REQUIRED: LibreOffice not available here; real-path + fallback PDF verified, text-layer PDF untested"
    return 0, "OK: real text-layer PDF produced by LibreOffice on the real certificate path"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--keep", action="store_true", help="keep the temporary work directory (prints its path to stderr)")
    parser.add_argument("--json", action="store_true", help="JSON only (no human summary line on stderr)")
    args = parser.parse_args(argv)
    workdir = Path(tempfile.mkdtemp(prefix="nora-cert-pdf-check-"))
    try:
        import contextlib  # noqa: PLC0415

        # لاگ‌های ربات روی stdout چاپ می‌شوند؛ تا JSON سالم بماند به stderr می‌روند.
        with contextlib.redirect_stdout(sys.stderr):
            report = run(workdir)
        code, summary = verdict(report)
        report["exit_code"] = code
        report["summary"] = summary
        print(json.dumps(report, ensure_ascii=False, indent=2))
        if not args.json:
            print(summary, file=sys.stderr)
        return code
    finally:
        if args.keep:
            print(f"work directory kept: {workdir}", file=sys.stderr)
        else:
            shutil.rmtree(workdir, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
